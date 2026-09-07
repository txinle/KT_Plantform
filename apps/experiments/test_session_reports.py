from io import BytesIO
from unittest.mock import patch
from urllib.parse import quote

from django.test import Client, TestCase
from django.urls import reverse
from docx import Document

from apps.accounts.models import ClassGroup, Course, CourseEnrollment, CourseTeacher, User
from apps.assessment.models import Attempt, ExperimentSession, StudentQuestionAssignment
from apps.audit.models import AuditLog
from apps.experiments.models import Experiment, ExperimentVersion, QuestionSnapshot
from apps.experiments.presentation import (
    OptionPresentation,
    answer_option_values,
    display_text,
    format_answer,
)
from apps.experiments.presentation_instances import get_or_create_default_item_instance
from apps.question_bank.models import Question, QuestionBank


class TrackingBytesIO(BytesIO):
    instances = []

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.was_closed = False
        self.instances.append(self)

    def close(self):
        self.was_closed = True
        super().close()


class CourseExperimentSessionReportTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(
            username="teacher",
            password="password123",
            role=User.Role.TEACHER,
            display_name="课程负责人",
        )
        self.collaborator = User.objects.create_user(
            username="collaborator",
            password="password123",
            role=User.Role.TEACHER,
        )
        self.outsider = User.objects.create_user(
            username="outsider",
            password="password123",
            role=User.Role.TEACHER,
        )
        self.admin = User.objects.create_user(
            username="admin",
            password="password123",
            role=User.Role.ADMIN,
        )
        self.superuser = User.objects.create_superuser(
            username="root",
            password="password123",
            role=User.Role.ADMIN,
        )
        self.course = Course.objects.create(
            owner=self.teacher,
            name="学习分析研究方法",
            code="LA-2026",
            term="2026 秋季",
        )
        CourseTeacher.objects.create(course=self.course, teacher=self.collaborator)
        self.class_group = ClassGroup.objects.create(
            course=self.course,
            owner=self.teacher,
            name="2026 实验班",
            academic_year="2026",
        )
        self.student = User.objects.create_user(
            username="S-001",
            password="password123",
            student_number="S-001",
            display_name="张小雨",
            role=User.Role.STUDENT,
        )
        CourseEnrollment.objects.create(
            course=self.course,
            class_group=self.class_group,
            student=self.student,
        )
        self.bank = QuestionBank.objects.create(
            course=self.course,
            created_by=self.teacher,
            name="报告题库",
        )
        self.experiment = Experiment.objects.create(
            course=self.course,
            created_by=self.teacher,
            question_bank=self.bank,
            name="课程会话报告实验",
        )
        self.version = ExperimentVersion.objects.create(
            experiment=self.experiment,
            version_number=7,
            published_by=self.teacher,
        )
        self.session = ExperimentSession.objects.create(
            version=self.version,
            student=self.student,
            class_snapshot_name="2026 实验班（发布快照）",
            status=ExperimentSession.Status.IN_PROGRESS,
            submission_count=1,
        )
        self._question_sequence = 0
        self.multiple_snapshot = self._create_snapshot(
            question_type=Question.Type.MULTIPLE,
            text="多选题：下列哪些描述符合知识追踪？",
            options=[
                {"key": "B", "text": "可以结合历史交互预测下一步表现"},
                {"key": "A", "text": "只记录最终分数，不关心过程"},
                {"key": "C", "text": "可用于诊断知识状态"},
            ],
            correct_answer=["B", "C"],
            stable_code="KT-C-001",
            code="BANK-001",
        )
        self.true_false_snapshot = self._create_snapshot(
            question_type=Question.Type.TRUE_FALSE,
            text="判断题：正式提交后仍可修改答案。",
            options=[],
            correct_answer=False,
            stable_code="",
            code="LEGACY-TF-002",
        )
        self.unanswered_snapshot = self._create_snapshot(
            question_type=Question.Type.SHORT,
            text="简答题：请概述你对题目序列的理解。",
            options=[],
            correct_answer="按分配顺序展示。",
            stable_code="KT-C-003",
            code="BANK-003",
        )
        self.submitted_assignment = self._create_assignment(
            self.multiple_snapshot,
            2,
            status=Attempt.Status.SUBMITTED,
            draft_answer=["A"],
            final_answer=["B", "C"],
            is_correct=True,
        )
        self.draft_assignment = self._create_assignment(
            self.true_false_snapshot,
            1,
            status=Attempt.Status.DRAFT,
            draft_answer=True,
            final_answer=None,
            is_correct=None,
        )
        self.unanswered_assignment = self._create_assignment(
            self.unanswered_snapshot,
            3,
            status=Attempt.Status.UNANSWERED,
            draft_answer=None,
            final_answer=None,
            is_correct=None,
        )
        self.detail_url = (
            f"/teacher/courses/{self.course.pk}/experiments/{self.experiment.pk}/"
            f"sessions/{self.session.pk}/"
        )
        self.word_url = f"{self.detail_url}word/"
        self.client.force_login(self.teacher)

    def _create_snapshot(
        self,
        *,
        question_type,
        text,
        options,
        correct_answer,
        stable_code,
        code,
    ):
        self._question_sequence += 1
        source = Question.objects.create(
            bank=self.bank,
            sequence=self._question_sequence,
            question_type=question_type,
            text=text,
            correct_answer=correct_answer,
        )
        return QuestionSnapshot.objects.create(
            version=self.version,
            source_question=source,
            source_sequence=self._question_sequence,
            stable_code=stable_code,
            code=code,
            question_type=question_type,
            text=text,
            options=options,
            correct_answer=correct_answer,
            primary_kc="不得出现在报告中",
            all_kcs=["不得出现在报告中"],
            explanation="不得出现在报告中",
            teacher_difficulty="0.500",
            empirical_difficulty="0.500",
        )

    def _create_assignment(
        self,
        snapshot,
        assigned_sequence_index,
        *,
        status,
        draft_answer,
        final_answer,
        is_correct,
    ):
        assignment = StudentQuestionAssignment.objects.create(
            session=self.session,
            snapshot=snapshot,
            item_instance=get_or_create_default_item_instance(snapshot),
            assigned_sequence_index=assigned_sequence_index,
        )
        Attempt.objects.create(
            assignment=assignment,
            session=self.session,
            status=status,
            draft_answer=draft_answer,
            final_answer=final_answer,
            is_correct=is_correct,
        )
        return assignment

    def test_detail_shows_real_session_content_in_assigned_order(self):
        response = self.client.get(self.detail_url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            reverse(
                "courses:experiment_session_detail",
                args=[self.course.pk, self.experiment.pk, self.session.pk],
            ),
            self.detail_url,
        )
        self.assertContains(response, "张小雨")
        self.assertContains(response, "S-001")
        self.assertContains(response, "2026 实验班（发布快照）")
        self.assertContains(response, "v7")
        self.assertContains(response, "进行中")
        self.assertContains(response, "实际题数：3")
        self.assertLess(
            response.content.find("LEGACY-TF-002".encode()),
            response.content.find("KT-C-001".encode()),
        )
        self.assertLess(
            response.content.find("KT-C-001".encode()),
            response.content.find("KT-C-003".encode()),
        )
        self.assertContains(response, "正确")
        self.assertContains(response, "错误")
        self.assertContains(response, "草稿，未正式提交")
        self.assertContains(response, "未作答")
        self.assertContains(response, "回答正确")
        self.assertContains(response, "B：可以结合历史交互预测下一步表现")
        self.assertContains(response, "C：可用于诊断知识状态")
        questions = response.context["report"].questions
        self.assertEqual([question.assigned_sequence_index for question in questions], [1, 2, 3])
        draft, submitted, unanswered = questions
        self.assertEqual(draft.stable_code, "LEGACY-TF-002")
        self.assertEqual(draft.student_answer, "正确")
        self.assertEqual(draft.correct_answer, "错误")
        self.assertEqual(draft.answer_status, "草稿，未正式提交")
        self.assertEqual(
            submitted.student_answer,
            "B：可以结合历史交互预测下一步表现；C：可用于诊断知识状态",
        )
        self.assertNotIn("A：只记录最终分数，不关心过程", submitted.student_answer)
        self.assertEqual(
            [option.display for option in submitted.options],
            [
                "B：可以结合历史交互预测下一步表现",
                "A：只记录最终分数，不关心过程",
                "C：可用于诊断知识状态",
            ],
        )
        self.assertEqual(unanswered.student_answer, "未作答")
        self.assertEqual(unanswered.answer_status, "未作答")

    def test_presentation_escapes_html_at_the_template_boundary_and_serializes_json_text(self):
        self.multiple_snapshot.text = "<script>不应执行</script>\x01"
        self.multiple_snapshot.save(update_fields=["text"])

        response = self.client.get(self.detail_url)

        self.assertContains(response, "&lt;script&gt;不应执行&lt;/script&gt;", html=False)
        self.assertNotContains(response, "<script>不应执行</script>", html=False)
        self.assertEqual(
            format_answer(
                "short_answer",
                {"answer": ["<b>原样文本</b>"]},
                (OptionPresentation(key="A", text="选项"),),
            ),
            '{"answer": ["<b>原样文本</b>"]}',
        )
        self.assertNotIn("\x00", display_text({"answer": "\x00安全文本"}))
        self.assertEqual(display_text("\ud800安全文本"), "安全文本")

    def test_answer_option_values_share_choice_and_true_false_normalization(self):
        self.assertEqual(
            answer_option_values("multiple_choice", ["c", "A", "C"]),
            ("C", "A"),
        )
        self.assertEqual(answer_option_values("true_false", "正确"), ("true",))
        self.assertEqual(answer_option_values("true_false", False), ("false",))
        self.assertEqual(answer_option_values("short_answer", "A"), ())

    def test_results_include_a_link_to_the_session_detail(self):
        response = self.client.get(
            reverse("courses:experiment_results", args=[self.course.pk, self.experiment.pk])
        )

        self.assertContains(response, "查看作答详情")
        self.assertContains(response, self.detail_url)

    def test_owner_admin_and_superuser_can_open_detail_but_collaborator_and_outsider_get_404(self):
        for user in (self.teacher, self.admin, self.superuser):
            with self.subTest(user=user.username):
                self.client.force_login(user)
                self.assertEqual(self.client.get(self.detail_url).status_code, 200)

        for user in (self.collaborator, self.outsider):
            with self.subTest(user=user.username):
                self.client.force_login(user)
                self.assertEqual(self.client.get(self.detail_url).status_code, 404)

    def test_detail_scope_rejects_session_from_another_experiment_or_course(self):
        other_experiment = Experiment.objects.create(
            course=self.course,
            created_by=self.teacher,
            question_bank=self.bank,
            name="同课程的另一实验",
        )
        other_version = ExperimentVersion.objects.create(
            experiment=other_experiment,
            version_number=1,
            published_by=self.teacher,
        )
        other_session = ExperimentSession.objects.create(
            version=other_version,
            student=self.student,
        )
        same_course_cross_experiment_url = (
            f"/teacher/courses/{self.course.pk}/experiments/{self.experiment.pk}/"
            f"sessions/{other_session.pk}/"
        )

        other_teacher = User.objects.create_user(
            username="other-teacher",
            password="password123",
            role=User.Role.TEACHER,
        )
        other_course = Course.objects.create(
            owner=other_teacher,
            name="另一门课程",
            code="OTHER-2026",
            term="2026 秋季",
        )
        cross_course_url = (
            f"/teacher/courses/{other_course.pk}/experiments/{self.experiment.pk}/"
            f"sessions/{self.session.pk}/"
        )

        self.assertEqual(self.client.get(same_course_cross_experiment_url).status_code, 404)
        self.assertEqual(self.client.get(cross_course_url).status_code, 404)
        self.assertEqual(self.client.post(f"{same_course_cross_experiment_url}word/").status_code, 404)
        self.assertEqual(self.client.post(f"{cross_course_url}word/").status_code, 404)

    def test_word_report_is_post_only_and_csrf_protected(self):
        self.assertEqual(self.client.get(self.word_url).status_code, 405)

        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.teacher)
        self.assertEqual(csrf_client.post(self.word_url).status_code, 403)

    def test_word_report_contains_only_report_fields_and_a_privacy_safe_audit_log(self):
        response = self.client.post(self.word_url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            reverse(
                "courses:experiment_session_word",
                args=[self.course.pk, self.experiment.pk, self.session.pk],
            ),
            self.word_url,
        )
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn(
            f"filename=student-session-report-{self.session.pk}.docx",
            response["Content-Disposition"],
        )
        self.assertIn("filename*=UTF-8''", response["Content-Disposition"])
        self.assertIn(
            quote(f"学生作答详情-会话-{self.session.pk}.docx"),
            response["Content-Disposition"],
        )
        self.assertTrue(response.content.startswith(b"PK"))

        document = Document(BytesIO(response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("张小雨", text)
        self.assertIn("S-001", text)
        self.assertIn("2026 实验班（发布快照）", text)
        self.assertIn("v7", text)
        self.assertIn("进行中", text)
        self.assertIn("实际题数：3", text)
        self.assertIn("KT-C-001", text)
        self.assertIn("LEGACY-TF-002", text)
        self.assertIn("B：可以结合历史交互预测下一步表现", text)
        self.assertIn("C：可用于诊断知识状态", text)
        self.assertIn("A：只记录最终分数，不关心过程", text)
        self.assertIn("草稿，未正式提交", text)
        self.assertIn("回答正确", text)
        self.assertIn("未作答", text)
        self.assertNotIn("不得出现在报告中", text)
        self.assertNotIn("解析", text)
        self.assertNotIn("导航", text)
        self.assertNotIn("用时", text)
        self.assertEqual(document.sections[0].top_margin.inches, 1.0)
        self.assertEqual(document.sections[0].left_margin.inches, 1.0)
        self.assertEqual(document.sections[0].page_width.inches, 8.5)
        self.assertEqual(document.sections[0].page_height.inches, 11.0)
        self.assertEqual(len(document.tables), 0)
        self.assertIn("PAGE", document.sections[0].footer._element.xml)
        self.assertEqual(document.styles["Normal"].font.name, "微软雅黑")

        audit = AuditLog.objects.get(action="experiment.session_word_report")
        self.assertEqual(
            audit.details,
            {
                "experiment_id": self.experiment.pk,
                "version": self.version.version_number,
                "session_id": self.session.pk,
                "student_id": self.student.pk,
                "question_count": 3,
            },
        )
        self.assertNotIn("student_number", audit.details)
        self.assertNotIn("student_name", audit.details)
        self.assertNotIn("answers", audit.details)
        self.assertNotIn("question_text", audit.details)

    def test_word_report_keeps_the_formal_answer_when_grading_is_pending(self):
        attempt = self.submitted_assignment.attempt
        attempt.draft_answer = ["A"]
        attempt.final_answer = ["B", "C"]
        attempt.is_correct = None
        attempt.save(update_fields=["draft_answer", "final_answer", "is_correct"])

        response = self.client.post(self.word_url)

        self.assertEqual(response.status_code, 200)
        text = "\n".join(
            paragraph.text for paragraph in Document(BytesIO(response.content)).paragraphs
        )
        self.assertIn("B：可以结合历史交互预测下一步表现", text)
        self.assertIn("C：可用于诊断知识状态", text)
        self.assertIn("已提交，正误未判定", text)

    def test_word_builder_closes_its_memory_stream_after_response_bytes_are_read(self):
        TrackingBytesIO.instances.clear()

        with patch("apps.experiments.session_report.BytesIO", TrackingBytesIO):
            response = self.client.post(self.word_url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(TrackingBytesIO.instances), 1)
        self.assertTrue(TrackingBytesIO.instances[0].was_closed)

    def test_word_route_reuses_owner_scope_for_downloads(self):
        for user in (self.collaborator, self.outsider):
            with self.subTest(user=user.username):
                self.client.force_login(user)
                self.assertEqual(self.client.post(self.word_url).status_code, 404)

    def test_detail_and_word_support_every_session_state(self):
        expected_labels = dict(ExperimentSession.Status.choices)

        for status in ExperimentSession.Status.values:
            with self.subTest(status=status):
                self.session.status = status
                self.session.save(update_fields=["status"])
                detail = self.client.get(self.detail_url)
                word = self.client.post(self.word_url)

                self.assertEqual(detail.status_code, 200)
                self.assertContains(detail, expected_labels[status])
                self.assertEqual(word.status_code, 200)

    def test_word_report_handles_an_empty_assignment_set(self):
        another_student = User.objects.create_user(
            username="S-002",
            password="password123",
            student_number="S-002",
            display_name="李明",
            role=User.Role.STUDENT,
        )
        empty_session = ExperimentSession.objects.create(
            version=self.version,
            student=another_student,
            class_snapshot_name="空题会话班级",
            status=ExperimentSession.Status.COMPLETED,
        )
        detail_url = (
            f"/teacher/courses/{self.course.pk}/experiments/{self.experiment.pk}/"
            f"sessions/{empty_session.pk}/"
        )
        response = self.client.post(f"{detail_url}word/")

        self.assertEqual(response.status_code, 200)
        text = "\n".join(
            paragraph.text for paragraph in Document(BytesIO(response.content)).paragraphs
        )
        self.assertIn("实际题数：0", text)
        self.assertNotIn("第 1 题", text)
        audit = AuditLog.objects.filter(action="experiment.session_word_report").latest("created_at")
        self.assertEqual(audit.details["question_count"], 0)
