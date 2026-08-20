"""DOCX builder for the teacher's read-only student session report."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .presentation import SessionReportPresentation, build_session_report_presentation


BUSINESS_BLUE = RGBColor(0x2E, 0x74, 0xB5)
BUSINESS_DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK_BLUE = RGBColor(0x0B, 0x25, 0x45)
PRIMARY_FONT = "微软雅黑"
EAST_ASIA_FALLBACK_FONT = "宋体"


def _set_font(target, *, size=None, color=None, bold=None, east_asia=EAST_ASIA_FALLBACK_FONT):
    target.font.name = PRIMARY_FONT
    if size is not None:
        target.font.size = Pt(size)
    if color is not None:
        target.font.color.rgb = color
    if bold is not None:
        target.font.bold = bold
    r_pr = target._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), PRIMARY_FONT)
    r_fonts.set(qn("w:hAnsi"), PRIMARY_FONT)
    r_fonts.set(qn("w:cs"), PRIMARY_FONT)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _configure_styles(document):
    normal = document.styles["Normal"]
    _set_font(normal, size=11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = document.styles.add_style("Session Report Title", WD_STYLE_TYPE.PARAGRAPH)
    _set_font(title, size=18, color=INK_BLUE, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    heading_one = document.styles["Heading 1"]
    _set_font(heading_one, size=16, color=BUSINESS_BLUE, bold=True)
    heading_one.paragraph_format.space_before = Pt(16)
    heading_one.paragraph_format.space_after = Pt(8)

    heading_two = document.styles["Heading 2"]
    _set_font(heading_two, size=13, color=BUSINESS_BLUE, bold=True)
    heading_two.paragraph_format.space_before = Pt(12)
    heading_two.paragraph_format.space_after = Pt(6)

    heading_three = document.styles["Heading 3"]
    _set_font(heading_three, size=12, color=BUSINESS_DARK_BLUE, bold=True)
    heading_three.paragraph_format.space_before = Pt(8)
    heading_three.paragraph_format.space_after = Pt(4)


def _configure_section(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    return section


def _clear_paragraph(paragraph):
    paragraph._element.clear_content()


def _add_page_number(paragraph):
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_separate = OxmlElement("w:fldChar")
    field_separate.set(qn("w:fldCharType"), "separate")
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend((field_begin, instruction, field_separate, field_end))


def _configure_memo_furniture(section):
    header = section.header
    header.is_linked_to_previous = False
    header_paragraph = header.paragraphs[0]
    _clear_paragraph(header_paragraph)
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_run = header_paragraph.add_run("学生作答详情")
    _set_font(header_run, size=9, color=BUSINESS_DARK_BLUE, bold=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    _clear_paragraph(footer_paragraph)
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    prefix = footer_paragraph.add_run("第 ")
    _set_font(prefix, size=9)
    _add_page_number(footer_paragraph)
    suffix = footer_paragraph.add_run(" 页")
    _set_font(suffix, size=9)


def _add_text(paragraph, text, *, bold=False):
    lines = str(text).splitlines() or [""]
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        _set_font(run, size=11, bold=bold)


def _add_labelled_paragraph(document, label, value):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(f"{label}：")
    _set_font(label_run, size=11, color=INK_BLUE, bold=True)
    _add_text(paragraph, value)
    return paragraph


def _add_question(document, question):
    heading = document.add_paragraph(style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.widow_control = True
    _add_text(heading, f"第 {question.assigned_sequence_index} 题 · {question.stable_code}", bold=True)

    type_paragraph = _add_labelled_paragraph(document, "题型", question.question_type_label)
    type_paragraph.paragraph_format.keep_with_next = True

    stem_label = document.add_paragraph()
    stem_label.paragraph_format.keep_with_next = True
    label_run = stem_label.add_run("题干：")
    _set_font(label_run, size=11, color=INK_BLUE, bold=True)
    stem = document.add_paragraph()
    stem.paragraph_format.widow_control = True
    _add_text(stem, question.text)

    options_label = document.add_paragraph()
    options_label.paragraph_format.keep_with_next = True
    label_run = options_label.add_run("选项：")
    _set_font(label_run, size=11, color=INK_BLUE, bold=True)
    if question.options:
        for option_index, option in enumerate(question.options):
            option_paragraph = document.add_paragraph()
            option_paragraph.paragraph_format.left_indent = Inches(0.25)
            option_paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            option_paragraph.paragraph_format.space_after = Pt(3)
            option_paragraph.paragraph_format.widow_control = True
            option_paragraph.paragraph_format.keep_with_next = (
                option_index < len(question.options) - 1
            )
            _add_text(option_paragraph, option.display)
    else:
        none_paragraph = document.add_paragraph()
        none_paragraph.paragraph_format.left_indent = Inches(0.25)
        _add_text(none_paragraph, "无")

    student_answer = _add_labelled_paragraph(document, "学生答案", question.student_answer)
    correct_answer = _add_labelled_paragraph(document, "标准答案", question.correct_answer)
    status = _add_labelled_paragraph(document, "作答状态", question.answer_status)
    status.paragraph_format.space_after = Pt(10)


def build_session_report_document(report: SessionReportPresentation) -> bytes:
    """Build and close an in-memory DOCX byte stream for one session."""

    document = Document()
    _configure_styles(document)
    section = _configure_section(document)
    _configure_memo_furniture(section)

    title = document.add_paragraph(style="Session Report Title")
    title.paragraph_format.keep_with_next = True
    _add_text(title, "学生作答详情", bold=True)
    _add_labelled_paragraph(document, "学生姓名", report.student_name)
    _add_labelled_paragraph(document, "学号", report.student_number)
    _add_labelled_paragraph(document, "发布时班级", report.class_snapshot_name)
    _add_labelled_paragraph(document, "实验版本", f"v{report.version_number}")
    _add_labelled_paragraph(document, "会话状态", report.session_status)
    metadata_end = _add_labelled_paragraph(document, "实际题数", report.question_count)
    metadata_end.paragraph_format.space_after = Pt(10)

    for question in report.questions:
        _add_question(document, question)

    with BytesIO() as stream:
        document.save(stream)
        return stream.getvalue()


def build_session_report(session, assignments) -> bytes:
    """Stable helper for CLI/QA callers that starts from a session and assignments."""

    return build_session_report_document(
        build_session_report_presentation(session, assignments)
    )
