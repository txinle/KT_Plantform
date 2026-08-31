from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook
from io import BytesIO

from apps.accounts.models import Course, User

from .models import Concept, Question, QuestionBank
from .services.file_reader import read_uploaded_workbook
from .services.import_parsing import parse_kc_workbook
from .services.import_service import import_kc_workbook
from .services.question_kc_mapping import project_question_kc_mapping


class KCWorkbookV2Tests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(
            username="kc-workbook-teacher",
            role=User.Role.TEACHER,
        )
        self.course = Course.objects.create(
            name="工作簿课程",
            code="WORKBOOK",
            term="2026 秋",
            owner=self.teacher,
        )
        self.bank = QuestionBank.objects.create(course=self.course, name="工作簿题库")
        self.question = Question.objects.create(
            bank=self.bank,
            sequence=1,
            question_type=Question.Type.SHORT,
            text="工作簿题目",
            correct_answer="答案",
            status=Question.Status.DRAFT,
        )

    def upload(self):
        workbook = Workbook()
        nodes = workbook.active
        nodes.title = "KC节点"
        nodes.append(["kc_code", "kc_name", "description", "status"])
        nodes.append(["KC1", "数学", "根节点", "active"])
        nodes.append(["KC2", "代数", "二级节点", "active"])
        nodes.append(["KC3", "线性方程", "三级节点", "active"])
        mappings = workbook.create_sheet("题目映射")
        mappings.append(["stable_question_code", "primary_kc_code"])
        mappings.append([self.question.course_identity.stable_code, "KC1"])
        relations = workbook.create_sheet("题目KC关系")
        relations.append([
            "stable_question_code", "parent_kc_code", "child_kc_code", "child_order"
        ])
        relations.append([self.question.course_identity.stable_code, "KC1", "KC2", 1])
        relations.append([self.question.course_identity.stable_code, "KC2", "KC3", 1])
        stream = BytesIO()
        workbook.save(stream)
        return SimpleUploadedFile(
            "kc-v2.xlsx",
            stream.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_legacy_three_sheet_workbook_resolves_existing_nodes_and_replaces_one_tree(self):
        Concept.objects.create(course=self.course, code="KC1", name="数学")
        Concept.objects.create(course=self.course, code="KC2", name="代数")
        Concept.objects.create(course=self.course, code="KC3", name="线性方程")
        untouched = Concept.objects.create(course=self.course, code="KC4", name="未列出的 KC")
        workbook = parse_kc_workbook(read_uploaded_workbook(self.upload()))

        result = import_kc_workbook(self.bank, workbook)

        self.assertEqual(result, {"nodes_validated": 3, "questions_updated": 1})
        untouched.refresh_from_db()
        self.assertEqual(untouched.status, Concept.Status.ACTIVE)
        mapping = project_question_kc_mapping(self.question)
        self.assertEqual(mapping["root_kc_code"], "KC1")
        self.assertEqual(
            [node["kc_code"] for node in mapping["nodes"]],
            ["KC1", "KC2", "KC3"],
        )

    def test_legacy_three_sheet_workbook_cannot_create_course_kcs(self):
        workbook = parse_kc_workbook(read_uploaded_workbook(self.upload()))

        with self.assertRaisesMessage(ValueError, "课程 KC 总库中不存在 KC1"):
            import_kc_workbook(self.bank, workbook)

        self.assertFalse(Concept.objects.filter(course=self.course).exists())

    def test_rejects_missing_sheet_without_mutating_data(self):
        uploaded = self.upload()
        sheets = read_uploaded_workbook(uploaded)
        sheets.pop("题目KC关系")

        with self.assertRaisesMessage(ValueError, "题目KC关系"):
            parse_kc_workbook(sheets)

        self.assertFalse(Concept.objects.filter(course=self.course, code="KC1").exists())
