from __future__ import annotations

import csv
from io import BytesIO, StringIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font


QUESTION_IMPORT_TEMPLATE_HEADERS = (
    "序号",
    "题目",
    "题型",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "选项E",
    "选项F",
    "答案",
    "可接受答案",
    "解析",
    "难度",
    "一级知识点",
    "二级知识点",
    "三级知识点",
)

KC_TEMPLATE_INSTRUCTION = (
    "KC1、KC2、KC3、KC5、KC6 仅为示例，请替换为当前课程中已存在的 "
    "KC 编码；导入不会创建课程 KC。二级知识点用“；”分隔，三级知识点使用“二级KC>三级KC”语法。"
)


def _sample_row(*values: str) -> dict[str, str]:
    return dict(zip(QUESTION_IMPORT_TEMPLATE_HEADERS, values, strict=True))


QUESTION_IMPORT_TEMPLATE_ROWS = (
    _sample_row(
        "1",
        "2 + 2 等于多少？\n请从下列选项中选择。",
        "单选题",
        "3",
        "4",
        "5",
        "6",
        "",
        "",
        "B",
        "",
        "4 是唯一正确答案。\n这是基础运算示例。",
        "0.100",
        "KC1",
        "KC2；KC5",
        "KC2>KC3；KC5>KC6",
    ),
    _sample_row(
        "2",
        "下列哪些数是质数？\n可多选。",
        "多选题",
        "2\n（质数）",
        "4\n（合数）",
        "5\n（质数）",
        "6",
        "",
        "",
        "A,C",
        "",
        "2、5 是质数。\n4、6 是合数。",
        "0.350",
        "KC1",
        "KC2；KC5",
        "KC2>KC3；KC5>KC6",
    ),
    _sample_row(
        "3",
        "所有学习者都必须获得完全相同的题目顺序。",
        "判断题",
        "",
        "",
        "",
        "",
        "",
        "",
        "错误",
        "",
        "题目顺序可依据实验设置随机化。",
        "0.200",
        "KC1",
        "KC5",
        "KC5>KC6",
    ),
    _sample_row(
        "4",
        "Knowledge Tracing 的中文名称是？",
        "填空题",
        "",
        "",
        "",
        "",
        "",
        "",
        "知识追踪",
        "知识跟踪；KT",
        "常译为知识追踪。",
        "0.450",
        "KC1",
        "KC2",
        "KC2>KC3",
    ),
    _sample_row(
        "5",
        "请简述知识追踪的研究目的。\n请使用一到两句话作答。",
        "简答题",
        "",
        "",
        "",
        "",
        "",
        "",
        "估计学习者知识状态",
        "",
        "用于估计学习者随时间变化的知识状态。",
        "",
        "KC1",
        "KC2；KC5",
        "KC2>KC3；KC5>KC6",
    ),
    _sample_row(
        "6",
        "圆周率保留两位小数约为多少？",
        "数值题",
        "",
        "",
        "",
        "",
        "",
        "",
        "3.14",
        "",
        "按四舍五入保留两位小数。\n答案须按数值填写。",
        "0.250",
        "KC1",
        "KC2",
        "KC2>KC3",
    ),
)


def build_question_import_template_csv() -> bytes:
    output = StringIO(newline="")
    writer = csv.DictWriter(
        output,
        fieldnames=QUESTION_IMPORT_TEMPLATE_HEADERS,
        lineterminator="\n",
    )
    writer.writeheader()
    writer.writerows(QUESTION_IMPORT_TEMPLATE_ROWS)
    return output.getvalue().encode("utf-8-sig")


def build_question_import_template_excel() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "题目导入"
    sheet.append(QUESTION_IMPORT_TEMPLATE_HEADERS)
    for row in QUESTION_IMPORT_TEMPLATE_ROWS:
        sheet.append([row[header] for header in QUESTION_IMPORT_TEMPLATE_HEADERS])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in ("N", "O", "P"):
        sheet[f"{column}1"].comment = Comment(KC_TEMPLATE_INSTRUCTION, "KT Data Lab")
    sheet.freeze_panes = "A2"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _add_labeled_paragraph(document: Document, label: str, value: str):
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(f"【{label}】")
    label_run.bold = True
    paragraph.add_run(value)
    return paragraph


def build_question_import_template_word() -> bytes:
    document = Document()
    title_style = document.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = title_style.element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    question_heading_style = document.styles["Heading 2"]
    question_heading_style.font.color.rgb = RGBColor(0, 0, 0)
    question_heading_style.paragraph_format.keep_with_next = True
    document.add_paragraph("题目与 KC 联合导入模板", style="Title")
    document.add_paragraph(f"填写说明：{KC_TEMPLATE_INSTRUCTION}")

    for row in QUESTION_IMPORT_TEMPLATE_ROWS:
        question_paragraphs = [document.add_heading(f"第{row['序号']}题", level=2)]
        for header in QUESTION_IMPORT_TEMPLATE_HEADERS[1:]:
            value = row[header]
            if value:
                question_paragraphs.append(
                    _add_labeled_paragraph(
                        document,
                        "题干" if header == "题目" else header,
                        value,
                    )
                )
        for paragraph in question_paragraphs[:-1]:
            paragraph.paragraph_format.keep_with_next = True

    output = BytesIO()
    document.save(output)
    return output.getvalue()
