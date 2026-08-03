from __future__ import annotations
import csv
import io
from openpyxl import load_workbook


def read_uploaded_rows(uploaded_file):
    name = uploaded_file.name.lower()
    raw = uploaded_file.read()
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    if name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(item or "").strip() for item in rows[0]]
        return [dict(zip(headers, values)) for values in rows[1:]]
    raise ValueError("仅支持 .csv 或 .xlsx 文件")
