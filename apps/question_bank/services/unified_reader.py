"""Read unified question records without treating source numbers as identities."""
import csv
import io
from pathlib import Path

from openpyxl import load_workbook

from .word_import import read_docx_rows


HEADERS = ('序号', '题目', '题型', *[f'选项{k}' for k in 'ABCDEF'],
           '答案', '可接受答案', '解析', '难度', '一级知识点', '二级知识点', '三级知识点')


def read_unified_rows(uploaded):
    suffix = Path(uploaded.name).suffix.lower()
    if suffix == '.docx':
        return read_docx_rows(uploaded)
    if suffix == '.csv':
        raw = uploaded.read()
        try:
            content = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                content = raw.decode('gb18030')
            except UnicodeDecodeError as error:
                raise ValueError('CSV 必须使用 UTF-8 或 GB18030 编码。') from error
        try:
            data = [(number, values, []) for number, values in enumerate(csv.reader(io.StringIO(content), strict=True), 1)]
        except csv.Error as error:
            raise ValueError(f'CSV 格式无效：{error}') from error
        location = '第 {number} 行'
    elif suffix == '.xlsx':
        try:
            workbook = load_workbook(uploaded, data_only=False)
            sheets = [sheet for sheet in workbook.worksheets
                      if any(cell.value is not None for row in sheet for cell in row)]
            if len(sheets) != 1:
                raise ValueError('Excel 必须且只能包含一个有数据的工作表。')
            sheet = sheets[0]
            if sheet._images or sheet._charts:
                raise ValueError('Excel 包含图片或图表，请改用纯文本题目。')
            data = [(number, [cell.value for cell in cells],
                     ['包含公式单元格，请改为明确的文本或数值。'] if any(cell.data_type == 'f' for cell in cells) else [])
                    for number, cells in enumerate(sheet.iter_rows(), 1)]
            location = f'{sheet.title} · 第 {{number}} 行'
            workbook.close()
        except ValueError:
            raise
        except Exception as error:
            raise ValueError('无法读取 Excel 文件，请使用有效的 .xlsx 文件。') from error
    else:
        raise ValueError('请选择 .xlsx、.csv 或 .docx 文件。')
    if not data:
        raise ValueError('文件为空。')
    headers = [str(value or '').strip() for value in data[0][1]]
    headers = ['题目' if name == '题干' else name for name in headers]
    # Empty trailing Excel columns are formatting, not additional fields.
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) != len(set(headers)) or set(headers) != set(HEADERS):
        raise ValueError(f'表头必须完整且不重复：{"、".join(HEADERS)}（题目也可写为题干，列可调整顺序）。')
    rows = []
    for number, values, errors in data[1:]:
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if any(value is not None and str(value).strip() for value in values[len(headers):]):
            errors = [*errors, '存在没有对应表头的数据列。']
        values = [*values, *([''] * max(0, len(headers) - len(values)))]
        rows.append({**dict(zip(headers, values)), '_location': location.format(number=number), '_errors': errors})
    if not rows:
        raise ValueError('文件中没有可导入的题目。')
    return rows
