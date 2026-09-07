import importlib
import io
from zipfile import ZIP_DEFLATED, ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from docx import Document
from docx.oxml import OxmlElement


CANONICAL_KEYS = (
    "序号",
    "题目",
    "题型",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "选项E",
    "选项F",
    "答案",
    "可接受答案",
    "解析",
    "难度",
    "一级知识点",
    "二级知识点",
    "三级知识点",
)


def _uploaded_docx(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return SimpleUploadedFile(
        "questions.docx",
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )


class WordQuestionImportTests(SimpleTestCase):
    def _reader(self):
        try:
            module = importlib.import_module("apps.question_bank.services.word_import")
        except ModuleNotFoundError:
            return None
        return getattr(module, "read_docx_rows", None)

    def test_reads_numbered_paragraph_question_with_split_label_and_multiline_stem(self):
        document = Document()
        document.add_paragraph("第1题")
        stem = document.add_paragraph()
        first = stem.add_run("【题")
        first.bold = True
        stem.add_run("干】：计算  x + y")
        document.add_paragraph("第二行")
        document.add_paragraph(
            "【题型】：单选题 【选项A】：甲 【选项B】：乙 【答案】：A"
        )

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(len(rows), 1)
        self.assertEqual(
            {key: rows[0][key] for key in CANONICAL_KEYS},
            {
                "序号": "1",
                "题目": "计算  x + y\n第二行",
                "题型": "单选题",
                "选项A": "甲",
                "选项B": "乙",
                "选项C": "",
                "选项D": "",
                "选项E": "",
                "选项F": "",
                "答案": "A",
                "可接受答案": "",
                "解析": "",
                "难度": "",
                "一级知识点": "",
                "二级知识点": "",
                "三级知识点": "",
            },
        )
        self.assertIn("第1段", rows[0]["_location"])
        self.assertEqual(rows[0]["_errors"], [])

    def test_reads_reordered_field_value_table_and_maps_knowledge_point_synonyms(self):
        document = Document()
        table = document.add_table(rows=10, cols=2)
        fields = (
            ("答案", "B"),
            ("序号", "7"),
            ("题型", "单选题"),
            ("选项B", "蓝"),
            ("次要知识点", "二级 KC"),
            ("【题干】", "天空是什么颜色？"),
            ("主要", "一级 KC"),
            ("三级知识点", "三级 KC"),
            ("选项A", "红"),
            ("解析", "晴天通常为蓝色。"),
        )
        for row, (label, value) in zip(table.rows, fields):
            row.cells[0].text = label
            row.cells[1].text = value

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["序号"], "7")
        self.assertEqual(rows[0]["题目"], "天空是什么颜色？")
        self.assertEqual(rows[0]["题型"], "单选题")
        self.assertEqual(rows[0]["选项A"], "红")
        self.assertEqual(rows[0]["选项B"], "蓝")
        self.assertEqual(rows[0]["选项E"], "")
        self.assertEqual(rows[0]["选项F"], "")
        self.assertEqual(rows[0]["答案"], "B")
        self.assertEqual(rows[0]["解析"], "晴天通常为蓝色。")
        self.assertEqual(rows[0]["一级知识点"], "一级 KC")
        self.assertEqual(rows[0]["二级知识点"], "二级 KC")
        self.assertEqual(rows[0]["三级知识点"], "三级 KC")
        self.assertEqual(rows[0]["_errors"], [])
        self.assertIn("表格1", rows[0]["_location"])

    def test_keeps_paragraph_and_header_table_questions_in_document_order(self):
        document = Document()
        document.add_paragraph("1.【题干】：段落题")
        document.add_paragraph("【题型】：单选题 【选项A】：甲 【选项B】：乙 【答案】：A")

        table = document.add_table(rows=2, cols=10)
        headers = (
            "序号",
            "题目",
            "题型",
            "选项A",
            "选项B",
            "选项E",
            "选项F",
            "答案",
            "知识点",
            "解析",
        )
        values = (
            "2",
            "表格题",
            "单选题",
            "春",
            "秋",
            "",
            "",
            "B",
            "季节",
            "秋天在春天之后。",
        )
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, values):
            cell.text = value

        document.add_paragraph("第3题")
        document.add_paragraph("【题目】：尾题 【答案】：文本答案")

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual([row["序号"] for row in rows], ["1", "2", "3"])
        self.assertEqual([row["题目"] for row in rows], ["段落题", "表格题", "尾题"])
        self.assertEqual(rows[1]["选项E"], "")
        self.assertEqual(rows[1]["选项F"], "")
        self.assertEqual(rows[1]["一级知识点"], "季节")
        self.assertEqual(rows[1]["解析"], "秋天在春天之后。")
        self.assertEqual(rows[1]["_errors"], [])
        self.assertIn("表格1", rows[1]["_location"])

    def test_returns_duplicate_unknown_and_ambiguous_boundary_errors_per_question(self):
        document = Document()
        document.add_paragraph("第1题")
        document.add_paragraph("【题干】：首次题干")
        document.add_paragraph("[题目]：重复题干")
        document.add_paragraph("［未知字段］：不应静默丢弃")
        document.add_paragraph("2. 没有字段标签的题目")
        document.add_paragraph("3. 【未知字段】：边界后的未知内容")

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual([row["序号"] for row in rows], ["1", "2", "3"])
        self.assertIn("重复字段：题目", rows[0]["_errors"])
        self.assertIn("无法识别的字段：未知字段", rows[0]["_errors"])
        self.assertIn("不应静默丢弃", rows[0]["题目"])
        self.assertEqual(rows[1]["题目"], "没有字段标签的题目")
        self.assertEqual(rows[1]["_errors"], ["题目边界不明确"])
        self.assertIn("第5段", rows[1]["_location"])
        self.assertIn("边界后的未知内容", rows[2]["题目"])
        self.assertEqual(
            rows[2]["_errors"],
            ["题目边界不明确", "无法识别的字段：未知字段"],
        )
        for row in rows:
            for error in row["_errors"]:
                self.assertNotRegex(error, r"第\d+段|表格\d")

    def test_reports_unsupported_ooxml_objects_without_external_media_fixtures(self):
        document = Document()
        document.add_paragraph("第1题")
        paragraph = document.add_paragraph("【题干】：当 x ≤ y 时，比较结果。")
        for tag in ("w:drawing", "m:oMath", "w:txbxContent"):
            run = paragraph.add_run()
            run._r.append(OxmlElement(tag))

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["题目"], "当 x ≤ y 时，比较结果。")
        self.assertEqual(
            rows[0]["_errors"],
            [
                "包含不支持的图片对象",
                "包含不支持的公式对象",
                "包含不支持的文本框对象",
            ],
        )
        self.assertIn("第2段", rows[0]["_location"])

    def test_raises_value_error_for_unrecoverable_or_non_docx_documents(self):
        reader = self._reader()
        self.assertIsNotNone(reader)

        with self.assertRaisesMessage(ValueError, "无法读取 DOCX 文档"):
            reader(SimpleUploadedFile("broken.docx", b"not a zip archive"))
        with self.assertRaisesMessage(ValueError, "仅支持 .docx 文件"):
            reader(SimpleUploadedFile("questions.txt", b"not a docx"))

    def test_reports_unsupported_object_in_a_table_cell_with_its_location(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "序号"
        table.rows[0].cells[1].text = "4"
        table.rows[1].cells[0].text = "题干"
        paragraph = table.rows[1].cells[1].paragraphs[0]
        paragraph.add_run("表格中的题干")
        image_run = paragraph.add_run()
        image_run._r.append(OxmlElement("w:drawing"))

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["题目"], "表格中的题干")
        self.assertEqual(rows[0]["_errors"], ["包含不支持的图片对象"])
        self.assertIn("表格1第2行第2列", rows[0]["_location"])

    def test_reports_unsupported_object_in_a_textless_paragraph(self):
        document = Document()
        document.add_paragraph("第1题")
        document.add_paragraph("【题干】：文字题干")
        object_only = document.add_paragraph()
        run = object_only.add_run()
        run._r.append(OxmlElement("w:drawing"))

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["题目"], "文字题干")
        self.assertEqual(rows[0]["_errors"], ["包含不支持的图片对象"])
        self.assertIn("第3段", rows[0]["_location"])

    def test_reports_unsupported_object_in_a_header_table_cell(self):
        document = Document()
        table = document.add_table(rows=2, cols=3)
        for cell, value in zip(table.rows[0].cells, ("序号", "题目", "答案")):
            cell.text = value
        image_run = table.rows[0].cells[1].paragraphs[0].add_run()
        image_run._r.append(OxmlElement("w:drawing"))
        for cell, value in zip(table.rows[1].cells, ("1", "表格题", "文本答案")):
            cell.text = value

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["题目"], "表格题")
        self.assertEqual(rows[0]["_errors"], ["包含不支持的图片对象"])
        self.assertIn("表格1第1行第2列", rows[0]["_location"])

    def test_returns_unknown_field_value_table_content_with_an_error(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "题干"
        table.rows[0].cells[1].text = "已保留的题干"
        table.rows[1].cells[0].text = "来源"
        table.rows[1].cells[1].text = "讲义第 3 页"

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["序号"], "1")
        self.assertIn("已保留的题干", rows[0]["题目"])
        self.assertIn("来源：讲义第 3 页", rows[0]["题目"])
        self.assertEqual(rows[0]["_errors"], ["无法识别的字段：来源"])

    def test_keeps_unlabeled_content_after_a_heading_out_of_the_sequence(self):
        document = Document()
        document.add_paragraph("第1题：未标注的题干内容")

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["序号"], "1")
        self.assertEqual(rows[0]["题目"], "未标注的题干内容")
        self.assertEqual(rows[0]["_errors"], [])

    def test_preserves_math_interval_brackets_as_question_text(self):
        document = Document()
        document.add_paragraph("第1题")
        document.add_paragraph("【题干】：设 x ∈ [0, 1]，且 x ≤ 1。")

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["题目"], "设 x ∈ [0, 1]，且 x ≤ 1。")
        self.assertEqual(rows[0]["_errors"], [])

    def test_preserves_letter_colon_in_a_stem_while_reading_bare_option_lines(self):
        document = Document()
        document.add_paragraph("第1题")
        document.add_paragraph("【题干】：若 A：B，则 A ≠ B。")
        document.add_paragraph("A. 选项甲")
        document.add_paragraph("B. 选项乙")
        document.add_paragraph("【答案】：A")

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["题目"], "若 A：B，则 A ≠ B。")
        self.assertEqual(rows[0]["选项A"], "选项甲")
        self.assertEqual(rows[0]["选项B"], "选项乙")
        self.assertEqual(rows[0]["_errors"], [])

    def test_ignores_the_current_word_template_preamble_before_its_questions(self):
        from .services.import_templates import build_question_import_template_word

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(
            SimpleUploadedFile(
                "template.docx",
                build_question_import_template_word(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
            )
        )

        self.assertEqual([row["序号"] for row in rows], ["1", "2", "3", "4", "5", "6"])
        self.assertEqual(rows[0]["题目"], "2 + 2 等于多少？\n请从下列选项中选择。")
        self.assertEqual(rows[0]["选项E"], "")
        self.assertEqual(rows[0]["选项F"], "")
        self.assertFalse(any(row["_errors"] for row in rows))

    def test_attaches_unsupported_template_preamble_objects_to_the_first_question(self):
        document = Document()
        title = document.add_paragraph("题目与 KC 联合导入模板")
        title_run = title.add_run()
        title_run._r.append(OxmlElement("w:drawing"))
        document.add_paragraph("填写说明：这是模板说明。")
        document.add_paragraph("第1题")
        document.add_paragraph("【题干】：实际题目")

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["题目"], "实际题目")
        self.assertEqual(rows[0]["_errors"], ["包含不支持的图片对象"])
        self.assertIn("第1段", rows[0]["_location"])

    def test_reads_structured_paragraphs_inside_a_one_cell_table_in_order(self):
        document = Document()
        document.add_paragraph("第1题")
        cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
        cell.paragraphs[0].text = "【题干】：表格题干"
        cell.add_paragraph("【答案】：表格答案")

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["题目"], "表格题干")
        self.assertEqual(rows[0]["答案"], "表格答案")
        self.assertEqual(rows[0]["_errors"], [])
        self.assertIn("表格1第1行第1列", rows[0]["_location"])

    def test_reports_unrecognized_table_with_unsupported_object_instead_of_ignoring_it(self):
        document = Document()
        document.add_paragraph("第1题")
        cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
        run = cell.paragraphs[0].add_run()
        run._r.append(OxmlElement("m:oMath"))

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertIn("包含不支持的公式对象", rows[0]["_errors"])
        self.assertIn("无法识别的表格结构", rows[0]["_errors"])
        self.assertIn("表格1第1行第1列", rows[0]["_location"])

    def test_marks_a_one_cell_field_label_without_value_as_an_unrecognized_table(self):
        document = Document()
        document.add_paragraph("第1题")
        document.add_table(rows=1, cols=1).rows[0].cells[0].text = "题干"

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertIn("无法识别的表格结构", rows[0]["_errors"])
        self.assertIn("表格1第1行第1列", rows[0]["_location"])

    def test_treats_an_empty_knowledge_point_marker_as_a_section_heading(self):
        document = Document()
        document.add_paragraph("第1题")
        document.add_paragraph("【题干】：知识点示例题")
        document.add_paragraph("【知识点】")
        document.add_paragraph("一级知识点：KC1")
        document.add_paragraph("二级知识点：KC2；KC5")
        document.add_paragraph("三级知识点：KC2>KC3；KC5>KC6")

        reader = self._reader()
        self.assertIsNotNone(reader)
        rows = reader(_uploaded_docx(document))

        self.assertEqual(rows[0]["一级知识点"], "KC1")
        self.assertEqual(rows[0]["二级知识点"], "KC2；KC5")
        self.assertEqual(rows[0]["三级知识点"], "KC2>KC3；KC5>KC6")
        self.assertEqual(rows[0]["_errors"], [])

    def test_raises_value_error_for_malformed_document_xml(self):
        document = Document()
        document.add_paragraph("第1题")
        document.add_paragraph("【题干】：有效题干")
        source = io.BytesIO(_uploaded_docx(document).read())
        malformed = io.BytesIO()
        with ZipFile(source) as source_archive, ZipFile(
            malformed, "w", ZIP_DEFLATED
        ) as target_archive:
            for info in source_archive.infolist():
                content = source_archive.read(info.filename)
                if info.filename == "word/document.xml":
                    content = b"<w:document"
                target_archive.writestr(info, content)

        reader = self._reader()
        self.assertIsNotNone(reader)
        with self.assertRaisesMessage(ValueError, "无法读取 DOCX 文档"):
            reader(SimpleUploadedFile("malformed.docx", malformed.getvalue()))
