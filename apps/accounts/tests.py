from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.audit.models import AuditLog
from .models import ClassGroup, User
from .student_import import parse_student_workbook

class HomeRoutingTests(TestCase):
    def test_teacher_is_routed_to_course_workspace(self):
        user = User.objects.create_user(username="t", password="x", role=User.Role.TEACHER)
        self.client.force_login(user)
        self.assertRedirects(self.client.get(reverse("home")), reverse("courses:list"))

    def test_student_is_routed_to_student_dashboard(self):
        user = User.objects.create_user(username="S1", password="x", role=User.Role.STUDENT, student_number="S1")
        self.client.force_login(user)
        self.assertRedirects(self.client.get(reverse("home")), reverse("assessment:dashboard"))


class LoginPresentationTests(TestCase):
    @override_settings(DEBUG=True)
    def test_demo_credentials_are_visible_in_debug(self):
        response = self.client.get(reverse("login"))
        self.assertContains(response, "teacher / teacher123")

    @override_settings(DEBUG=False)
    def test_demo_credentials_are_hidden_outside_debug(self):
        response = self.client.get(reverse("login"))
        self.assertNotContains(response, "teacher / teacher123")

    def test_login_page_explains_identified_and_legacy_export_boundary(self):
        response = self.client.get(reverse("login"))

        self.assertContains(response, "实名主导出仅限授权场景")
        self.assertContains(response, "匿名导出保留兼容")
        self.assertNotContains(response, "研究导出默认使用匿名编号")


class HeaderIdentityTests(TestCase):
    def test_teacher_header_shows_name_and_login_account(self):
        teacher = User.objects.create_user(
            username="teacher-account",
            password="x",
            role=User.Role.TEACHER,
            display_name="陈老师",
        )
        self.client.force_login(teacher)

        response = self.client.get(reverse("question_bank:dashboard"))

        self.assertContains(response, '<span class="account-name">陈老师</span>', html=True)
        self.assertContains(response, '<span class="account-username mono">teacher-account</span>', html=True)
        self.assertNotContains(response, "陈老师 / 教师")

    def test_student_header_uses_the_same_name_and_account_format(self):
        student = User.objects.create_user(
            username="S101",
            password="x",
            role=User.Role.STUDENT,
            student_number="S101",
            display_name="林同学",
        )
        self.client.force_login(student)

        response = self.client.get(reverse("assessment:dashboard"))

        self.assertContains(response, '<span class="account-name">林同学</span>', html=True)
        self.assertContains(response, '<span class="account-username mono">S101</span>', html=True)
        self.assertNotContains(response, "林同学 / 学生")

    def test_header_does_not_repeat_account_when_name_falls_back_to_username(self):
        teacher = User.objects.create_user(
            username="fallback-account",
            password="x",
            role=User.Role.TEACHER,
        )
        self.client.force_login(teacher)

        response = self.client.get(reverse("question_bank:dashboard"))
        account_label = response.content.decode("utf-8").split("account-label", 1)[1].split("</span>", 1)[0]

        self.assertEqual(account_label.count("fallback-account"), 1)


class StudentListFilterTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(
            username="teacher",
            password="x",
            role=User.Role.TEACHER,
        )
        self.other_teacher = User.objects.create_user(
            username="other",
            password="x",
            role=User.Role.TEACHER,
        )
        self.class_a = ClassGroup.objects.create(
            owner=self.teacher,
            academic_year="2026",
            name="认知追踪实验班",
        )
        self.class_b = ClassGroup.objects.create(
            owner=self.teacher,
            academic_year="2026",
            name="数据分析班",
        )
        self.other_class = ClassGroup.objects.create(
            owner=self.other_teacher,
            academic_year="2026",
            name="其他教师班级",
        )
        self.student_a = User.objects.create_user(
            username="00123",
            student_number="00123",
            display_name="张晓明",
            class_group=self.class_a,
            role=User.Role.STUDENT,
        )
        self.student_b = User.objects.create_user(
            username="S002",
            student_number="S002",
            display_name="李华",
            class_group=self.class_b,
            role=User.Role.STUDENT,
        )
        self.other_student = User.objects.create_user(
            username="O001",
            student_number="O001",
            display_name="越权学生",
            class_group=self.other_class,
            role=User.Role.STUDENT,
        )
        self.client.force_login(self.teacher)

    def test_searches_by_student_number_or_chinese_name(self):
        url = reverse("accounts:student_list")

        number_response = self.client.get(url, {"q": "012"})
        name_response = self.client.get(url, {"q": "晓明"})

        self.assertContains(number_response, "00123")
        self.assertNotContains(number_response, "S002")
        self.assertContains(name_response, "张晓明")
        self.assertNotContains(name_response, "李华")

    def test_combines_search_and_owned_class_filter(self):
        response = self.client.get(
            reverse("accounts:student_list"),
            {"q": "张", "class_id": str(self.class_a.pk)},
        )

        self.assertContains(response, "张晓明")
        self.assertNotContains(response, "李华")
        self.assertContains(response, 'value="张"')
        self.assertContains(response, f'value="{self.class_a.pk}" selected')

    def test_foreign_class_filter_never_exposes_foreign_students(self):
        response = self.client.get(
            reverse("accounts:student_list"),
            {"class_id": str(self.other_class.pk)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "越权学生")
        self.assertNotContains(response, "其他教师班级")

    def test_malformed_class_id_falls_back_without_server_error(self):
        response = self.client.get(
            reverse("accounts:student_list"),
            {"class_id": "not-a-number"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "张晓明")
        self.assertContains(response, "李华")
        self.assertNotContains(response, 'value="not-a-number"')

    def test_no_results_preserves_filter_controls_and_shows_filtered_empty_state(self):
        response = self.client.get(reverse("accounts:student_list"), {"q": "不存在"})

        self.assertContains(response, 'value="不存在"')
        self.assertContains(response, "没有符合筛选条件的学生")
        self.assertContains(response, "清除筛选")

    def test_student_page_links_to_batch_import(self):
        response = self.client.get(reverse("accounts:student_list"))

        self.assertContains(response, reverse("accounts:student_import"))
        self.assertContains(response, "批量导入学生")


class StudentBatchImportCoreTests(TestCase):
    template_url = "/teacher/accounts/students/import-template.xlsx"
    import_url = "/teacher/accounts/students/import/"

    def setUp(self):
        self.teacher = User.objects.create_user(
            username="teacher",
            password="x",
            role=User.Role.TEACHER,
        )
        self.class_group = ClassGroup.objects.create(
            owner=self.teacher,
            academic_year="2026",
            name="认知追踪实验班",
        )
        self.client.force_login(self.teacher)

    def _upload(self, headers, rows, name="学生名单.xlsx"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        stream = BytesIO()
        workbook.save(stream)
        return SimpleUploadedFile(
            name,
            stream.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_downloads_unicode_template_with_instructions_and_owned_classes(self):
        response = self.client.get(self.template_url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("filename*=UTF-8''", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        self.assertEqual(workbook.sheetnames, ["学生导入", "填写说明"])
        self.assertEqual(
            [cell.value for cell in next(workbook["学生导入"].iter_rows(max_row=1))],
            ["学号", "姓名", "班级"],
        )
        instructions = "\n".join(
            str(cell.value or "") for row in workbook["填写说明"].iter_rows() for cell in row
        )
        self.assertIn("2026 认知追踪实验班", instructions)
        self.assertIn("S + 学号", instructions)
        self.assertIn("填写示例", instructions)
        self.assertIn("20260001", instructions)
        self.assertIn("示例学生", instructions)
        self.assertEqual(workbook["学生导入"].column_dimensions["A"].number_format, "@")
        self.assertEqual(workbook["学生导入"].column_dimensions["B"].number_format, "@")
        self.assertEqual(workbook["学生导入"].column_dimensions["C"].number_format, "@")

    def test_previews_chinese_rows_with_columns_in_any_order(self):
        upload = self._upload(
            ["班级", "姓名", "学号"],
            [["2026 认知追踪实验班", "林晓雨", "00123"]],
        )

        response = self.client.post(self.import_url, {"file": upload})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "校验通过")
        self.assertContains(response, "00123")
        self.assertContains(response, "林晓雨")
        self.assertContains(response, "2026 认知追踪实验班")
        self.assertFalse(User.objects.filter(student_number="00123").exists())

    def test_import_page_has_a_real_accessible_region_label(self):
        response = self.client.get(self.import_url)

        self.assertContains(response, 'id="student-import-title"')
        self.assertContains(response, 'aria-labelledby="student-import-title"')

    def test_confirm_creates_students_atomically_with_documented_password(self):
        upload = self._upload(
            ["学号", "姓名", "班级"],
            [
                ["00123", "林晓雨", "2026 认知追踪实验班"],
                ["A-02", "王强", "2026 认知追踪实验班"],
            ],
        )
        preview = self.client.post(self.import_url, {"file": upload})
        self.assertEqual(preview.status_code, 200)

        response = self.client.post(self.import_url, {"action": "confirm"})

        self.assertRedirects(response, reverse("accounts:student_list"))
        first = User.objects.get(student_number="00123")
        second = User.objects.get(student_number="A-02")
        self.assertEqual(first.display_name, "林晓雨")
        self.assertEqual(first.class_group, self.class_group)
        self.assertTrue(first.check_password("S00123"))
        self.assertTrue(second.check_password("SA-02"))
        log = AuditLog.objects.get(action="student.batch_import")
        self.assertEqual(log.details, {"count": 2})
        self.assertNotIn("password", str(log.details).lower())


class StudentBatchImportValidationTests(StudentBatchImportCoreTests):
    def _assert_upload_error(self, headers, rows, expected):
        before = User.objects.filter(role=User.Role.STUDENT).count()
        response = self.client.post(self.import_url, {"file": self._upload(headers, rows)})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, expected)
        self.assertEqual(User.objects.filter(role=User.Role.STUDENT).count(), before)

    def test_rejects_duplicate_existing_unknown_class_and_blank_fields(self):
        User.objects.create_user(
            username="EXISTING",
            student_number="EXISTING",
            display_name="已有学生",
            class_group=self.class_group,
            role=User.Role.STUDENT,
        )
        cases = [
            (
                ["学号", "姓名", "班级"],
                [["DUP", "甲", str(self.class_group)], ["DUP", "乙", str(self.class_group)]],
                "学号重复：DUP",
            ),
            (["学号", "姓名", "班级"], [["EXISTING", "甲", str(self.class_group)]], "学号已存在"),
            (["学号", "姓名", "班级"], [["NEW", "甲", "不存在的班级"]], "班级无法匹配"),
            (["学号", "姓名", "班级"], [["NEW", "", str(self.class_group)]], "存在空字段"),
        ]
        for headers, rows, expected in cases:
            with self.subTest(expected=expected):
                self._assert_upload_error(headers, rows, expected)

    def test_rejects_missing_duplicate_unknown_headers_and_non_xlsx_files(self):
        cases = [
            (["学号", "姓名"], [["S1", "甲"]]),
            (["学号", "姓名", "班级", "学号"], [["S1", "甲", str(self.class_group), "S1"]]),
            (["学号", "姓名", "班级", "备注"], [["S1", "甲", str(self.class_group), "x"]]),
        ]
        for headers, rows in cases:
            with self.subTest(headers=headers):
                self._assert_upload_error(headers, rows, "表头必须且只能包含")

        response = self.client.post(
            self.import_url,
            {"file": SimpleUploadedFile("名单.csv", b"student,name,class")},
        )
        self.assertContains(response, "仅支持 .xlsx 文件")

    def test_rejects_values_longer_than_model_fields(self):
        self._assert_upload_error(
            ["学号", "姓名", "班级"],
            [["S" * 65, "甲", str(self.class_group)]],
            "学号不能超过 64 个字符",
        )
        self._assert_upload_error(
            ["学号", "姓名", "班级"],
            [["NEW", "甲" * 81, str(self.class_group)]],
            "姓名不能超过 80 个字符",
        )

    def test_confirm_revalidates_conflicts_without_partial_writes(self):
        upload = self._upload(
            ["学号", "姓名", "班级"],
            [
                ["RACE", "并发学生", str(self.class_group)],
                ["SAFE", "不应写入", str(self.class_group)],
            ],
        )
        self.client.post(self.import_url, {"file": upload})
        User.objects.create_user(
            username="RACE",
            student_number="RACE",
            display_name="先创建",
            class_group=self.class_group,
            role=User.Role.STUDENT,
        )

        response = self.client.post(self.import_url, {"action": "confirm"})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "部分学号已存在")
        self.assertFalse(User.objects.filter(student_number="SAFE").exists())
        self.assertEqual(User.objects.filter(student_number="RACE").count(), 1)

    def test_student_role_cannot_download_template_or_open_import(self):
        student = User.objects.create_user(
            username="STUDENT",
            student_number="STUDENT",
            role=User.Role.STUDENT,
        )
        self.client.force_login(student)

        self.assertEqual(self.client.get(self.template_url).status_code, 403)
        self.assertEqual(self.client.get(self.import_url).status_code, 403)

    def test_failed_new_upload_invalidates_previous_preview(self):
        valid = self._upload(
            ["学号", "姓名", "班级"],
            [["STALE", "旧预览学生", str(self.class_group)]],
        )
        self.client.post(self.import_url, {"file": valid})
        invalid = SimpleUploadedFile("错误名单.csv", b"student,name,class")
        failed_upload = self.client.post(self.import_url, {"file": invalid})
        self.assertContains(failed_upload, "仅支持 .xlsx 文件")

        response = self.client.post(self.import_url, {"action": "confirm"})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "导入预览已失效")
        self.assertFalse(User.objects.filter(student_number="STALE").exists())

    def test_existing_account_check_uses_constant_query_count(self):
        upload = self._upload(
            ["学号", "姓名", "班级"],
            [
                ["B001", "甲", str(self.class_group)],
                ["B002", "乙", str(self.class_group)],
                ["B003", "丙", str(self.class_group)],
            ],
        )

        with self.assertNumQueries(2):
            parsed = parse_student_workbook(
                upload,
                ClassGroup.objects.filter(pk=self.class_group.pk),
            )

        self.assertEqual(len(parsed), 3)
