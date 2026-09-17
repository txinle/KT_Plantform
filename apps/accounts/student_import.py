import csv
from dataclasses import dataclass
from io import BytesIO, StringIO

from django.db.models import Q
from openpyxl import Workbook, load_workbook

from .models import User


EXPECTED_HEADERS = ("学号", "姓名", "班级")


class StudentImportError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedStudent:
    student_number: str
    display_name: str
    class_id: int
    class_name: str
    source_class_id: int | None = None


def build_student_import_workbook(classes, scope_label=None, *, allow_csv=False):
    classes = list(classes)
    workbook = Workbook()
    import_sheet = workbook.active
    import_sheet.title = "学生导入"
    import_sheet.append(["学号", "姓名", "班级"])
    import_sheet.freeze_panes = "A2"
    for column in ("A", "B", "C"):
        import_sheet.column_dimensions[column].number_format = "@"

    instructions = workbook.create_sheet("填写说明")
    instructions.append(["填写规则"])
    if scope_label:
        instructions.append([f"适用范围：{scope_label}"])
    formats = ".xlsx 或 .csv" if allow_csv else ".xlsx"
    instructions.append([f"支持 {formats}；学号、姓名、班级均为必填字段。"])
    if len(classes) == 1:
        instructions.append(
            [
                "本页导入可填写当前班级名称或完整名称；"
                "校验后将统一显示为完整名称。"
            ]
        )
    else:
        instructions.append(["班级必须与下列可用班级完整名称精确一致。"])
    instructions.append(["新账号初始密码规则：S + 学号；已有账号密码保持不变。"])
    instructions.append([])
    instructions.append(["填写示例（仅供参考，不会从本页导入）"])
    instructions.append(["学号", "姓名", "班级"])
    instructions.append(["20260001", "示例学生", str(classes[0]) if classes else "2026 示例班级"])
    instructions.append([])
    instructions.append(["可用班级完整名称"])
    for class_group in classes:
        instructions.append([str(class_group)])

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def parse_student_workbook(uploaded_file, classes, *, allow_existing=False, allow_csv=False):
    name = uploaded_file.name.lower()
    if name.endswith(".csv") and allow_csv:
        content = uploaded_file.read()
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                decoded = content.decode("gb18030")
            except UnicodeDecodeError as exc:
                raise StudentImportError("无法读取 CSV 文件，请使用 UTF-8 或 GB18030 编码") from exc
        rows = csv.reader(StringIO(decoded, newline=""), strict=True)
    elif name.endswith(".xlsx"):
        try:
            workbook = load_workbook(BytesIO(uploaded_file.read()), read_only=True, data_only=True)
        except Exception as exc:
            raise StudentImportError("无法读取 Excel 文件，请使用下载的参考模板") from exc
        rows = workbook.active.iter_rows(values_only=True)
    else:
        raise StudentImportError("仅支持 .xlsx 或 .csv 文件" if allow_csv else "仅支持 .xlsx 文件")

    try:
        return _parse_student_rows(rows, classes, allow_existing=allow_existing)
    except csv.Error as exc:
        raise StudentImportError("CSV 文件格式无效，请检查引号和分隔符") from exc


def _parse_student_rows(rows, classes, *, allow_existing):
    header_row = next(rows, None)
    headers = [str(value or "").strip() for value in (header_row or ())]
    if set(headers) != set(EXPECTED_HEADERS) or len(headers) != len(EXPECTED_HEADERS):
        raise StudentImportError("表头必须且只能包含：学号、姓名、班级")

    classes = list(classes)
    class_map = {str(class_group): class_group for class_group in classes}
    if len(classes) == 1:
        class_map[classes[0].name] = classes[0]
    parsed = []
    seen_numbers = set()
    for row_number, values in enumerate(rows, start=2):
        record = dict(zip(headers, values))
        if all(value is None or str(value).strip() == "" for value in record.values()):
            continue
        student_number = str(record["学号"] or "").strip()
        display_name = str(record["姓名"] or "").strip()
        class_name = str(record["班级"] or "").strip()
        if not student_number or not display_name or not class_name:
            raise StudentImportError(f"第 {row_number} 行存在空字段")
        if len(student_number) > 64:
            raise StudentImportError(f"第 {row_number} 行学号不能超过 64 个字符")
        if len(display_name) > 80:
            raise StudentImportError(f"第 {row_number} 行姓名不能超过 80 个字符")
        if student_number in seen_numbers:
            raise StudentImportError(f"第 {row_number} 行学号重复：{student_number}")
        class_group = class_map.get(class_name)
        if class_group is None:
            raise StudentImportError(f"第 {row_number} 行班级无法匹配：{class_name}")
        seen_numbers.add(student_number)
        parsed.append(
            ParsedStudent(student_number, display_name, class_group.pk, str(class_group))
        )
    if not parsed:
        raise StudentImportError("文件中没有可导入的学生")
    if allow_existing:
        return parsed

    numbers = [item.student_number for item in parsed]
    existing_rows = User.objects.filter(
        Q(student_number__in=numbers) | Q(username__in=numbers)
    ).values_list("student_number", "username")
    existing_numbers = {
        value
        for student_number, username in existing_rows
        for value in (student_number, username)
        if value in seen_numbers
    }
    for student_number in numbers:
        if student_number in existing_numbers:
            raise StudentImportError(f"学号已存在：{student_number}")
    return parsed
