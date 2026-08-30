import hashlib
import json
import csv
import io
import os
import subprocess
import zipfile
from datetime import timedelta
from types import SimpleNamespace

from django.conf import settings
from django.db import connection, models, transaction
from django.utils import timezone
from openpyxl import Workbook

from apps.assessment.models import (
    AssessmentEvent, Attempt, ExperimentSession, EventSchemaArtifact,
    StudentQuestionAssignment, ConfidenceAssignment, SimulationSnapshot,
    GradeRecord, ResponseInputObservation, ResponseSnapshot, SubmissionEvent,
    CaptureContractArtifact, ClientCaptureCheckpoint, ClientTimeSync,
)
from apps.experiments.models import (
    ClassExperimentTask,
    ConditionAssignment,
    ExperimentClassSnapshot,
    ExperimentVersion,
    ExperimentVersionParticipantSnapshot,
    ItemInstance,
    PublishPreviewClassResult,
    PublishPreviewRun,
    QuestionSnapshot,
    ReleaseSnapshotQuestionMapping,
    ScheduleRevision,
)
from .export_catalog import (
    EXPOSURE_EVENT_TYPES,
    PACKAGE_TABLES,
    PLACEHOLDER_TABLES,
    RESPONSE_EVENT_TYPES,
    SIMULATION_EVENT_TYPES,
    TABLE_PATHS,
)

from .models import (
    ExportJob, ExportSnapshot, RecordingCoverageAttempt,
    RecordingCoverageSession,
)
from .data_scope import DataScope


class ExportContractError(ValueError):
    pass


UNRECORDED_CLASS_FILTER = "__unrecorded__"


def _class_snapshot_filter_value(class_name):
    return "" if class_name == UNRECORDED_CLASS_FILTER else class_name


def _validate_export_scope(scope, package):
    if package in {"teaching", "kt", "cd"} and not scope.is_single_version:
        raise ExportContractError(f"{package.upper()} 导出必须选择单一实验版本")


def create_export_snapshot(*, course=None, experiment=None, research_study=None,
                           created_by, filter_definition=None, scope=None,
                           package=None):
    data_scope = scope
    if scope is not None:
        if scope.course.pk != course.pk:
            raise ExportContractError("导出范围不属于当前课程")
        if experiment is not None and scope.experiment_id not in {None, experiment.pk}:
            raise ExportContractError("导出范围与实验不一致")
        _validate_export_scope(scope, package)
        filter_definition = scope.serialize()
    versions = []
    version_queryset = ExperimentVersion.objects.none()
    if experiment is not None:
        version_queryset = experiment.versions.order_by("version_number")
    elif course is not None:
        version_queryset = ExperimentVersion.objects.filter(
            experiment__course=course
        ).order_by("experiment_id", "version_number")
    if scope is not None and scope.version_id is not None:
        version_queryset = version_queryset.filter(pk=scope.version_id)
    if scope is not None and scope.experiment_id is not None:
        version_queryset = version_queryset.filter(experiment_id=scope.experiment_id)
    versions = list(version_queryset.values_list("pk", "version_number"))
    cutoff = timezone.now()
    event_scope = {"session__version__experiment_id": experiment.pk} if experiment is not None else {"session__version__experiment__course_id": course.pk}
    if research_study is not None:
        course_ids = list(research_study.course_links.values_list("course_id", flat=True))
        event_scope = {"session__version__experiment__course_id__in": course_ids}
    watermark = AssessmentEvent.objects.filter(**event_scope, server_received_at__lte=cutoff).aggregate(max_seq=models.Max("server_ingest_sequence_no"))["max_seq"] or 0
    schema_hashes = {
        f"{row.event_type}:v{row.event_version}": row.schema_sha256
        for row in EventSchemaArtifact.objects.all().order_by("event_type", "event_version")
    }
    schema_manifest_hash = hashlib.sha256(json.dumps(schema_hashes, sort_keys=True).encode()).hexdigest()
    canonical_filter_definition = filter_definition or {}
    scope_hash = hashlib.sha256(
        json.dumps(canonical_filter_definition, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()
    coverage_run = None
    if course is not None and research_study is None:
        coverage_mapping = dict(filter_definition or {})
        if experiment is not None and not coverage_mapping.get("experiment_id"):
            coverage_mapping["experiment_id"] = experiment.pk
        coverage_scope = data_scope or DataScope.from_mapping(course, coverage_mapping)
        finalized_sessions = list(
            coverage_scope.sessions()
            .filter(
                status__in=("completed", "expired", "invalidated"),
                terminal_at__lte=cutoff - timedelta(seconds=300),
            )
            .select_related("capture_contract")
            .order_by("pk")
        )
        if finalized_sessions:
            from .recording_coverage import generate_scope_coverage

            coverage_run = generate_scope_coverage(
                finalized_sessions,
                scope_type="export",
                scope_definition={
                    "scope_hash": scope_hash,
                    "session_ids": [session.pk for session in finalized_sessions],
                },
                source_cutoff_at=cutoff,
            )

    return ExportSnapshot.objects.create(
        course=course,
        experiment=experiment,
        research_study=research_study,
        created_by=created_by,
        cutoff_at=cutoff,
        source_cutoff_at=cutoff,
        event_high_watermark=watermark,
        schema_manifest_hash=schema_manifest_hash,
        selected_versions=[{"id": pk, "version": version} for pk, version in versions],
        filter_definition=canonical_filter_definition,
        scope_hash=scope_hash,
        coverage_run=coverage_run,
    )


def create_export_job(*, course, snapshot, created_by, experiment=None,
                      research_study=None, kind=ExportJob.Kind.RESEARCH,
                      package="complete", selected_tables=None,
                      selected_fields=None):
    """Create, never reuse, an auditable export request."""
    return ExportJob.objects.create(
        course=course,
        experiment=experiment,
        research_study=research_study,
        snapshot=snapshot,
        created_by=created_by,
        kind=kind,
        package=package,
        selected_tables=selected_tables or [],
        selected_fields=selected_fields or [],
        status=ExportJob.Status.RUNNING,
        started_at=timezone.now(),
        progress_stage="snapshot",
    )


def remove_export_artifact(job, *, reason, invalidated=False):
    """Remove a materialized artifact while retaining its audit identity."""
    if job.artifact_path and os.path.exists(job.artifact_path):
        artifact_dir = os.path.dirname(job.artifact_path)
        os.remove(job.artifact_path)
        if artifact_dir and os.path.isdir(artifact_dir) and not os.listdir(artifact_dir):
            os.rmdir(artifact_dir)
    job.status = ExportJob.Status.INVALIDATED if invalidated else ExportJob.Status.EXPIRED
    job.artifact_path = ""
    job.artifact_removed_at = timezone.now()
    job.artifact_removal_reason = reason
    job.save(update_fields=[
        "status", "artifact_path", "artifact_removed_at", "artifact_removal_reason"
    ])
    return job


def _git_commit():
    try:
        return subprocess.check_output(["git", "rev-parse", "HEAD"], text=True, stderr=subprocess.DEVNULL).strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def build_export_manifest(snapshot, *, selected_tables, row_counts=None, data_status=None,
                          experiment_name=None, version_number=None,
                          session_ids=None, attempt_ids=None):
    schemas = {
        f"{row.event_type}:v{row.event_version}": row.schema_sha256
        for row in EventSchemaArtifact.objects.all()
    }
    config = {
        "selected_tables": list(selected_tables),
        "selected_versions": snapshot.selected_versions,
        "filter_definition": snapshot.filter_definition,
    }
    filter_display = dict(snapshot.filter_definition)
    if filter_display.get("class_name") == UNRECORDED_CLASS_FILTER:
        filter_display["class_name"] = "未记录班级"
    provenance = {
        "export_snapshot": {"id": str(snapshot.snapshot_id), "source_cutoff_at": snapshot.source_cutoff_at.isoformat(), "event_high_watermark": snapshot.event_high_watermark},
        "schema": {"manifest_hash": snapshot.schema_manifest_hash, "event_schema_hashes": schemas},
        "derivation": {"version": "v1", "deterministic": True, "code_sha": _git_commit()},
        "source_records": {"sessions": "assessment.ExperimentSession", "attempts": "assessment.Attempt", "events": "assessment.AssessmentEvent", "grades": "assessment.GradeRecord", "confidence": "assessment.ConfidenceAssignment"},
    }
    coverage_filter = {
        "session__version__experiment_id": snapshot.experiment_id
    } if snapshot.experiment_id else {
        "session__version__experiment__course_id": snapshot.course_id
    }
    coverage_query = RecordingCoverageSession.objects.filter(
            **coverage_filter,
            run__run_status="completed",
            run__source_cutoff_at__lte=snapshot.source_cutoff_at,
        )
    if getattr(snapshot, "coverage_run_id", None):
        coverage_query = coverage_query.filter(run_id=snapshot.coverage_run_id)
    if session_ids is not None:
        coverage_query = coverage_query.filter(session_id__in=session_ids)
    all_coverage_results = list(
        coverage_query.select_related("run", "session__capture_contract").order_by(
            "session_id", "-run__generated_at"
        )
    )
    latest_by_session = {}
    for item in all_coverage_results:
        latest_by_session.setdefault(item.session_id, item)
    coverage_results = list(latest_by_session.values())
    coverage_runs = {str(item.run_id): item.run for item in coverage_results}
    contract_rows = {
        str(item.session.capture_contract_id): item.session.capture_contract
        for item in coverage_results
        if item.session.capture_contract_id
    }
    status_counts = {}
    for item in coverage_results:
        status_counts[item.overall_status] = status_counts.get(item.overall_status, 0) + 1
    run_ids = sorted(coverage_runs)
    coverage_versions = sorted(
        {run.coverage_version for run in coverage_runs.values()}
    )
    parameter_hashes = sorted({run.parameter_hash for run in coverage_runs.values()})
    coverage_attempt_query = RecordingCoverageAttempt.objects.filter(
        run_id__in=run_ids
    )
    if attempt_ids is not None:
        coverage_attempt_query = coverage_attempt_query.filter(
            attempt_id__in=attempt_ids
        )
    time_quality_counts = {}
    for quality in coverage_attempt_query.values_list("time_quality", flat=True):
        time_quality_counts[quality] = time_quality_counts.get(quality, 0) + 1
    checkpoint_query = ClientCaptureCheckpoint.objects.filter(
        server_received_at__lte=snapshot.source_cutoff_at
    )
    if session_ids is not None:
        checkpoint_query = checkpoint_query.filter(session_id__in=session_ids)
    elif snapshot.experiment_id:
        checkpoint_query = checkpoint_query.filter(
            session__version__experiment_id=snapshot.experiment_id
        )
    else:
        checkpoint_query = checkpoint_query.filter(
            session__version__experiment__course_id=snapshot.course_id
        )
    latest_checkpoint_by_instance = {}
    for checkpoint in checkpoint_query.order_by(
        "session_id", "client_instance_id", "-server_received_at", "-checkpoint_id"
    ):
        latest_checkpoint_by_instance.setdefault(
            (checkpoint.session_id, checkpoint.client_instance_id), checkpoint
        )
    return {
        "manifest_version": 2,
        "contains_direct_identifiers": True,
        "identity_columns": ["student_number", "student_name"],
        "primary_question_identifier": "question_code",
        "export_snapshot_id": str(snapshot.snapshot_id),
        "course_id": snapshot.course_id,
        "experiment_id": snapshot.experiment_id,
        "experiment_name": experiment_name,
        "version_number": version_number,
        "research_study_id": snapshot.research_study_id,
        "cutoff_at": snapshot.cutoff_at.isoformat(),
        "source_cutoff_at": snapshot.source_cutoff_at.isoformat(),
        "event_high_watermark": snapshot.event_high_watermark,
        "schema_manifest_hash": snapshot.schema_manifest_hash,
        "scope_hash": snapshot.scope_hash,
        "database_snapshot_strategy": snapshot.database_snapshot_strategy,
        "filter_definition": snapshot.filter_definition,
        "filter_display": filter_display,
        "selected_tables": list(selected_tables),
        "event_schema_hashes": schemas,
        "application_git_commit_sha": _git_commit(),
        "export_config_hash": hashlib.sha256(json.dumps(config, sort_keys=True).encode()).hexdigest(),
        "row_counts": row_counts or {},
        "data_status": data_status or {},
        "provenance": provenance,
        "capture_contract_ids": sorted(contract_rows),
        "capture_contract_hashes": sorted({row.schema_sha256 for row in contract_rows.values()}),
        "coverage_run_id": run_ids[0] if len(run_ids) == 1 else None,
        "coverage_run_ids": run_ids,
        "coverage_version": coverage_versions[0] if len(coverage_versions) == 1 else None,
        "coverage_parameter_hash": parameter_hashes[0] if len(parameter_hashes) == 1 else None,
        "coverage_source_cutoff": max((run.source_cutoff_at for run in coverage_runs.values()), default=None).isoformat() if run_ids else None,
        "coverage_source_high_water_marks": {run_id: coverage_runs[run_id].source_high_water_marks for run_id in run_ids},
        "coverage_status_counts": status_counts,
        "sequence_gap_count": sum(item.sequence_gap_count for item in coverage_results),
        "tail_gap_count": sum(bool(item.tail_gap_detected) for item in coverage_results),
        "dead_letter_count": sum(
            checkpoint.failed_event_count
            for checkpoint in latest_checkpoint_by_instance.values()
        ),
        "time_quality_counts": time_quality_counts,
    }


def manifest_bytes(manifest):
    return json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8")


def sha256_bytes(value):
    return hashlib.sha256(value).hexdigest()


def claim_export_job(*, worker_id, now=None):
    now = now or timezone.now()
    with transaction.atomic():
        job = (
            ExportJob.objects.select_for_update()
            .filter(
                models.Q(status=ExportJob.Status.QUEUED)
                | models.Q(status=ExportJob.Status.RUNNING, lease_expires_at__lt=now)
            )
            .order_by("created_at")
            .first()
        )
        if not job:
            return None
        job.status = ExportJob.Status.RUNNING
        job.locked_by = worker_id
        job.started_at = now
        job.heartbeat_at = now
        job.lease_expires_at = now + timedelta(seconds=90)
        job.progress_stage = "snapshot"
        job.save(update_fields=["status", "locked_by", "started_at", "heartbeat_at", "lease_expires_at", "progress_stage"])
        return job


def renew_export_job_lease(job, *, worker_id, now=None):
    now = now or timezone.now()
    updated = ExportJob.objects.filter(
        pk=job.pk,
        status=ExportJob.Status.RUNNING,
        locked_by=worker_id,
    ).update(
        heartbeat_at=now,
        lease_expires_at=now + timedelta(seconds=90),
    )
    if not updated:
        return None
    job.refresh_from_db(fields=["heartbeat_at", "lease_expires_at"])
    return job


def _csv_bytes(headers, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def run_export_job(job):
    """Build a deterministic, identified ZIP from the job's frozen cutoff."""
    snapshot = job.snapshot
    if snapshot.invalidated_at is not None:
        raise ExportContractError("ExportSnapshot has been invalidated and cannot be rerun.")
    root = os.path.join(settings.MEDIA_ROOT, "exports", str(job.job_id))
    os.makedirs(root, exist_ok=True)
    artifact_extension = "xlsx" if job.package == "teaching" else "zip"
    artifact_path = os.path.join(root, f"{job.package}-{job.job_id}.{artifact_extension}")
    study_course_ids = list(job.research_study.course_links.values_list("course_id", flat=True)) if job.research_study_id else []
    study_student_ids = list(job.research_study.participants.values_list("student_id", flat=True)) if job.research_study_id else []
    filter_definition = snapshot.filter_definition or {}
    scope = None
    if filter_definition.get("scope_version"):
        try:
            scope = DataScope.from_mapping(snapshot.course, filter_definition)
        except (LookupError, ValueError) as error:
            raise ExportContractError(f"冻结的数据范围无效：{error}") from error
        session_queryset = scope.sessions()
    else:
        session_filter = {"version__experiment_id": job.experiment_id} if job.experiment_id else ({"version__experiment__course_id__in": study_course_ids, "student_id__in": study_student_ids} if study_course_ids else {"version__experiment__course_id": job.course_id})
        if filter_definition.get("version_id"):
            session_filter["version_id"] = int(filter_definition["version_id"])
        if filter_definition.get("class_name"):
            session_filter["class_snapshot_name"] = _class_snapshot_filter_value(filter_definition["class_name"])
        if filter_definition.get("student_number"):
            session_filter["student__student_number"] = filter_definition["student_number"]
        if filter_definition.get("session_id"):
            session_filter["pk"] = int(filter_definition["session_id"])
        if filter_definition.get("status") in {choice for choice, _label in ExperimentSession.Status.choices}:
            session_filter["status"] = filter_definition["status"]
        session_queryset = ExperimentSession.objects.filter(**session_filter)
    with transaction.atomic():
        if connection_vendor := connection.vendor:
            if connection_vendor == "postgresql":
                with connection.cursor() as cursor:
                    cursor.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
        sessions = list(
            session_queryset.filter(started_at__lte=snapshot.source_cutoff_at)
            .select_related("student", "version")
        )
        session_ids = [s.pk for s in sessions]
        attempt_queryset = (
            scope.attempts() if scope is not None else Attempt.objects.filter(session_id__in=session_ids)
        )
        attempt_queryset = attempt_queryset.filter(
            session_id__in=session_ids,
            session__started_at__lte=snapshot.source_cutoff_at,
        ).filter(
            models.Q(submitted_at__isnull=True)
            | models.Q(submitted_at__lte=snapshot.source_cutoff_at)
        )
        legacy_attempt_status = filter_definition.get("status") in {
            choice for choice, _label in Attempt.Status.choices
        }
        if scope is None and legacy_attempt_status:
            attempt_queryset = attempt_queryset.filter(status=filter_definition["status"])
        attempts = list(attempt_queryset.select_related("session", "assignment__snapshot", "assignment__item_instance"))
        if (scope is not None and scope.attempt_status is not None) or legacy_attempt_status:
            selected_session_ids = {attempt.session_id for attempt in attempts}
            sessions = [session for session in sessions if session.pk in selected_session_ids]
            session_ids = [session.pk for session in sessions]
        event_queryset = scope.events() if scope is not None else AssessmentEvent.objects.all()
        events = list(
            event_queryset.filter(
                session_id__in=session_ids,
                server_received_at__lte=snapshot.source_cutoff_at,
                server_ingest_sequence_no__lte=snapshot.event_high_watermark,
            ).order_by("session_id", "server_ingest_sequence_no")
        )
        assignments_queryset = StudentQuestionAssignment.objects.filter(session_id__in=session_ids).select_related("snapshot", "item_instance")
        if (scope is not None and scope.attempt_status is not None) or legacy_attempt_status:
            assignments_queryset = assignments_queryset.filter(attempt__in=attempts)
        assignments = list(assignments_queryset)
        _validate_question_codes(assignments)
        confidence_assignments = list(
            ConfidenceAssignment.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                assigned_at__lte=snapshot.source_cutoff_at,
            )
        )
        simulations = list(
            SimulationSnapshot.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                created_at__lte=snapshot.source_cutoff_at,
            )
        )
        response_snapshots = list(
            ResponseSnapshot.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                created_at__lte=snapshot.source_cutoff_at,
            ).order_by("attempt_id", "created_at", "snapshot_id")
        )
        input_observations = list(
            ResponseInputObservation.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                observed_at__lte=snapshot.source_cutoff_at,
            ).select_related("source_event").order_by(
                "attempt_id", "client_instance_id", "client_sequence_no"
            )
        )
        submissions = list(
            SubmissionEvent.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                server_received_at__lte=snapshot.source_cutoff_at,
            )
        )
        grade_records = list(
            GradeRecord.objects.filter(
                submission__attempt_id__in=[a.pk for a in attempts],
                created_at__lte=snapshot.source_cutoff_at,
            ).select_related("submission")
        )
        version_ids = {session.version_id for session in sessions}
        class_snapshots = list(
            ExperimentClassSnapshot.objects.filter(
                version_id__in=version_ids,
                created_at__lte=snapshot.source_cutoff_at,
            ).order_by("version_id", "created_at")
        )
        participant_snapshots = list(
            ExperimentVersionParticipantSnapshot.objects.filter(
                version_id__in=version_ids,
                student_id__in=[session.student_id for session in sessions],
                created_at__lte=snapshot.source_cutoff_at,
            ).order_by("version_id", "student_id")
        )
        condition_assignments = list(
            ConditionAssignment.objects.filter(
                models.Q(session_id__in=session_ids)
                | models.Q(
                    subject_id__in=[session.student_id for session in sessions],
                    condition__experiment_id__in=[
                        session.version.experiment_id for session in sessions
                    ],
                ),
                assigned_at__lte=snapshot.source_cutoff_at,
            ).select_related("condition")
        )
    roster_rows = []
    if job.kind == ExportJob.Kind.IDENTITY:
        selected_tables = ["student_roster"]
        students = list(job.experiment.students.select_related("class_group")) if job.experiment_id else []
        roster_rows = [[student.student_number, student.shown_name, str(student.class_group or "")] for student in students]
    elif job.selected_tables:
        selected_tables = list(job.selected_tables)
    else:
        selected_tables = list(PACKAGE_TABLES.get(job.package, PACKAGE_TABLES["complete"]))
    rows = _table_rows(
        job,
        sessions,
        assignments,
        attempts,
        events,
        confidence_assignments,
        simulations,
        response_snapshots=response_snapshots,
        input_observations=input_observations,
        submissions=submissions,
        grade_records=grade_records,
        condition_assignments=condition_assignments,
        class_snapshots=class_snapshots,
        participant_snapshots=participant_snapshots,
        snapshot=snapshot,
    )
    row_counts = {
        table: len(roster_rows) if table == "student_roster" else len(rows[table][1])
        for table in selected_tables
        if table == "student_roster" or table in rows
    }
    data_status = {
        table: ("no_observation" if row_counts.get(table, 0) == 0 else "observed")
        for table in selected_tables
    }
    selected_version_number = None
    if filter_definition.get("version_id"):
        selected_version_number = next(
            (item["version"] for item in snapshot.selected_versions if str(item["id"]) == str(filter_definition["version_id"])),
            None,
        )
    elif len(snapshot.selected_versions) == 1:
        selected_version_number = snapshot.selected_versions[0]["version"]
    manifest = build_export_manifest(
        snapshot,
        selected_tables=selected_tables,
        row_counts=row_counts,
        data_status=data_status,
        experiment_name=job.experiment.name if job.experiment_id else None,
        version_number=selected_version_number,
        session_ids=session_ids,
        attempt_ids=[attempt.pk for attempt in attempts],
    )
    if job.package == "teaching":
        _write_teaching_workbook(
            artifact_path,
            rows["attempts"],
            manifest,
            attempts,
            events=events,
            submissions=submissions,
        )
        return _finalize_export_job(job, artifact_path, manifest)
    with zipfile.ZipFile(artifact_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("00_manifest/manifest.json", manifest_bytes(manifest))
        archive.writestr("00_manifest/derivation_runs.json", manifest_bytes({
            "derivation_version": "v1",
            "deterministic": True,
            "source_snapshot_id": str(snapshot.snapshot_id),
            "source_cutoff_at": snapshot.source_cutoff_at.isoformat(),
            "event_high_watermark": snapshot.event_high_watermark,
            "schema_manifest_hash": snapshot.schema_manifest_hash,
            "application_git_commit_sha": _git_commit(),
        }))
        archive.writestr("01_design_metadata/README.txt", "This package contains direct student identifiers and must be handled as identified research data.\n")
        if job.kind == ExportJob.Kind.IDENTITY:
            archive.writestr(
                "01_identity/student_roster.csv",
                _csv_bytes(
                    ["student_number", "student_name", "class_name"],
                    roster_rows,
                ),
            )
            return _finalize_export_job(job, artifact_path, manifest)
        for table in selected_tables:
            if table == "identity_mapping" or table not in rows:
                continue
            path = TABLE_PATHS.get(table)
            if path:
                headers, values = rows[table]
                archive.writestr(path, _csv_bytes(headers, values))
    return _finalize_export_job(job, artifact_path, manifest)


def _write_teaching_workbook(
    artifact_path, attempt_table, manifest, attempts, *, events, submissions
):
    """Write the teaching-facing grade workbook from the same frozen rows.

    The source is the canonical attempts export table, whose grade columns are
    already restricted to the effective GradeRecord.  The second sheet keeps
    the snapshot provenance with the downloadable workbook rather than
    requiring a separate, mutable dashboard lookup.
    """
    headers, rows = attempt_table
    requested = (
        "student_number", "student_name", "question_code", "session_id",
        "status", "grade_score", "grade_max_score", "grade_is_correct",
        "submission_grade_status", "graded_at",
    )
    positions = [headers.index(column) for column in requested if column in headers]
    attempt_id_position = headers.index("attempt_id")
    events_by_attempt = {}
    for event in events:
        events_by_attempt.setdefault(event.attempt_id, []).append(event)
    submissions_by_attempt = {
        submission.attempt_id: submission
        for submission in submissions
    }
    from .teaching_data import summarize_attempt_timing

    timing_by_attempt = {
        attempt.pk: summarize_attempt_timing(
            attempt,
            events=events_by_attempt.get(attempt.pk, []),
            submission=submissions_by_attempt.get(attempt.pk),
        )
        for attempt in attempts
    }
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("教学成绩")
    sheet.append([
        *[headers[position] for position in positions],
        "display_duration_ms", "display_kind", "interval_quality",
    ])
    for row in rows:
        timing = timing_by_attempt.get(row[attempt_id_position])
        sheet.append([
            *[row[position] for position in positions],
            timing.display_duration_ms if timing else None,
            timing.display_kind if timing else "unavailable",
            timing.interval_quality if timing else "unknown",
        ])
    lineage = workbook.create_sheet("导出血缘")
    lineage.append(["field", "value"])
    for key in (
        "export_snapshot_id", "source_cutoff_at", "event_high_watermark",
        "schema_manifest_hash", "scope_hash", "contains_direct_identifiers",
    ):
        lineage.append([key, json.dumps(manifest.get(key), ensure_ascii=False)])
    workbook.save(artifact_path)


def _validate_question_codes(assignments):
    seen = {}
    for assignment in assignments:
        snapshot = assignment.snapshot
        code = snapshot.code.strip()
        if not code:
            raise ExportContractError(f"题目快照 {snapshot.pk} 缺少 question_code")
        key = (snapshot.version_id, code)
        previous_snapshot_id = seen.setdefault(key, snapshot.pk)
        if previous_snapshot_id != snapshot.pk:
            raise ExportContractError(f"同一实验版本题目编码重复：{code}")


def preview_export_table(*, course, experiment, table, version_id=None,
                         class_name=None, student_number=None, session_id=None, status=None,
                         page=1, page_size=20, scope=None):
    """Return a paginated export preview.

    ``scope`` is the data-center path and is deliberately preferred over the
    legacy scalar parameters.  Keeping the preview on the same ``DataScope``
    as teaching, quality, and ZIP export prevents a table view from showing a
    wider population than the visible range bar claims.
    """
    if scope is not None:
        if scope.course.pk != course.pk:
            raise ExportContractError("预览范围不属于当前课程")
        if experiment is not None and scope.experiment_id not in {None, experiment.pk}:
            raise ExportContractError("预览范围与实验不一致")
        filter_definition = scope.serialize()
        sessions = list(
            scope.sessions().select_related("student", "version").order_by("pk")
        )
    else:
        filter_definition = {
            key: value
            for key, value in {
                "version_id": version_id,
                "class_name": class_name,
                "student_number": student_number,
                "session_id": session_id,
                "status": status,
            }.items()
            if value not in (None, "")
        }
        session_filter = {"version__experiment_id": experiment.pk}
        if version_id:
            session_filter["version_id"] = int(version_id)
        if class_name:
            session_filter["class_snapshot_name"] = _class_snapshot_filter_value(class_name)
        if student_number:
            session_filter["student__student_number"] = student_number
        if session_id:
            session_filter["pk"] = int(session_id)
        if status in {choice for choice, _label in ExperimentSession.Status.choices}:
            session_filter["status"] = status
        sessions = list(
            ExperimentSession.objects.filter(**session_filter)
            .select_related("student", "version")
            .order_by("pk")
        )
    session_ids = [session.pk for session in sessions]
    attempt_queryset = (scope.attempts() if scope is not None else Attempt.objects.filter(
        session_id__in=session_ids
    )).select_related(
        "session", "assignment__snapshot", "assignment__item_instance"
    )
    has_attempt_status = (
        scope is not None and scope.attempt_status is not None
    ) or (scope is None and status in {choice for choice, _label in Attempt.Status.choices})
    if scope is None and status in {choice for choice, _label in Attempt.Status.choices}:
        attempt_queryset = attempt_queryset.filter(status=status)
    attempts = list(attempt_queryset.order_by("pk"))
    if has_attempt_status:
        selected_session_ids = {attempt.session_id for attempt in attempts}
        sessions = [session for session in sessions if session.pk in selected_session_ids]
        session_ids = [session.pk for session in sessions]
    attempt_ids = {attempt.pk for attempt in attempts}
    assignments = list(
        StudentQuestionAssignment.objects.filter(session_id__in=session_ids)
        .select_related("snapshot", "item_instance")
        .order_by("session_id", "assigned_sequence_index")
    )
    if has_attempt_status:
        assignment_ids = {attempt.assignment_id for attempt in attempts}
        assignments = [assignment for assignment in assignments if assignment.pk in assignment_ids]
    events = list(
        (scope.events() if scope is not None else AssessmentEvent.objects.all()).filter(
            session_id__in=session_ids,
            server_received_at__lte=timezone.now(),
        ).order_by("session_id", "server_ingest_sequence_no")
    )
    confidence_assignments = list(ConfidenceAssignment.objects.filter(attempt_id__in=attempt_ids))
    simulations = list(SimulationSnapshot.objects.filter(attempt_id__in=attempt_ids))
    response_snapshots = list(
        ResponseSnapshot.objects.filter(attempt_id__in=attempt_ids).order_by(
            "attempt_id", "created_at", "snapshot_id"
        )
    )
    input_observations = list(
        ResponseInputObservation.objects.filter(
            attempt_id__in=attempt_ids
        ).select_related("source_event").order_by(
            "attempt_id", "client_instance_id", "client_sequence_no"
        )
    )
    submissions = list(SubmissionEvent.objects.filter(attempt_id__in=attempt_ids))
    grade_records = list(
        GradeRecord.objects.filter(submission__attempt_id__in=attempt_ids).select_related("submission")
    )
    version_ids = {session.version_id for session in sessions}
    class_snapshots = ExperimentClassSnapshot.objects.filter(version_id__in=version_ids)
    if scope is not None and scope.class_snapshot_id is not None:
        class_snapshots = class_snapshots.filter(pk=scope.class_snapshot_id)
    class_snapshots = list(class_snapshots.order_by("version_id", "created_at"))
    participant_snapshots = list(
        ExperimentVersionParticipantSnapshot.objects.filter(
            version_id__in=version_ids,
            student_id__in=[session.student_id for session in sessions],
        ).order_by("version_id", "student_id")
    )
    condition_assignments = list(
        ConditionAssignment.objects.filter(
            models.Q(session_id__in=session_ids)
            | models.Q(
                subject_id__in=[session.student_id for session in sessions],
                condition__experiment_id__in=[session.version.experiment_id for session in sessions],
            )
        ).select_related("condition")
    )
    fake_job = SimpleNamespace(course_id=course.pk, research_study_id=None)
    rows = _table_rows(
        fake_job,
        sessions,
        assignments,
        attempts,
        events,
        confidence_assignments,
        simulations,
        response_snapshots=response_snapshots,
        input_observations=input_observations,
        submissions=submissions,
        grade_records=grade_records,
        condition_assignments=condition_assignments,
        class_snapshots=class_snapshots,
        participant_snapshots=participant_snapshots,
    )
    headers, values = rows.get(table, ((), []))
    total = len(values)
    try:
        page = max(1, int(page or 1))
    except (TypeError, ValueError):
        page = 1
    try:
        page_size = max(1, min(100, int(page_size or 20)))
    except (TypeError, ValueError):
        page_size = 20
    start = (page - 1) * page_size
    return {
        "table": table,
        "headers": headers,
        "rows": values[start:start + page_size],
        "total": total,
        "page": page,
        "page_size": page_size,
        "has_previous": page > 1,
        "has_next": start + page_size < total,
        "previous_page": max(1, page - 1),
        "next_page": page + 1,
        "filter_definition": filter_definition,
        "data_status": "placeholder" if table in PLACEHOLDER_TABLES else ("no_observation" if total == 0 else "observed"),
    }


def _table_rows(job, sessions, assignments, attempts, events, confidence_assignments, simulations,
                *, response_snapshots=(), input_observations=(), submissions=(),
                grade_records=(), condition_assignments=(), class_snapshots=(),
                participant_snapshots=(), snapshot=None):
    identity_headers = ("student_number", "student_name")
    session_by_id = {session.pk: session for session in sessions}
    attempt_by_assignment = {attempt.assignment_id: attempt for attempt in attempts}
    attempt_by_id = {attempt.pk: attempt for attempt in attempts}
    snapshot_code_by_id = {assignment.snapshot_id: assignment.snapshot.code for assignment in assignments}

    def identity_for_session(session_id):
        session = session_by_id[session_id]
        return [session.student.student_number or "", session.student.shown_name]

    def question_code_for_snapshot(snapshot_id):
        return snapshot_code_by_id.get(snapshot_id, "")

    subject_values = []
    seen_student_numbers = set()
    for session in sessions:
        student_number = session.student.student_number or ""
        if student_number not in seen_student_numbers:
            subject_values.append([student_number, session.student.shown_name])
            seen_student_numbers.add(student_number)

    selected_attempt_ids = set(attempt_by_id)
    selected_events = [event for event in events if event.attempt_id is None or event.attempt_id in selected_attempt_ids]
    selected_confidence = [assignment for assignment in confidence_assignments if assignment.attempt_id in selected_attempt_ids]
    selected_simulations = [simulation for simulation in simulations if simulation.attempt_id in selected_attempt_ids]

    cutoff = snapshot.source_cutoff_at if snapshot is not None else timezone.now()
    capture_contracts = list(
        CaptureContractArtifact.objects.filter(
            capture_contract_id__in={
                session.capture_contract_id
                for session in sessions
                if session.capture_contract_id
            }
        ).order_by("contract_name")
    )
    time_syncs = list(
        ClientTimeSync.objects.filter(
            session_id__in=session_by_id,
            created_at__lte=cutoff,
        ).order_by("session_id", "client_instance_id", "created_at")
    )
    valid_syncs_by_instance = {}
    for sync in time_syncs:
        if (
            sync.sync_quality == ClientTimeSync.SyncQuality.VALID
            and sync.estimated_clock_offset_ms is not None
        ):
            valid_syncs_by_instance.setdefault(
                (str(sync.client_instance_id), sync.client_time_origin), []
            ).append(sync)

    def derived_event_time(event):
        if event.estimated_occurred_at_utc is not None:
            return event.estimated_occurred_at_utc
        if (
            event.client_instance_id is None
            or event.client_time_origin is None
            or event.client_monotonic_ms is None
        ):
            return None
        candidates = valid_syncs_by_instance.get(
            (str(event.client_instance_id), event.client_time_origin), []
        )
        if not candidates:
            return None
        sample = min(
            candidates,
            key=lambda row: abs(row.client_monotonic_ms - event.client_monotonic_ms),
        )
        delta_ms = event.client_monotonic_ms - sample.client_monotonic_ms
        if abs(delta_ms) > 20 * 60 * 1000:
            return None
        return (
            sample.client_wall_time_utc
            + timedelta(milliseconds=sample.estimated_clock_offset_ms)
            + timedelta(milliseconds=float(delta_ms))
        )
    checkpoints = list(
        ClientCaptureCheckpoint.objects.filter(
            session_id__in=session_by_id,
            server_received_at__lte=cutoff,
        ).order_by("session_id", "client_instance_id", "server_received_at")
    )
    coverage_sessions = list(
        RecordingCoverageSession.objects.filter(
            session_id__in=session_by_id,
            run__run_status="completed",
            run__source_cutoff_at__lte=cutoff,
        ).filter(
            **({"run_id": snapshot.coverage_run_id} if getattr(snapshot, "coverage_run_id", None) else {})
        ).select_related("run").order_by("session_id", "-run__generated_at")
    )
    coverage_attempts = list(
        RecordingCoverageAttempt.objects.filter(
            attempt_id__in=selected_attempt_ids,
            run__run_status="completed",
            run__source_cutoff_at__lte=cutoff,
        ).filter(
            **({"run_id": snapshot.coverage_run_id} if getattr(snapshot, "coverage_run_id", None) else {})
        ).select_related("run", "attempt").order_by("attempt_id", "-run__generated_at")
    )
    latest_coverage_session = {}
    for result in coverage_sessions:
        latest_coverage_session.setdefault(result.session_id, result)
    latest_coverage_attempt = {}
    for result in coverage_attempts:
        latest_coverage_attempt.setdefault(result.attempt_id, result)

    event_headers = (
        *identity_headers, "student_id", "question_code", "event_id", "event_type", "event_version", "session_id", "attempt_id",
        "client_instance_id", "client_sequence_no", "server_ingest_sequence_no", "correlation_id",
        "server_received_at", "capture_mode", "quality_flags", "payload", "event_domain",
        "actor_source", "envelope_version", "assignment_id", "question_snapshot_id", "item_instance_id",
        "client_wall_time_utc", "client_monotonic_ms", "client_time_origin", "time_sync_id",
        "estimated_clock_offset_ms", "estimated_rtt_ms", "estimated_occurred_at_utc", "derived_occurred_at_utc", "server_committed_at",
        "app_version",
    )

    def event_row(event):
        return [
            *identity_for_session(event.session_id), event.student_id,
            question_code_for_snapshot(event.question_snapshot_id),
            event.event_id, event.event_type, event.event_version, event.session_id, event.attempt_id,
            event.client_instance_id, event.client_sequence_no, event.server_ingest_sequence_no,
            event.correlation_id, event.server_received_at.isoformat(), event.capture_mode,
            json.dumps(event.quality_flags, ensure_ascii=False), json.dumps(event.payload, ensure_ascii=False),
            event.event_domain, event.actor_source, event.envelope_version, event.assignment_id,
            event.question_snapshot_id, event.item_instance_id,
            event.client_wall_time_utc.isoformat() if event.client_wall_time_utc else "",
            event.client_monotonic_ms, event.client_time_origin.isoformat() if event.client_time_origin else "",
            event.time_sync_id, event.estimated_clock_offset_ms, event.estimated_rtt_ms,
            event.estimated_occurred_at_utc.isoformat() if event.estimated_occurred_at_utc else "",
            derived_event_time(event).isoformat() if derived_event_time(event) else "",
            event.server_committed_at.isoformat() if event.server_committed_at else "", event.app_version,
        ]

    event_rows = [event_row(event) for event in selected_events]
    compact_event_headers = (*identity_headers, "question_code", "event_id", "event_type", "session_id", "attempt_id", "server_received_at", "payload")

    def compact_event_row(event):
        return [
            *identity_for_session(event.session_id), question_code_for_snapshot(event.question_snapshot_id),
            event.event_id, event.event_type, event.session_id, event.attempt_id,
            event.server_received_at.isoformat(), json.dumps(event.payload, ensure_ascii=False),
        ]

    projection_headers = (*event_headers, "time_quality", "coverage_run_id")

    def projection_row(event):
        coverage = latest_coverage_attempt.get(event.attempt_id) or latest_coverage_session.get(event.session_id)
        return [
            *event_row(event),
            getattr(coverage, "time_quality", None) or event.interval_quality,
            coverage.run_id if coverage else "",
        ]

    response_headers = (
        *projection_headers, "before_state", "after_state", "before_state_hash",
        "after_state_hash", "response_snapshot_id", "question_type", "change_source",
    )
    response_rows = [
        [
            *projection_row(event),
            json.dumps(event.payload.get("response_state_before"), ensure_ascii=False),
            json.dumps(event.payload.get("response_state_after"), ensure_ascii=False),
            event.payload.get("before_state_hash", ""),
            event.payload.get("after_state_hash") or event.payload.get("state_hash", ""),
            event.payload.get("response_snapshot_id", ""),
            event.payload.get("question_type", ""),
            event.payload.get("change_source", ""),
        ]
        for event in selected_events if event.event_type in RESPONSE_EVENT_TYPES
    ]
    exposure_rows = [projection_row(event) for event in selected_events if event.event_type in EXPOSURE_EVENT_TYPES]
    navigation_headers = (
        *projection_headers, "from_assignment_id", "to_assignment_id",
        "entry_reason", "leave_reason",
    )
    navigation_rows = [
        [
            *projection_row(event), event.payload.get("from_assignment_id", ""),
            event.payload.get("to_assignment_id", ""), event.payload.get("entry_reason", ""),
            event.payload.get("leave_reason", ""),
        ]
        for event in selected_events
        if event.event_type in ("item_entered", "item_rendered", "item_left")
    ]
    telemetry_rows = [projection_row(event) for event in selected_events if event.event_domain == AssessmentEvent.EventDomain.TELEMETRY]
    simulation_event_rows = [compact_event_row(event) for event in selected_events if event.event_type in SIMULATION_EVENT_TYPES]

    def identity_and_question_for_attempt(attempt_id):
        attempt = attempt_by_id[attempt_id]
        return [*identity_for_session(attempt.session_id), attempt.assignment.snapshot.code]

    def input_observation_row(item):
        event = item.source_event if item.source_event_id else None
        coverage = latest_coverage_attempt.get(item.attempt_id)
        derived_time = derived_event_time(event) if event is not None else None
        return [
            *identity_and_question_for_attempt(item.attempt_id),
            item.observation_id,
            item.source_event_id or "",
            event.event_type if event else "",
            event.event_version if event else "",
            event.envelope_version if event else "",
            item.attempt_id,
            item.client_instance_id,
            item.client_sequence_no,
            item.client_monotonic_ms,
            item.batch_start_client_monotonic_ms,
            item.batch_end_client_monotonic_ms,
            event.client_wall_time_utc.isoformat()
            if event and event.client_wall_time_utc else "",
            event.client_time_origin.isoformat()
            if event and event.client_time_origin else "",
            event.time_sync_id if event else "",
            event.estimated_occurred_at_utc.isoformat()
            if event and event.estimated_occurred_at_utc else "",
            derived_time.isoformat() if derived_time else "",
            event.server_received_at.isoformat() if event else "",
            event.server_committed_at.isoformat()
            if event and event.server_committed_at else "",
            event.server_ingest_sequence_no if event else "",
            event.actor_source if event else "",
            event.capture_mode if event else "",
            json.dumps(event.quality_flags, ensure_ascii=False) if event else "",
            (coverage.time_quality if coverage else event.interval_quality)
            if event else "unknown",
            coverage.run_id if coverage else "",
            event.app_version if event else "",
            item.edit_operation,
            item.content_hash_before,
            item.content_hash_after,
            item.content_length_before,
            item.content_length_after,
            item.input_method,
            item.is_ime,
            item.observed_at.isoformat(),
        ]

    grade_by_attempt = {}
    if not grade_records and selected_attempt_ids:
        grade_records = GradeRecord.objects.filter(
            submission__attempt_id__in=selected_attempt_ids,
            created_at__lte=(snapshot.source_cutoff_at if snapshot is not None else timezone.now()),
        ).select_related("submission")
    for grade in grade_records:
        if grade.validity_status == GradeRecord.ValidityStatus.EFFECTIVE:
            grade_by_attempt[grade.submission.attempt_id] = grade
    submission_by_attempt = {submission.attempt_id: submission for submission in submissions}
    task_ids = {session.class_task_id for session in sessions if session.class_task_id}
    class_tasks = list(
        ClassExperimentTask.objects.filter(pk__in=task_ids)
        .select_related(
            "class_group",
            "release_snapshot",
            "randomization_identity",
            "preview_class_result__preview_run",
        )
        .order_by("pk")
    )
    schedule_revisions = list(
        ScheduleRevision.objects.filter(class_task_id__in=task_ids)
        .select_related("class_task")
        .order_by("class_task_id", "revision_no")
    )
    release_ids = {session.version_id for session in sessions} | {
        task.release_snapshot_id for task in class_tasks
    }
    releases = list(
        ExperimentVersion.objects.filter(pk__in=release_ids)
        .select_related("experiment")
        .order_by("experiment_id", "version_number")
    )
    release_mappings = list(
        ReleaseSnapshotQuestionMapping.objects.filter(release_snapshot_id__in=release_ids)
        .select_related("source_question", "question_snapshot")
        .order_by("release_snapshot_id", "source_question_id")
    )
    preview_results = [
        task.preview_class_result
        for task in class_tasks
        if task.preview_class_result_id
    ]
    preview_runs = {
        result.preview_run_id: result.preview_run
        for result in preview_results
    }

    def export_time(value):
        return value.isoformat() if value else ""

    rows = {
        "subjects": (identity_headers, subject_values),
        "sessions": ((*identity_headers, "session_id", "experiment_id", "version", "class_task_id", "participant_snapshot_id", "schedule_revision_id", "status", "started_at", "completed_at", "task_opens_at", "task_closes_at", "total_time_limit_seconds", "initial_effective_deadline_at", "timer_semantics_version", "terminal_reason", "authoritative_terminal_at", "terminal_recorded_at"), [[*identity_for_session(session.pk), session.pk, session.version.experiment_id, session.version.version_number, session.class_task_id or "", session.participant_snapshot_id or "", session.schedule_revision_id or "", session.status, session.started_at.isoformat(), session.completed_at.isoformat() if session.completed_at else "", export_time(session.task_opens_at), export_time(session.task_closes_at), session.total_time_limit_seconds or "", export_time(session.initial_effective_deadline_at), session.timer_semantics_version, session.terminal_reason, export_time(session.authoritative_terminal_at), export_time(session.terminal_recorded_at)] for session in sessions]),
        "assignments": ((*identity_headers, "question_code", "assignment_id", "session_id", "question_snapshot_id", "item_instance_id", "assigned_sequence_index"), [[*identity_for_session(assignment.session_id), assignment.snapshot.code, assignment.pk, assignment.session_id, assignment.snapshot_id, assignment.item_instance_id, assignment.assigned_sequence_index] for assignment in assignments]),
        "item_instances": (("question_code", "item_instance_id", "question_snapshot_id", "presentation_variant", "option_order", "parameter_values", "random_seed", "presentation_hash"), [[assignment.snapshot.code, assignment.item_instance_id, assignment.snapshot_id, assignment.item_instance.presentation_variant, json.dumps(assignment.item_instance.option_order), json.dumps(assignment.item_instance.parameter_values), assignment.item_instance.random_seed, assignment.item_instance.presentation_hash] for assignment in _unique_assignments(assignments, "item_instance_id")]),
        "attempts": ((*identity_headers, "question_code", "attempt_id", "session_id", "question_snapshot_id", "item_instance_id", "status", "first_opened_at", "last_opened_at", "visit_count", "draft_answer", "draft_saved_at", "final_answer", "submitted_at", "first_to_submit_duration_ms", "last_visit_to_submission_ms", "is_correct", "score", "first_open_sequence_index", "submission_sequence_index", "current_response_snapshot_id", "final_response_snapshot_id", "submission_event_id", "submission_grade_status", "grade_version", "grade_is_correct", "grade_score", "grade_max_score", "normalized_response", "grading_method", "grading_version", "grader_version", "graded_at", "item_started_at", "per_item_time_limit_seconds", "item_deadline_at", "initial_effective_deadline_at", "close_reason", "authoritative_terminal_at", "terminal_recorded_at"), [_attempt_export_row(assignment, attempt_by_assignment.get(assignment.pk), session_by_id, grade_by_attempt) for assignment in assignments if assignment.session_id in session_by_id]),
        "assessment_events": (event_headers, event_rows),
        "response_events": (response_headers, response_rows),
        "exposure_events": (projection_headers, exposure_rows),
        "navigation_events": (navigation_headers, navigation_rows),
        "telemetry_events": (projection_headers, telemetry_rows),
        "simulation_events": (compact_event_headers, simulation_event_rows),
        "simulation_snapshots": ((*identity_headers, "question_code", "snapshot_id", "attempt_id", "simulation_version", "state_metadata", "content_hash", "content_size", "created_at"), [[*identity_and_question_for_attempt(simulation.attempt_id), simulation.snapshot_id, simulation.attempt_id, simulation.simulation_version, json.dumps(simulation.state_metadata), simulation.content_hash, simulation.content_size, simulation.created_at.isoformat()] for simulation in selected_simulations]),
        "confidence_assignments": ((*identity_headers, "question_code", "confidence_assignment_id", "attempt_id", "condition", "assignment_unit", "assignment_version", "response_status", "probability", "assigned_at", "committed_at"), [[*identity_and_question_for_attempt(confidence.attempt_id), confidence.confidence_assignment_id, confidence.attempt_id, confidence.condition, confidence.assignment_unit, confidence.assignment_version, confidence.response_status, confidence.probability, confidence.assigned_at.isoformat(), confidence.committed_at.isoformat() if confidence.committed_at else ""] for confidence in selected_confidence]),
        "item_snapshots": (("question_code", "question_snapshot_id", "version_id", "stable_code", "question_type", "primary_kc", "all_kcs"), [[assignment.snapshot.code, assignment.snapshot_id, assignment.snapshot.version_id, assignment.snapshot.stable_code, assignment.snapshot.question_type, assignment.snapshot.primary_kc, json.dumps(assignment.snapshot.all_kcs, ensure_ascii=False)] for assignment in _unique_assignments(assignments, "snapshot_id")]),
        "item_kc_mapping": (("question_code", "question_snapshot_id", "primary_kc", "all_kcs"), [[assignment.snapshot.code, assignment.snapshot_id, assignment.snapshot.primary_kc, json.dumps(assignment.snapshot.all_kcs, ensure_ascii=False)] for assignment in _unique_assignments(assignments, "snapshot_id")]),
        "q_matrix": (("question_code", "question_snapshot_id", "primary_kc", "all_kcs"), [[assignment.snapshot.code, assignment.snapshot_id, assignment.snapshot.primary_kc, json.dumps(assignment.snapshot.all_kcs, ensure_ascii=False)] for assignment in _unique_assignments(assignments, "snapshot_id")]),
        "kc_dictionary": (("kc_id", "kc_name"), [[assignment.snapshot.primary_kc, assignment.snapshot.primary_kc] for assignment in _unique_assignments(assignments, "snapshot_id") if assignment.snapshot.primary_kc]),
        "response_snapshots": (
            (*identity_headers, "question_code", "snapshot_id", "attempt_id", "content", "content_hash", "content_length", "state_json", "state_hash", "schema_version", "created_at"),
            [
                [
                    *identity_and_question_for_attempt(item.attempt_id),
                    item.snapshot_id,
                    item.attempt_id,
                    item.content,
                    item.content_hash,
                    item.content_length,
                    json.dumps(item.state_json, ensure_ascii=False),
                    item.state_hash,
                    item.schema_version,
                    item.created_at.isoformat(),
                ]
                for item in response_snapshots
                if item.attempt_id in attempt_by_id
            ],
        ),
        "response_input_observations": (
            (*identity_headers, "question_code", "observation_id", "source_event_id", "source_event_type", "event_version", "envelope_version", "attempt_id", "client_instance_id", "client_sequence_no", "client_monotonic_ms", "batch_start_client_monotonic_ms", "batch_end_client_monotonic_ms", "client_wall_time_utc", "client_time_origin", "time_sync_id", "estimated_occurred_at_utc", "derived_occurred_at_utc", "server_received_at", "server_committed_at", "server_ingest_sequence_no", "actor_source", "capture_mode", "quality_flags", "time_quality", "coverage_run_id", "app_version", "edit_operation", "content_hash_before", "content_hash_after", "content_length_before", "content_length_after", "input_method", "is_ime", "observed_at"),
            [
                input_observation_row(item)
                for item in input_observations
                if item.attempt_id in attempt_by_id
            ],
        ),
        "submissions": (
            (*identity_headers, "question_code", "submission_id", "attempt_id", "submission_token", "final_response_snapshot_id", "grading_status", "score", "is_correct", "client_clicked_at", "server_received_at", "committed_at"),
            [
                [
                    *identity_and_question_for_attempt(item.attempt_id),
                    item.pk,
                    item.attempt_id,
                    item.submission_token,
                    item.final_response_snapshot_id,
                    item.grading_status,
                    item.score,
                    item.is_correct,
                    item.client_clicked_at.isoformat() if item.client_clicked_at else "",
                    item.server_received_at.isoformat(),
                    item.committed_at.isoformat(),
                ]
                for item in submissions
                if item.attempt_id in attempt_by_id
            ],
        ),
        "grade_records": (
            (*identity_headers, "question_code", "grade_id", "submission_id", "attempt_id", "version", "status", "validity_status", "score", "max_score", "is_correct", "grading_method", "grading_version", "graded_at", "supersedes_grade_id", "voided_at", "void_reason"),
            [
                [
                    *identity_and_question_for_attempt(item.submission.attempt_id),
                    item.pk,
                    item.submission_id,
                    item.submission.attempt_id,
                    item.version,
                    item.status,
                    item.validity_status,
                    item.score,
                    item.max_score,
                    item.is_correct,
                    item.grading_method,
                    item.grading_version,
                    item.graded_at.isoformat() if item.graded_at else "",
                    item.supersedes_grade_id or "",
                    item.voided_at.isoformat() if item.voided_at else "",
                    item.void_reason,
                ]
                for item in grade_records
                if item.submission.attempt_id in attempt_by_id
            ],
        ),
        "condition_assignments": (
            ("condition_assignment_id", "condition_id", "condition_key", "condition_label", "subject_id", "session_id", "assignment_unit_type", "assignment_algorithm", "assignment_version", "randomization_seed", "randomization_draw", "stratum", "assignment_probability", "assigned_at"),
            [
                [
                    item.pk,
                    item.condition_id,
                    item.condition.key,
                    item.condition.label,
                    item.subject_id or "",
                    item.session_id or "",
                    "subject" if item.subject_id else "session",
                    item.algorithm,
                    item.algorithm_version,
                    item.seed,
                    item.draw,
                    item.stratum,
                    item.probability,
                    item.assigned_at.isoformat(),
                ]
                for item in condition_assignments
            ],
        ),
        "class_snapshots": (
            ("class_snapshot_id", "version_id", "source_class_group_id", "name_snapshot", "academic_year_snapshot", "capture_mode", "created_at"),
            [
                [
                    item.class_snapshot_id,
                    item.version_id,
                    item.source_class_group_id or "",
                    item.name_snapshot,
                    item.academic_year_snapshot,
                    item.capture_mode,
                    item.created_at.isoformat(),
                ]
                for item in class_snapshots
            ],
        ),
        "participant_snapshots": (
            ("participant_snapshot_id", "version_id", "student_id", "class_snapshot_id", "created_at"),
            [
                [item.pk, item.version_id, item.student_id, item.class_snapshot_id or "", item.created_at.isoformat()]
                for item in participant_snapshots
            ],
        ),
        "capture_contracts": (
            ("capture_contract_id", "contract_name", "contract_version", "schema_document", "schema_sha256", "created_at"),
            [[item.pk, item.contract_name, item.contract_version, json.dumps(item.schema_document, ensure_ascii=False), item.schema_sha256, item.created_at.isoformat()] for item in capture_contracts],
        ),
        "client_time_syncs": (
            ("time_sync_id", "session_id", "client_instance_id", "client_wall_time_utc", "client_monotonic_ms", "client_time_origin", "server_received_at", "server_responded_at", "client_received_wall_time_utc", "client_received_monotonic_ms", "estimated_clock_offset_ms", "estimated_rtt_ms", "sync_quality", "finalized_at"),
            [[item.pk, item.session_id, item.client_instance_id, item.client_wall_time_utc.isoformat(), item.client_monotonic_ms, item.client_time_origin.isoformat(), item.server_received_at.isoformat(), item.server_responded_at.isoformat(), item.client_received_wall_time_utc.isoformat() if item.client_received_wall_time_utc else "", item.client_received_monotonic_ms or "", item.estimated_clock_offset_ms or "", item.estimated_rtt_ms or "", item.sync_quality, item.finalized_at.isoformat() if item.finalized_at else ""] for item in time_syncs],
        ),
        "client_capture_checkpoints": (
            ("checkpoint_id", "session_id", "student_id", "client_instance_id", "highest_emitted_sequence_no", "highest_acked_sequence_no", "highest_contiguous_acked_sequence_no", "pending_event_count", "failed_event_count", "missing_local_sequence_ranges", "client_monotonic_ms", "client_wall_time_utc", "client_time_origin", "time_sync_id", "trigger", "quality_flags", "app_version", "server_received_at"),
            [[item.pk, item.session_id, item.student_id, item.client_instance_id, item.highest_emitted_sequence_no, item.highest_acked_sequence_no, item.highest_contiguous_acked_sequence_no, item.pending_event_count, item.failed_event_count, json.dumps(item.missing_local_sequence_ranges), item.client_monotonic_ms, item.client_wall_time_utc.isoformat(), item.client_time_origin.isoformat(), item.time_sync_id or "", item.trigger, json.dumps(item.quality_flags), item.app_version, item.server_received_at.isoformat()] for item in checkpoints],
        ),
        "recording_coverage_sessions": (
            ("coverage_run_id", "coverage_version", "coverage_parameter_hash", "coverage_source_cutoff", "session_id", "overall_status", "lifecycle_status", "navigation_status", "client_sequence_status", "time_sync_status", "telemetry_status", "client_instance_count", "sequence_gap_count", "missing_sequence_ranges", "tail_gap_detected", "missing_reasons", "analysis_eligibility"),
            [[item.run_id, item.run.coverage_version, item.run.parameter_hash, item.run.source_cutoff_at.isoformat(), item.session_id, item.overall_status, item.lifecycle_status, item.navigation_status, item.client_sequence_status, item.time_sync_status, item.telemetry_status, item.client_instance_count, item.sequence_gap_count, json.dumps(item.missing_sequence_ranges), item.tail_gap_detected, json.dumps(item.missing_reasons), json.dumps(item.analysis_eligibility, ensure_ascii=False)] for item in latest_coverage_session.values()],
        ),
        "recording_coverage_attempts": (
            ("coverage_run_id", "coverage_version", "coverage_parameter_hash", "coverage_source_cutoff", "attempt_id", "overall_status", "enter_status", "render_status", "leave_status", "response_chain_status", "final_state_status", "time_quality", "response_state_count", "sequence_gap_count", "missing_reasons", "analysis_eligibility", "derived_features"),
            [[item.run_id, item.run.coverage_version, item.run.parameter_hash, item.run.source_cutoff_at.isoformat(), item.attempt_id, item.overall_status, item.enter_status, item.render_status, item.leave_status, item.response_chain_status, item.final_state_status, item.time_quality, item.response_state_count, item.sequence_gap_count, json.dumps(item.missing_reasons), json.dumps(item.analysis_eligibility, ensure_ascii=False), json.dumps(item.derived_features, ensure_ascii=False)] for item in latest_coverage_attempt.values()],
        ),
    }
    rows.update(
        {
            "release_metadata": (
                (
                    "release_snapshot_id", "experiment_id", "version_number",
                    "published_at", "timer_semantics_version",
                    "randomization_semantics_version", "definition_fingerprint",
                    "canonicalization_version",
                ),
                [
                    [
                        release.pk, release.experiment_id, release.version_number,
                        export_time(release.published_at), release.timer_semantics_version,
                        release.randomization_semantics_version,
                        release.definition_fingerprint, release.canonicalization_version,
                    ]
                    for release in releases
                ],
            ),
            "class_experiment_tasks": (
                (
                    "class_task_id", "experiment_id", "release_snapshot_id",
                    "class_group_id", "class_snapshot_id", "published_at",
                    "class_published_at", "randomization_algorithm_version",
                    "root_seed_fingerprint", "preview_class_result_id",
                    "cancelled_at", "manual_closed_at", "invalidated_at",
                    "terminal_reason", "provenance",
                ),
                [
                    [
                        task.pk, task.experiment_id, task.release_snapshot_id,
                        task.class_group_id, task.class_snapshot_id,
                        export_time(task.published_at), export_time(task.class_published_at),
                        task.randomization_identity.algorithm_version
                        if task.randomization_identity_id else "",
                        task.randomization_identity.root_seed_fingerprint
                        if task.randomization_identity_id else "",
                        task.preview_class_result_id or "", export_time(task.cancelled_at),
                        export_time(task.manual_closed_at), export_time(task.invalidated_at),
                        task.terminal_reason or "", task.provenance or "",
                    ]
                    for task in class_tasks
                ],
            ),
            "schedule_revisions": (
                (
                    "class_task_id", "revision_no", "opens_at", "closes_at",
                    "timezone", "reason", "provenance", "created_at",
                ),
                [
                    [
                        revision.class_task_id, revision.revision_no,
                        export_time(revision.opens_at), export_time(revision.closes_at),
                        revision.timezone_name, revision.reason, revision.provenance,
                        export_time(revision.created_at),
                    ]
                    for revision in schedule_revisions
                ],
            ),
            "release_question_mappings": (
                (
                    "release_snapshot_id", "source_question_id",
                    "question_release_fingerprint", "question_snapshot_id",
                    "created_at",
                ),
                [
                    [
                        mapping.release_snapshot_id, mapping.source_question_id,
                        mapping.question_release_fingerprint,
                        mapping.question_snapshot_id, export_time(mapping.created_at),
                    ]
                    for mapping in release_mappings
                ],
            ),
            "publish_preview_runs": (
                (
                    "preview_run_id", "plan_revision", "status",
                    "definition_fingerprint", "input_hash", "output_hash",
                    "algorithm_version", "created_at", "invalidated_at",
                    "invalidation_reason",
                ),
                [
                    [
                        run.pk, run.plan_revision, run.status,
                        run.definition_fingerprint, run.input_hash, run.output_hash,
                        run.algorithm_version, export_time(run.created_at),
                        export_time(run.invalidated_at), run.invalidation_reason,
                    ]
                    for run in preview_runs.values()
                ],
            ),
            "publish_preview_class_results": (
                (
                    "preview_class_result_id", "preview_run_id",
                    "class_task_id", "participant_count", "quota_valid",
                    "input_hash", "output_hash", "algorithm_version",
                    "root_seed_fingerprint", "created_at",
                ),
                [
                    [
                        result.pk, result.preview_run_id,
                        next(
                            task.pk for task in class_tasks
                            if task.preview_class_result_id == result.pk
                        ),
                        result.participant_count, result.quota_valid,
                        result.input_hash, result.output_hash,
                        result.randomization_identity.algorithm_version,
                        result.randomization_identity.root_seed_fingerprint,
                        export_time(result.created_at),
                    ]
                    for result in preview_results
                ],
            ),
        }
    )
    derivation_id = str(getattr(snapshot, "snapshot_id", "preview"))
    base = (*identity_headers, "question_code", "attempt_id", "derivation_run_id", "derivation_version")
    feature_rows, kt_rows, cd_rows, cal_rows = [], [], [], []
    for attempt in attempts:
        ident = [*identity_and_question_for_attempt(attempt.pk), attempt.pk, derivation_id, "v1"]
        feature_rows.append([*ident, attempt.is_correct, attempt.score, attempt.visit_count, attempt.first_opened_at.isoformat() if attempt.first_opened_at else "", attempt.submitted_at.isoformat() if attempt.submitted_at else ""])
        outcome = "" if attempt.is_correct is None else int(bool(attempt.is_correct))
        kt_rows.append([*ident, attempt.assignment.assigned_sequence_index, outcome, attempt.assignment.snapshot.primary_kc or ""])
        response_value = attempt.final_answer if attempt.final_answer is not None else attempt.draft_answer
        cd_rows.append([*ident, json.dumps(response_value, ensure_ascii=False) if response_value is not None else "", json.dumps(attempt.assignment.snapshot.all_kcs, ensure_ascii=False)])
        confidence = next((c for c in selected_confidence if c.attempt_id == attempt.pk), None)
        probability = confidence.probability if confidence and confidence.probability is not None else ""
        error = (float(probability) / 100.0 - float(bool(attempt.is_correct))) if probability != "" and attempt.is_correct is not None else ""
        cal_rows.append([*ident, probability, attempt.is_correct, error, confidence.condition if confidence else ""])
    rows.update({
        "attempt_features": (base + ("is_correct", "score", "visit_count", "first_opened_at", "submitted_at"), feature_rows),
        "kt_sequences": (base + ("sequence_index", "outcome", "primary_kc"), kt_rows),
        "cd_response_matrix": (base + ("response", "all_kcs"), cd_rows),
        "calibration_features": (base + ("confidence_probability", "is_correct", "calibration_error", "condition"), cal_rows),
    })
    return rows


def _unique_assignments(assignments, field):
    seen = set()
    result = []
    for assignment in assignments:
        value = getattr(assignment, field)
        if value not in seen:
            result.append(assignment)
            seen.add(value)
    return result


def _attempt_export_row(assignment, attempt, session_by_id, grade_by_attempt=None):
    session = session_by_id[assignment.session_id]
    identity = [session.student.student_number or "", session.student.shown_name, assignment.snapshot.code]
    linkage = [assignment.session_id, assignment.snapshot.pk, assignment.item_instance_id]
    if attempt is None:
        return [*identity, "", *linkage] + [""] * 28 + [""] * 7
    grade = (grade_by_attempt or {}).get(attempt.pk)
    submission = getattr(attempt, "submission_event", None)
    if attempt is None:
        return [*identity, "", *linkage, *([""] * 14)]
    return [
        *identity,
        attempt.pk,
        *linkage,
        attempt.status,
        attempt.first_opened_at.isoformat() if attempt.first_opened_at else "",
        attempt.last_opened_at.isoformat() if attempt.last_opened_at else "",
        attempt.visit_count,
        json.dumps(attempt.draft_answer, ensure_ascii=False) if attempt.draft_answer is not None else "",
        attempt.draft_saved_at.isoformat() if attempt.draft_saved_at else "",
        json.dumps(attempt.final_answer, ensure_ascii=False) if attempt.final_answer is not None else "",
        attempt.submitted_at.isoformat() if attempt.submitted_at else "",
        attempt.first_to_submit_duration_ms,
        attempt.last_visit_to_submission_ms,
        attempt.is_correct,
        attempt.score,
        attempt.first_open_sequence_index,
        attempt.submission_sequence_index,
        attempt.current_response_snapshot_id,
        attempt.final_response_snapshot_id,
        submission.pk if submission else "",
        submission.grading_status if submission else "",
        grade.version if grade else "",
        grade.is_correct if grade else "",
        grade.score if grade else "",
        grade.max_score if grade else "",
        json.dumps(grade.normalized_response, ensure_ascii=False) if grade and grade.normalized_response is not None else "",
        grade.grading_method if grade else "",
        grade.grading_version if grade else "",
        grade.grader_version if grade else "",
        grade.graded_at.isoformat() if grade and grade.graded_at else "",
        attempt.item_started_at.isoformat() if attempt.item_started_at else "",
        attempt.per_item_time_limit_seconds or "",
        attempt.item_deadline_at.isoformat() if attempt.item_deadline_at else "",
        attempt.initial_effective_deadline_at.isoformat()
        if attempt.initial_effective_deadline_at else "",
        attempt.close_reason,
        attempt.authoritative_terminal_at.isoformat()
        if attempt.authoritative_terminal_at else "",
        attempt.terminal_recorded_at.isoformat()
        if attempt.terminal_recorded_at else "",
    ]


def _finalize_export_job(job, zip_path, manifest):
    digest = hashlib.sha256(open(zip_path, "rb").read()).hexdigest()
    job.status = ExportJob.Status.SUCCEEDED
    job.finished_at = timezone.now()
    job.progress_stage = "complete"
    job.artifact_path = zip_path
    job.artifact_sha256 = digest
    job.save(update_fields=["status", "finished_at", "progress_stage", "artifact_path", "artifact_sha256"])
    return job


def execute_purge_request(purge_request, *, executed_by):
    from apps.accounts.models import DataPurgeRequest
    from apps.assessment.models import (
        AssessmentEvent, Attempt, ClientTimeSync, ConfidenceAssignment,
        DraftSaveEvent, ExperimentSession, QuestionVisitEvent, ResponseInputObservation, ResponseSnapshot,
        SimulationSnapshot, StudentQuestionAssignment, SubmissionEvent, GradeRecord,
    )
    course_id = purge_request.scope_definition.get("course_id")
    if not course_id:
        raise ValueError("purge scope missing course_id")
    with transaction.atomic():
        affected_snapshots = list(
            ExportSnapshot.objects.filter(course_id=course_id)
        )
        session_ids = list(ExperimentSession.objects.filter(version__experiment__course_id=course_id).values_list("pk", flat=True))
        attempt_ids = list(Attempt.objects.filter(session_id__in=session_ids).values_list("pk", flat=True))
        counts = {"sessions": len(session_ids), "attempts": len(attempt_ids), "events": AssessmentEvent.objects.filter(session_id__in=session_ids).count()}
        # Clear durable pointers before removing immutable snapshots that they
        # intentionally protect during ordinary operation.
        Attempt.objects.filter(pk__in=attempt_ids).update(
            current_response_snapshot_id=None,
            final_response_snapshot_id=None,
        )
        for model, field, ids in (
            (AssessmentEvent, "session_id", session_ids),
            (ConfidenceAssignment, "attempt_id", attempt_ids),
            (SimulationSnapshot, "attempt_id", attempt_ids),
            (ResponseInputObservation, "attempt_id", attempt_ids),
            (DraftSaveEvent, "attempt_id", attempt_ids),
            (QuestionVisitEvent, "session_id", session_ids),
            (GradeRecord, "submission__attempt_id", attempt_ids),
            (SubmissionEvent, "attempt_id", attempt_ids),
            # The durable snapshot pointers must be nulled before their
            # referenced immutable records can be physically erased.
            (ResponseSnapshot, "attempt_id", attempt_ids),
            (Attempt, "pk", attempt_ids),
            (StudentQuestionAssignment, "session_id", session_ids),
            (ClientTimeSync, "session_id", session_ids),
            (ExperimentSession, "pk", session_ids),
        ):
            if ids:
                counts[model.__name__.lower()] = model.objects.filter(**{f"{field}__in": ids}).count() if field != "pk" else len(ids)
                model.objects.filter(**{f"{field}__in": ids})._raw_delete(using="default")
        affected_jobs = list(ExportJob.objects.filter(course_id=course_id))
        invalidated_at = timezone.now()
        for snapshot in affected_snapshots:
            ExportSnapshot.objects.filter(pk=snapshot.pk).update(
                invalidated_at=invalidated_at,
                invalidated_reason=f"purge:{purge_request.reason_code}",
                filter_definition={"scope_redacted": True},
                scope_hash="",
                redacted_at=invalidated_at,
            )
        for job in affected_jobs:
            remove_export_artifact(
                job,
                reason=f"purge:{purge_request.reason_code}",
                invalidated=True,
            )
        purge_request.status = DataPurgeRequest.Status.EXECUTED
        purge_request.executed_by = executed_by
        purge_request.executed_at = timezone.now()
        purge_request.result_counts = counts
        purge_request.save(update_fields=["status", "executed_by", "executed_at", "result_counts"])
    return counts


def cleanup_expired_exports(*, now=None):
    now = now or timezone.now()
    removed = 0
    for job in ExportJob.objects.filter(expires_at__lte=now).exclude(status=ExportJob.Status.EXPIRED):
        if job.artifact_path and os.path.exists(job.artifact_path):
            removed += 1
        remove_export_artifact(job, reason="retention_expired")
    return removed
