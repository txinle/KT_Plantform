import os
import tempfile
import zipfile
from unittest.mock import patch

from openpyxl import load_workbook

from django.test import TransactionTestCase

from apps.accounts.models import Course, User
from apps.assessment.models import Attempt, ExperimentSession, StudentQuestionAssignment
from apps.experiments.models import Experiment, ExperimentVersion, QuestionSnapshot
from apps.experiments.presentation_instances import get_or_create_default_item_instance
from apps.question_bank.models import Question, QuestionBank

from .data_scope import DataScope
from .models import ExportJob
from .services import (
    ExportContractError,
    _write_teaching_workbook,
    create_export_job,
    create_export_snapshot,
    run_export_job,
)


class ExportScopeV2Tests(TransactionTestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(
            username="export-v2-teacher", role=User.Role.TEACHER
        )
        self.student = User.objects.create_user(
            username="export-v2-student", role=User.Role.STUDENT, student_number="EXV2-1"
        )
        self.course = Course.objects.create(
            owner=self.teacher, name="导出范围课程", code="EXPORTV2", term="2026 秋"
        )
        bank = QuestionBank.objects.create(
            course=self.course, created_by=self.teacher, name="导出范围题库"
        )
        self.experiment = Experiment.objects.create(
            course=self.course,
            created_by=self.teacher,
            question_bank=bank,
            name="导出范围实验",
        )
        self.version = ExperimentVersion.objects.create(
            experiment=self.experiment, version_number=1, published_by=self.teacher
        )
        question = Question.objects.create(
            bank=bank,
            sequence=1,
            question_type=Question.Type.SHORT,
            text="导出题",
            correct_answer="答案",
        )
        snapshot = QuestionSnapshot.objects.create(
            version=self.version,
            source_question=question,
            source_sequence=1,
            code="EXPORT-Q1",
            question_type=Question.Type.SHORT,
            text="导出题",
            correct_answer="答案",
            primary_kc="导出 KC",
            all_kcs=["导出 KC"],
        )
        session = ExperimentSession.objects.create(version=self.version, student=self.student)
        assignment = StudentQuestionAssignment.objects.create(
            session=session,
            snapshot=snapshot,
            item_instance=get_or_create_default_item_instance(snapshot),
            assigned_sequence_index=1,
        )
        Attempt.objects.create(session=session, assignment=assignment)

    def _single_version_scope(self):
        return DataScope.from_mapping(
            self.course,
            {
                "experiment_id": self.experiment.pk,
                "version_mode": "single",
                "version_id": self.version.pk,
            },
        )

    def test_every_request_creates_a_new_snapshot_and_job_with_canonical_scope(self):
        scope = self._single_version_scope()
        first_snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
            scope=scope,
        )
        first_job = create_export_job(
            course=self.course,
            experiment=self.experiment,
            snapshot=first_snapshot,
            created_by=self.teacher,
            package="process",
        )
        second_snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
            scope=scope,
        )
        second_job = create_export_job(
            course=self.course,
            experiment=self.experiment,
            snapshot=second_snapshot,
            created_by=self.teacher,
            package="process",
        )

        self.assertNotEqual(first_snapshot.pk, second_snapshot.pk)
        self.assertNotEqual(first_job.pk, second_job.pk)
        self.assertEqual(first_snapshot.filter_definition, scope.serialize())

    def test_process_package_includes_raw_response_grade_condition_and_class_contracts(self):
        scope = self._single_version_scope()
        export_snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
            scope=scope,
        )
        job = create_export_job(
            course=self.course,
            experiment=self.experiment,
            snapshot=export_snapshot,
            created_by=self.teacher,
            package="process",
        )

        run_export_job(job)

        with zipfile.ZipFile(job.artifact_path) as archive:
            names = set(archive.namelist())
        self.assertTrue(
            {
                "03_raw_events/response_snapshots.csv",
                "03_raw_events/response_input_observations.csv",
                "02_core/submissions.csv",
                "02_core/grade_records.csv",
                "01_design_metadata/condition_assignments.csv",
                "01_design_metadata/class_snapshots.csv",
            }.issubset(names)
        )

    def test_kt_and_cd_reject_all_version_scope(self):
        all_versions = DataScope.from_mapping(
            self.course,
            {"experiment_id": self.experiment.pk, "version_mode": "all"},
        )

        with self.assertRaises(ExportContractError):
            create_export_snapshot(
                course=self.course,
                experiment=self.experiment,
                created_by=self.teacher,
                scope=all_versions,
                package="kt",
            )

    def test_artifact_removal_keeps_an_expired_audit_tombstone(self):
        scope = self._single_version_scope()
        export_snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
            scope=scope,
        )
        job = create_export_job(
            course=self.course,
            experiment=self.experiment,
            snapshot=export_snapshot,
            created_by=self.teacher,
            package="process",
        )
        run_export_job(job)
        artifact_path = job.artifact_path

        job.expire_artifact(reason="user_requested_removal")

        job.refresh_from_db()
        self.assertEqual(job.status, ExportJob.Status.EXPIRED)
        self.assertEqual(job.artifact_path, "")
        self.assertTrue(job.artifact_sha256)
        self.assertIsNotNone(job.artifact_removed_at)
        self.assertFalse(os.path.exists(artifact_path))

    def test_teaching_package_is_a_single_version_xlsx_with_lineage(self):
        scope = self._single_version_scope()
        export_snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
            scope=scope,
            package="teaching",
        )
        job = create_export_job(
            course=self.course,
            experiment=self.experiment,
            snapshot=export_snapshot,
            created_by=self.teacher,
            package="teaching",
        )

        run_export_job(job)

        self.assertTrue(job.artifact_path.endswith(".xlsx"))
        workbook = load_workbook(job.artifact_path, read_only=True)
        self.assertEqual(workbook.sheetnames, ["教学成绩", "导出血缘"])
        self.assertIn("grade_score", [cell.value for cell in next(workbook["教学成绩"].iter_rows())])

    def test_teaching_workbook_uses_only_records_frozen_by_export_snapshot(self):
        attempt = Attempt.objects.get()
        headers = (
            "attempt_id",
            "student_number",
            "student_name",
            "question_code",
            "session_id",
            "status",
            "grade_score",
            "grade_max_score",
            "grade_is_correct",
            "submission_grade_status",
            "graded_at",
        )
        rows = [[attempt.pk, "EXV2-1", "学生", "EXPORT-Q1", attempt.session_id, "unvisited", None, None, None, "pending", None]]
        manifest = {
            "export_snapshot_id": "snapshot",
            "source_cutoff_at": "2026-08-28T00:00:00+00:00",
            "event_high_watermark": 0,
            "schema_manifest_hash": "hash",
            "scope_hash": "scope",
            "contains_direct_identifiers": True,
        }
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "teaching.xlsx")
            with patch(
                "apps.analytics.services.AssessmentEvent.objects.filter",
                side_effect=AssertionError("must not query live events"),
            ), patch(
                "apps.analytics.services.SubmissionEvent.objects.filter",
                side_effect=AssertionError("must not query live submissions"),
            ):
                _write_teaching_workbook(
                    path,
                    (headers, rows),
                    manifest,
                    [attempt],
                    events=[],
                    submissions=[],
                )
            self.assertTrue(os.path.exists(path))
