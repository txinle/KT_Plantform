import base64
import csv
from decimal import Decimal
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import uuid
import zipfile
from datetime import timedelta, timezone as dt_timezone
from io import StringIO
from unittest.mock import patch

from django.conf import settings
from django.db import transaction
from django.test import TransactionTestCase
from django.utils import timezone

from apps.accounts.models import (
    ClassGroup, Course, CourseEnrollment, DataPurgeRequest, User,
)
from apps.analytics.models import ExportJob, ExportSnapshot
from apps.analytics.services import (
    ExportContractError,
    create_export_snapshot,
    execute_purge_request,
    run_export_job,
)
from apps.analytics.recording_coverage import generate_session_coverage
from apps.assessment.models import (
    AssessmentEvent,
    Attempt,
    ClientTimeSync,
    ExperimentSession,
    GradeRecord,
    StudentQuestionAssignment,
    SubmissionEvent,
)
from apps.experiments.models import (
    Experiment, ExperimentStudent, ExperimentVersion, QuestionSnapshot,
)
from apps.experiments.presentation_instances import get_or_create_default_item_instance
from apps.question_bank.models import QuestionBank, Question


class ExportLineageV2Tests(TransactionTestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(username="lineage-teacher", role=User.Role.TEACHER)
        self.course = Course.objects.create(owner=self.teacher, name="C", code="C", term="T")
        bank = QuestionBank.objects.create(owner=self.teacher, course=self.course, name="B")
        self.experiment = Experiment.objects.create(course=self.course, created_by=self.teacher, question_bank=bank, name="E")
        self.version = ExperimentVersion.objects.create(experiment=self.experiment, version_number=1, published_by=self.teacher)
        q = Question.objects.create(bank=bank, sequence=1, question_type=Question.Type.SHORT, text="Q", correct_answer="A")
        self.qs = QuestionSnapshot.objects.create(version=self.version, source_question=q, source_sequence=1, code="Q1", question_type=Question.Type.SHORT, text="Q", correct_answer="A", primary_kc="KC", all_kcs=["KC"])

    def _session(self, number, *, status=Attempt.Status.SUBMITTED, draft_answer=None,
                 capture_contract=None):
        student = User.objects.create_user(username=number, role=User.Role.STUDENT, student_number=number, display_name=number)
        session = ExperimentSession.objects.create(
            version=self.version,
            student=student,
            capture_contract=capture_contract,
        )
        assignment = StudentQuestionAssignment.objects.create(session=session, snapshot=self.qs, item_instance=get_or_create_default_item_instance(self.qs), assigned_sequence_index=1)
        return Attempt.objects.create(
            session=session,
            assignment=assignment,
            status=status,
            draft_answer=draft_answer,
            final_answer="A" if status == Attempt.Status.SUBMITTED else None,
            is_correct=True if status == Attempt.Status.SUBMITTED else None,
            score=1 if status == Attempt.Status.SUBMITTED else None,
        )

    def test_snapshot_freezes_event_boundary_fields(self):
        with patch("apps.analytics.services.timezone.now", return_value=timezone.datetime(2026, 1, 1, tzinfo=dt_timezone.utc)):
            snapshot = create_export_snapshot(course=self.course, experiment=self.experiment, created_by=self.teacher)
        self.assertEqual(snapshot.source_cutoff_at, snapshot.cutoff_at)
        self.assertIsNotNone(snapshot.event_high_watermark)
        self.assertTrue(snapshot.schema_manifest_hash)
        snapshot.event_high_watermark += 1
        with self.assertRaises(ValueError):
            snapshot.save(update_fields=["event_high_watermark"])

    def test_outer_read_committed_transaction_is_rejected_for_source_freeze(self):
        with transaction.atomic():
            with self.assertRaisesRegex(ExportContractError, "REPEATABLE READ"):
                create_export_snapshot(
                    course=self.course,
                    experiment=self.experiment,
                    created_by=self.teacher,
                )

    def test_export_contains_deterministic_derived_tables_and_provenance(self):
        self._session("S1")
        snapshot = create_export_snapshot(course=self.course, experiment=self.experiment, created_by=self.teacher)
        job = ExportJob.objects.create(course=self.course, experiment=self.experiment, snapshot=snapshot, created_by=self.teacher, package="complete")
        run_export_job(job)
        with zipfile.ZipFile(job.artifact_path) as archive:
            names = set(archive.namelist())
            self.assertIn("00_manifest/derivation_runs.json", names)
            manifest = json.loads(archive.read("00_manifest/manifest.json"))
            self.assertIn("provenance", manifest)
            for table in ("attempt_features", "kt_sequences", "cd_response_matrix", "calibration_features"):
                path = "06_derived/" + table + ".csv"
                rows = list(csv.reader(StringIO(archive.read(path).decode("utf-8-sig"))))
                self.assertEqual(rows[0][0:3], ["student_number", "student_name", "question_code"])
                self.assertNotEqual(manifest["data_status"][table], "placeholder")

    def test_manifest_freezes_actual_capture_contract_without_coverage(self):
        from apps.assessment.recording_capture_services import ensure_capture_contract_v2

        capture_contract = ensure_capture_contract_v2()
        self._session(
            "capture-v2-no-coverage",
            status=Attempt.Status.DRAFT,
            draft_answer="draft",
            capture_contract=capture_contract,
        )

        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )

        self.assertIsNone(snapshot.coverage_run_id)
        manifest = snapshot.source_bundle["manifest"]
        self.assertEqual(manifest["capture_contract_ids"], [str(capture_contract.pk)])
        self.assertEqual(
            manifest["capture_contract_hashes"],
            [capture_contract.schema_sha256],
        )

    def test_source_bundle_freezes_resources_and_table_paths(self):
        self._session("resource-freeze")
        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )

        resource_manifest = snapshot.source_bundle["resource_manifest"]
        frozen_resources = snapshot.source_bundle["resources"]
        frozen_dictionary = base64.b64decode(
            frozen_resources["00_manifest/field_dictionary.json"]
        )
        frozen_replay_script = base64.b64decode(
            frozen_resources["06_derived/replay_response_trajectory.py"]
        )
        frozen_attempt_path = resource_manifest["table_paths"]["attempts"]
        frozen_attempts = base64.b64decode(
            snapshot.source_bundle["tables"]["attempts"]
        )
        frozen_resource_hashes = dict(resource_manifest["resource_sha256"])
        self.assertEqual(
            frozen_resource_hashes["00_manifest/field_dictionary.json"],
            hashlib.sha256(frozen_dictionary).hexdigest(),
        )
        self.assertEqual(
            frozen_resource_hashes["06_derived/replay_response_trajectory.py"],
            hashlib.sha256(frozen_replay_script).hexdigest(),
        )

        job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        with (
            patch.dict(
                "apps.analytics.export_catalog.TABLE_FIELD_DESCRIPTIONS",
                {"attempts": {"attempt_id": "mutated catalog"}},
                clear=True,
            ),
            patch(
                "apps.analytics.replay_artifact.standalone_response_replay_script",
                return_value="mutated replay script",
            ),
            patch.dict(
                "apps.analytics.services.TABLE_PATHS",
                {"attempts": "99_mutated/attempts.csv"},
                clear=True,
            ),
        ):
            run_export_job(job)

        with zipfile.ZipFile(job.artifact_path) as archive:
            names = set(archive.namelist())
            manifest = json.loads(archive.read("00_manifest/manifest.json"))
            self.assertEqual(
                archive.read("00_manifest/field_dictionary.json"),
                frozen_dictionary,
            )
            self.assertEqual(
                archive.read("06_derived/replay_response_trajectory.py"),
                frozen_replay_script,
            )
            self.assertEqual(archive.read(frozen_attempt_path), frozen_attempts)

        self.assertNotIn("99_mutated/attempts.csv", names)
        self.assertEqual(
            manifest["resource_manifest"]["resource_sha256"],
            frozen_resource_hashes,
        )
        self.assertEqual(
            manifest["resource_manifest"]["resource_freeze_status"],
            "frozen",
        )
        field_dictionary = json.loads(frozen_dictionary)
        attempt_entry = next(
            entry for entry in field_dictionary["tables"]
            if entry["table"] == "attempts"
        )
        attempt_id = next(
            field for field in attempt_entry["columns"]
            if field["name"] == "attempt_id"
        )
        self.assertEqual(attempt_id["definition"], "作答记录唯一标识")
        self.assertEqual(attempt_id["type"], None)
        self.assertEqual(attempt_id["nullable"], None)

    def test_derivation_version_is_frozen_for_replay_and_legacy_bundles(self):
        self._session("derivation-v2")
        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        self.assertEqual(
            snapshot.source_bundle["manifest"]["provenance"]["derivation"]["version"],
            "v2",
        )
        source_feature_rows = list(csv.DictReader(StringIO(
            base64.b64decode(
                snapshot.source_bundle["tables"]["attempt_features"]
            ).decode("utf-8-sig")
        )))
        self.assertEqual(
            {row["derivation_version"] for row in source_feature_rows},
            {"v2"},
        )

        job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        with patch("apps.analytics.services.EXPORT_DERIVATION_VERSION", "v99"):
            run_export_job(job)
        with zipfile.ZipFile(job.artifact_path) as archive:
            manifest = json.loads(archive.read("00_manifest/manifest.json"))
            derivation_run = json.loads(
                archive.read("00_manifest/derivation_runs.json")
            )
        self.assertEqual(
            manifest["provenance"]["derivation"]["version"],
            "v2",
        )
        self.assertEqual(derivation_run["derivation_version"], "v2")

        legacy_bundle = json.loads(json.dumps(snapshot.source_bundle))
        legacy_bundle["manifest"]["provenance"]["derivation"]["version"] = "v1"
        legacy_bundle.pop("resources")
        legacy_bundle["resource_manifest"].pop("resource_sha256")
        legacy_bundle["resource_manifest"].pop("resource_freeze_status")
        legacy_bundle_hash = hashlib.sha256(
            json.dumps(
                legacy_bundle,
                ensure_ascii=False,
                separators=(",", ":"),
                sort_keys=True,
            ).encode("utf-8")
        ).hexdigest()
        ExportSnapshot.objects.filter(pk=snapshot.pk).update(
            source_bundle=legacy_bundle,
            source_bundle_hash=legacy_bundle_hash,
        )
        snapshot.refresh_from_db()
        legacy_job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        with patch("apps.analytics.services.EXPORT_DERIVATION_VERSION", "v100"):
            run_export_job(legacy_job)
        with zipfile.ZipFile(legacy_job.artifact_path) as archive:
            legacy_derivation_run = json.loads(
                archive.read("00_manifest/derivation_runs.json")
            )
            legacy_manifest = json.loads(archive.read("00_manifest/manifest.json"))
            legacy_names = set(archive.namelist())
        self.assertEqual(legacy_derivation_run["derivation_version"], "v1")
        self.assertEqual(
            legacy_manifest["resource_manifest"]["resource_freeze_status"],
            "legacy_unfrozen",
        )
        self.assertNotIn("00_manifest/field_dictionary.json", legacy_names)
        self.assertNotIn("06_derived/replay_response_trajectory.py", legacy_names)

    def test_grade_record_history_freezes_provenance_and_original_comment(self):
        from apps.analytics.export_catalog import (
            COMMON_FIELD_DESCRIPTIONS,
            TABLE_FIELD_DESCRIPTIONS,
        )

        attempt = self._session("grade-history")
        submission = SubmissionEvent.objects.create(
            attempt=attempt,
            submission_token=uuid.uuid4(),
            final_answer="A",
            server_received_at=timezone.now(),
            committed_at=timezone.now(),
        )
        voided_grade = GradeRecord.objects.create(
            submission=submission,
            version=1,
            status=GradeRecord.Status.COMPLETED,
            validity_status=GradeRecord.ValidityStatus.VOIDED,
            is_correct=True,
            score=1,
            max_score=1,
            grader_version="rubric-v1",
            grader=self.teacher,
            comment="请说明推理过程。",
            graded_at=timezone.now(),
            voided_at=timezone.now(),
            voided_by=self.teacher,
            void_reason="需要按修订量规重评。",
        )
        effective_grade = GradeRecord.objects.create(
            submission=submission,
            version=2,
            status=GradeRecord.Status.COMPLETED,
            validity_status=GradeRecord.ValidityStatus.EFFECTIVE,
            is_correct=True,
            score=1,
            max_score=1,
            grader_version="rubric-v2",
            grader=self.teacher,
            comment="按修订量规评分。",
            graded_at=timezone.now(),
            supersedes_grade=voided_grade,
        )

        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        run_export_job(job)

        with zipfile.ZipFile(job.artifact_path) as archive:
            rows = list(csv.DictReader(StringIO(
                archive.read("02_core/grade_records.csv").decode("utf-8-sig")
            )))

        rows_by_grade_id = {row["grade_id"]: row for row in rows}
        voided_row = rows_by_grade_id[str(voided_grade.pk)]
        effective_row = rows_by_grade_id[str(effective_grade.pk)]
        self.assertEqual(voided_row["grader_id"], str(self.teacher.pk))
        self.assertEqual(voided_row["comment"], "请说明推理过程。")
        self.assertEqual(voided_row["grader_version"], "rubric-v1")
        self.assertEqual(voided_row["voided_by_id"], str(self.teacher.pk))
        self.assertEqual(voided_row["created_at"], voided_grade.created_at.isoformat())
        self.assertEqual(effective_row["grader_version"], "rubric-v2")
        self.assertEqual(effective_row["voided_by_id"], "")
        self.assertEqual(
            effective_row["created_at"], effective_grade.created_at.isoformat(),
        )
        for field in (
            "grader_id", "grader_version", "voided_by_id", "created_at", "comment",
        ):
            self.assertIn(field, TABLE_FIELD_DESCRIPTIONS["grade_records"])
            self.assertIn(field, COMMON_FIELD_DESCRIPTIONS)

    def test_complete_export_freezes_coverage_run_provenance(self):
        from apps.assessment.recording_capture_services import ensure_capture_contract_v2

        attempt = self._session(
            "coverage-provenance",
            status=Attempt.Status.DRAFT,
            draft_answer="draft",
            capture_contract=ensure_capture_contract_v2(),
        )
        frozen_run = generate_session_coverage(attempt.session, force=True)
        self.assertEqual(frozen_run.coverage_version, "coverage-v3")
        self.assertEqual(frozen_run.algorithm_version, "recording-coverage-v3")
        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        later_run = generate_session_coverage(attempt.session, force=True)
        self.assertNotEqual(later_run.pk, frozen_run.pk)
        job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        run_export_job(job)

        with zipfile.ZipFile(job.artifact_path) as archive:
            coverage_csv = archive.read(
                "07_quality/recording_coverage_runs.csv"
            )
            coverage_rows = list(csv.DictReader(StringIO(
                coverage_csv.decode("utf-8-sig")
            )))
            field_dictionary = json.loads(archive.read(
                "00_manifest/field_dictionary.json"
            ))
            manifest = json.loads(archive.read("00_manifest/manifest.json"))

        self.assertEqual([row["coverage_run_id"] for row in coverage_rows], [
            str(frozen_run.pk),
        ])
        row = coverage_rows[0]
        self.assertEqual(row["scope_type"], frozen_run.scope_type)
        self.assertEqual(json.loads(row["scope_definition"]), frozen_run.scope_definition)
        self.assertEqual(row["scope_hash"], frozen_run.scope_hash)
        self.assertEqual(row["coverage_version"], frozen_run.coverage_version)
        self.assertEqual(row["algorithm_version"], frozen_run.algorithm_version)
        self.assertEqual(row["code_version"], frozen_run.code_version)
        self.assertEqual(
            json.loads(row["parameter_document"]), frozen_run.parameter_document,
        )
        self.assertEqual(row["parameter_hash"], frozen_run.parameter_hash)
        self.assertEqual(
            row["source_cutoff_at"], frozen_run.source_cutoff_at.isoformat(),
        )
        self.assertEqual(
            json.loads(row["source_high_water_marks"]),
            frozen_run.source_high_water_marks,
        )
        self.assertEqual(
            json.loads(row["capture_contract_hashes"]),
            frozen_run.capture_contract_hashes,
        )
        self.assertEqual(row["source_hash"], frozen_run.source_hash)
        self.assertEqual(row["run_status"], frozen_run.run_status)
        self.assertEqual(row["generated_at"], frozen_run.generated_at.isoformat())
        coverage_run_table = next(
            table for table in field_dictionary["tables"]
            if table["table"] == "recording_coverage_runs"
        )
        self.assertEqual(
            coverage_run_table["path"],
            "07_quality/recording_coverage_runs.csv",
        )
        self.assertIn("recording_coverage_runs", manifest["selected_tables"])

        for package, selected_tables in (
            ("process", []),
            ("complete", ["recording_coverage_runs"]),
        ):
            with self.subTest(package=package, selected_tables=selected_tables):
                scoped_job = ExportJob.objects.create(
                    course=self.course,
                    experiment=self.experiment,
                    snapshot=snapshot,
                    created_by=self.teacher,
                    package=package,
                    selected_tables=selected_tables,
                )
                run_export_job(scoped_job)
                with zipfile.ZipFile(scoped_job.artifact_path) as archive:
                    self.assertEqual(
                        archive.read("07_quality/recording_coverage_runs.csv"),
                        coverage_csv,
                    )

    def test_all_advertised_export_artifacts_materialize_frozen_catalog_members(self):
        from openpyxl import load_workbook

        from apps.analytics.export_catalog import PACKAGE_TABLES, TABLE_PATHS
        from apps.assessment.models import ConfidenceAssignment
        from apps.assessment.recording_capture_services import ensure_capture_contract_v2
        from apps.assessment.snapshot_services import create_simulation_snapshot

        attempt = self._session(
            "all-package-audit",
            capture_contract=ensure_capture_contract_v2(),
        )
        audit_class = ClassGroup.objects.create(course=self.course, name="Audit")
        CourseEnrollment.objects.create(
            course=self.course,
            class_group=audit_class,
            student=attempt.session.student,
        )
        attempt.session.student.class_group = audit_class
        attempt.session.student.save(update_fields=["class_group"])
        ExperimentStudent.objects.create(
            experiment=self.experiment,
            student=attempt.session.student,
            class_snapshot_name=str(audit_class),
        )
        submitted_at = timezone.now()
        submission = SubmissionEvent.objects.create(
            attempt=attempt,
            submission_token=uuid.uuid4(),
            final_answer="A",
            is_correct=True,
            score=1,
            server_received_at=submitted_at,
            committed_at=submitted_at,
        )
        GradeRecord.objects.create(
            submission=submission,
            version=1,
            status=GradeRecord.Status.COMPLETED,
            validity_status=GradeRecord.ValidityStatus.EFFECTIVE,
            is_correct=True,
            score=1,
            max_score=1,
            grading_method="automatic",
            grader_version="audit-rubric-v1",
            graded_at=submitted_at,
        )
        ConfidenceAssignment.objects.create(
            attempt=attempt,
            condition=ConfidenceAssignment.Condition.PROMPT,
            assignment_unit=ConfidenceAssignment.AssignmentUnit.ATTEMPT,
            assignment_version="audit-confidence-v1",
            response_status=ConfidenceAssignment.ResponseStatus.COMMITTED,
            probability=0,
            committed_at=submitted_at,
        )
        create_simulation_snapshot(
            attempt=attempt,
            simulation_version="audit-simulation-v1",
            state_metadata={"mode": "audit"},
            blob_storage_key="audit/simulation.bin",
            content_bytes=b"simulation-audit-content",
        )
        frozen_coverage_run = generate_session_coverage(attempt.session, force=True)
        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        self.assertEqual(frozen_coverage_run.run_status, "completed")

        package_attempt_bytes = {}
        package_jobs = {}
        zip_packages = (
            "complete", "kt", "cd", "process", "confidence", "simulation",
            "longitudinal",
        )
        for package in zip_packages:
            with self.subTest(package=package):
                job = ExportJob.objects.create(
                    course=self.course,
                    experiment=self.experiment,
                    snapshot=snapshot,
                    created_by=self.teacher,
                    package=package,
                )
                run_export_job(job)
                package_jobs[package] = job
                self.assertTrue(job.artifact_path.endswith(".zip"))
                with open(job.artifact_path, "rb") as artifact:
                    self.assertEqual(
                        hashlib.sha256(artifact.read()).hexdigest(),
                        job.artifact_sha256,
                    )
                with zipfile.ZipFile(job.artifact_path) as archive:
                    self.assertIsNone(archive.testzip())
                    manifest = json.loads(archive.read("00_manifest/manifest.json"))
                    self.assertEqual(
                        manifest["selected_tables"], list(PACKAGE_TABLES[package]),
                    )
                    members = set(archive.namelist())
                    self.assertTrue({
                        "00_manifest/manifest.json",
                        "00_manifest/derivation_runs.json",
                        "00_manifest/field_dictionary.json",
                        "01_design_metadata/README.txt",
                    }.issubset(members))
                    for table in PACKAGE_TABLES[package]:
                        path = TABLE_PATHS[table]
                        self.assertIn(path, members)
                        body = archive.read(path)
                        rows = list(csv.reader(StringIO(body.decode("utf-8-sig"))))
                        self.assertTrue(rows[0], table)
                        self.assertEqual(
                            len(rows) - 1,
                            manifest["row_counts"][table],
                            table,
                        )
                        self.assertEqual(
                            hashlib.sha256(body).hexdigest(),
                            snapshot.source_bundle["resource_manifest"]["table_sha256"][table],
                            table,
                        )
                    if package == "complete":
                        self.assertIn(
                            "06_derived/replay_response_trajectory.py", members,
                        )
                    package_attempt_bytes[package] = archive.read(
                        "02_core/attempts.csv"
                    )

        self.assertEqual(
            package_attempt_bytes["complete"],
            package_attempt_bytes["kt"],
        )
        with zipfile.ZipFile(package_jobs["confidence"].artifact_path) as archive:
            confidence_rows = list(csv.DictReader(StringIO(
                archive.read("05_self_report/confidence_assignments.csv").decode("utf-8-sig")
            )))
        self.assertEqual(confidence_rows[0]["probability"], "0")

        custom_job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
            selected_tables=["attempts", "recording_coverage_runs"],
        )
        run_export_job(custom_job)
        with zipfile.ZipFile(custom_job.artifact_path) as archive:
            custom_manifest = json.loads(archive.read("00_manifest/manifest.json"))
            self.assertEqual(
                custom_manifest["selected_tables"],
                ["attempts", "recording_coverage_runs"],
            )
            self.assertEqual(
                {
                    "02_core/attempts.csv",
                    "07_quality/recording_coverage_runs.csv",
                },
                {
                    name for name in archive.namelist()
                    if name.endswith(".csv")
                },
            )
            self.assertNotIn("06_derived/replay_response_trajectory.py", archive.namelist())

        identity_job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            kind=ExportJob.Kind.IDENTITY,
            package="identity",
        )
        run_export_job(identity_job)
        with zipfile.ZipFile(identity_job.artifact_path) as archive:
            identity_manifest = json.loads(archive.read("00_manifest/manifest.json"))
            roster_rows = list(csv.DictReader(StringIO(
                archive.read("01_identity/student_roster.csv").decode("utf-8-sig")
            )))
            self.assertEqual(identity_manifest["selected_tables"], ["student_roster"])
            self.assertEqual(len(roster_rows), 1)
            self.assertEqual(roster_rows[0]["student_number"], "all-package-audit")

        teaching_job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="teaching",
        )
        run_export_job(teaching_job)
        self.assertTrue(teaching_job.artifact_path.endswith(".xlsx"))
        workbook = load_workbook(teaching_job.artifact_path, read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["教学成绩", "导出血缘"])
        teaching_rows = list(workbook["教学成绩"].iter_rows(values_only=True))
        self.assertIn("student_number", teaching_rows[0])
        self.assertIn("grade_score", teaching_rows[0])
        self.assertNotIn("final_answer", teaching_rows[0])
        self.assertEqual(len(teaching_rows), 2)
        lineage_rows = list(workbook["导出血缘"].iter_rows(values_only=True))
        self.assertIn(
            ("export_snapshot_id", json.dumps(str(snapshot.snapshot_id), ensure_ascii=False)),
            lineage_rows,
        )

        attempt.draft_answer = {"text": "post-cutoff mutation"}
        attempt.save(update_fields=["draft_answer"])
        rerun_job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        run_export_job(rerun_job)
        with zipfile.ZipFile(rerun_job.artifact_path) as archive:
            self.assertEqual(
                archive.read("02_core/attempts.csv"),
                package_attempt_bytes["complete"],
            )

    def test_kt_sequences_only_include_scored_formal_submissions(self):
        submitted = self._session("KT-submitted")
        submitted.submission_sequence_index = 7
        submitted.save(update_fields=["submission_sequence_index"])
        self._session("KT-draft", status=Attempt.Status.DRAFT, draft_answer="draft")
        self._session("KT-timeout", status=Attempt.Status.TIMED_OUT)
        self._session("KT-skipped", status=Attempt.Status.UNANSWERED)
        pending = self._session("KT-pending")
        pending.is_correct = None
        pending.score = None
        pending.submission_sequence_index = 8
        pending.save(update_fields=["is_correct", "score", "submission_sequence_index"])

        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        run_export_job(job)

        with zipfile.ZipFile(job.artifact_path) as archive:
            kt_rows = list(csv.DictReader(StringIO(
                archive.read("06_derived/kt_sequences.csv").decode("utf-8-sig")
            )))
            calibration_headers = next(csv.reader(StringIO(
                archive.read("06_derived/calibration_features.csv").decode("utf-8-sig")
            )))

        self.assertEqual(len(kt_rows), 1)
        self.assertEqual(kt_rows[0]["student_number"], "KT-submitted")
        self.assertEqual(kt_rows[0]["sequence_type"], "formal_submission")
        self.assertEqual(kt_rows[0]["sequence_index"], "7")
        self.assertEqual(kt_rows[0]["outcome"], "1")
        self.assertIn("signed_prediction_error", calibration_headers)
        self.assertIn("calibration_error", calibration_headers)

    def test_source_bundle_catalog_covers_feedback_and_research_tool_metadata(self):
        from apps.analytics.export_catalog import (
            COMMON_FIELD_DESCRIPTIONS,
            EXPOSURE_EVENT_TYPES,
            TABLE_FIELD_DESCRIPTIONS,
        )

        self._session("META")
        self.version.configuration = {
            "feedback_mode": "full",
            "feedback_release_policy": "manual_release",
            "research_tools": {"doubt_enabled": True},
        }
        self.version.save(update_fields=["configuration"])

        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        release_rows = list(csv.DictReader(StringIO(
            base64.b64decode(
                snapshot.source_bundle["tables"]["release_metadata"]
            ).decode("utf-8-sig")
        )))
        class_task_headers = next(csv.reader(StringIO(
            base64.b64decode(
                snapshot.source_bundle["tables"]["class_experiment_tasks"]
            ).decode("utf-8-sig")
        )))

        release_row = next(
            row for row in release_rows
            if row["version_number"] == str(self.version.version_number)
        )
        self.assertEqual(release_row["feedback_mode"], "full")
        self.assertEqual(
            release_row["feedback_release_policy"],
            "manual_release",
        )
        self.assertTrue(
            {
                "feedback_exposure",
                "research_tool",
                "research_tool_requested",
                "research_tool_observed",
                "confidence_observed",
            }.issubset(EXPOSURE_EVENT_TYPES)
        )
        self.assertTrue(
            {"feedback_released_at", "feedback_released_by_id"}
            .issubset(class_task_headers)
        )
        self.assertIn(
            "research_tools_sha256",
            TABLE_FIELD_DESCRIPTIONS["research_tools"],
        )
        for field in (
            "feedback_mode",
            "feedback_release_policy",
            "feedback_released_at",
            "feedback_released_by_id",
            "signed_prediction_error",
        ):
            self.assertIn(field, COMMON_FIELD_DESCRIPTIONS)

    def test_source_bundle_keeps_csv_bytes_after_live_draft_identity_and_question_drift(self):
        attempt = self._session(
            "S1",
            status=Attempt.Status.DRAFT,
            draft_answer={"text": "pre-cutoff draft"},
        )
        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        self.assertTrue(hasattr(snapshot, "source_bundle"))

        attempt.draft_answer = {"text": "post-cutoff draft"}
        attempt.save(update_fields=["draft_answer"])
        attempt.session.student.display_name = "renamed after cutoff"
        attempt.session.student.save(update_fields=["display_name"])
        self.qs.text = "question body changed after cutoff"
        self.qs.save(update_fields=["text"])

        first_job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        with patch(
            "apps.analytics.services.ExperimentSession.objects.filter",
            side_effect=AssertionError("worker must read the immutable source bundle"),
        ):
            run_export_job(first_job)
        with zipfile.ZipFile(first_job.artifact_path) as archive:
            first_attempts = archive.read("02_core/attempts.csv")

        first_rows = list(csv.reader(StringIO(first_attempts.decode("utf-8-sig"))))
        headers = first_rows[0]
        self.assertEqual(first_rows[1][headers.index("student_name")], "S1")
        self.assertEqual(
            first_rows[1][headers.index("draft_answer")],
            '{"text": "pre-cutoff draft"}',
        )
        self.assertEqual(
            hashlib.sha256(first_attempts).hexdigest(),
            snapshot.source_bundle["resource_manifest"]["table_sha256"]["attempts"],
        )

        second_job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        run_export_job(second_job)
        with zipfile.ZipFile(second_job.artifact_path) as archive:
            self.assertEqual(archive.read("02_core/attempts.csv"), first_attempts)

    def test_source_bundle_keeps_a_pre_cutoff_attempt_after_later_submission(self):
        attempt = self._session(
            "S2",
            status=Attempt.Status.DRAFT,
            draft_answer={"text": "draft captured before submission"},
        )
        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        self.assertTrue(hasattr(snapshot, "source_bundle"))

        attempt.status = Attempt.Status.SUBMITTED
        attempt.final_answer = {"text": "submitted after cutoff"}
        attempt.submitted_at = snapshot.source_cutoff_at + timedelta(seconds=1)
        attempt.save(update_fields=["status", "final_answer", "submitted_at"])

        job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        run_export_job(job)
        with zipfile.ZipFile(job.artifact_path) as archive:
            rows = list(csv.reader(StringIO(archive.read("02_core/attempts.csv").decode("utf-8-sig"))))

        headers = rows[0]
        row = next(row for row in rows[1:] if row[headers.index("attempt_id")] == str(attempt.pk))
        self.assertEqual(row[headers.index("status")], Attempt.Status.DRAFT)
        self.assertEqual(
            row[headers.index("draft_answer")],
            '{"text": "draft captured before submission"}',
        )
        self.assertEqual(row[headers.index("final_answer")], "")

    def test_client_time_sync_export_distinguishes_zero_from_missing_measurements(self):
        attempt = self._session("sync-zero")
        observed_at = timezone.now()
        zero_measurement = ClientTimeSync.objects.create(
            session=attempt.session,
            client_instance_id=uuid.uuid4(),
            client_wall_time_utc=observed_at,
            client_monotonic_ms=Decimal("0.000"),
            client_time_origin=observed_at,
            server_received_at=observed_at,
            server_responded_at=observed_at,
            client_received_wall_time_utc=observed_at,
            client_received_monotonic_ms=Decimal("0.000"),
            estimated_clock_offset_ms=0,
            estimated_rtt_ms=0,
            sync_quality=ClientTimeSync.SyncQuality.VALID,
            finalized_at=observed_at,
        )
        missing_measurement = ClientTimeSync.objects.create(
            session=attempt.session,
            client_instance_id=uuid.uuid4(),
            client_wall_time_utc=observed_at,
            client_monotonic_ms=Decimal("1.000"),
            client_time_origin=observed_at,
            server_received_at=observed_at,
            server_responded_at=observed_at,
            client_received_wall_time_utc=observed_at,
            client_received_monotonic_ms=Decimal("1.000"),
            estimated_clock_offset_ms=None,
            estimated_rtt_ms=None,
            sync_quality=ClientTimeSync.SyncQuality.INVALID,
            finalized_at=observed_at,
        )
        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        run_export_job(job)

        with zipfile.ZipFile(job.artifact_path) as archive:
            rows = list(csv.DictReader(StringIO(
                archive.read("03_raw_events/client_time_syncs.csv").decode("utf-8-sig")
            )))

        zero_row = next(row for row in rows if row["time_sync_id"] == str(zero_measurement.pk))
        self.assertEqual(zero_row["client_received_monotonic_ms"], "0.000")
        self.assertEqual(zero_row["estimated_clock_offset_ms"], "0")
        self.assertEqual(zero_row["estimated_rtt_ms"], "0")
        missing_row = next(row for row in rows if row["time_sync_id"] == str(missing_measurement.pk))
        self.assertEqual(missing_row["estimated_clock_offset_ms"], "")
        self.assertEqual(missing_row["estimated_rtt_ms"], "")

    def test_legacy_snapshot_requires_a_new_freeze_and_keeps_existing_artifact(self):
        snapshot = ExportSnapshot.objects.create(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        artifact_path = os.path.join(settings.MEDIA_ROOT, "exports", "legacy-frozen-artifact.zip")
        os.makedirs(os.path.dirname(artifact_path), exist_ok=True)
        self.addCleanup(lambda: os.path.exists(artifact_path) and os.remove(artifact_path))
        with open(artifact_path, "wb") as artifact:
            artifact.write(b"already-materialized artifact")
        job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
            artifact_path=artifact_path,
        )

        with self.assertRaisesRegex(ExportContractError, "不可变源包"):
            run_export_job(job)

        self.assertTrue(os.path.exists(artifact_path))
        with open(artifact_path, "rb") as artifact:
            self.assertEqual(artifact.read(), b"already-materialized artifact")

    def test_complete_export_freezes_question_content_and_capture_v2_replay(self):
        from apps.assessment.capture_v2_services import (
            ResponseVersionConflict,
            ingest_capture_batch,
            record_response_pointer_conflict,
        )
        from apps.assessment.event_services import append_assessment_event
        from apps.assessment.recording_capture_services import ensure_capture_contract_v2
        from apps.assessment.response_snapshot_services import ResponseState

        capture_contract = ensure_capture_contract_v2()
        attempt = self._session(
            "S3",
            status=Attempt.Status.DRAFT,
            draft_answer="AB",
            capture_contract=capture_contract,
        )
        self.qs.text = "冻结后的完整题干"
        self.qs.options = [{"key": "A", "text": "甲"}, {"key": "B", "text": "乙"}]
        self.qs.correct_answer = {"answer": "AB"}
        self.qs.accepted_answers = ["AB", "A B"]
        self.qs.explanation = "冻结解析"
        self.qs.release_fingerprint = "question-release-fingerprint"
        self.qs.kc_mapping_snapshot = {"nodes": [{"kc_code": "KC", "name": "知识点", "depth": 0, "order": 1}]}
        self.qs.kc_mapping_schema_version = "question-kc-tree-v2"
        self.qs.kc_mapping_hash = "a" * 64
        self.qs.save()
        frozen_research_tools = {
            "calculator": {"enabled": True, "content_version": "calculator-v1"},
        }
        self.version.configuration = {"research_tools": frozen_research_tools}
        self.version.save(update_fields=["configuration"])

        primary_branch = uuid.uuid4()
        alternate_branch = uuid.uuid4()

        def transition(before, after, *, parent_event_id, branch_id, sequence):
            event_id = uuid.uuid4()
            kind = attempt.assignment.snapshot.question_type
            return {
                "event_id": str(event_id),
                "event_type": "response_transition",
                "event_version": 1,
                "session_id": attempt.session_id,
                "attempt_id": attempt.pk,
                "client_sequence_no": sequence,
                "client_wall_time_utc": timezone.now().isoformat(),
                "client_monotonic_ms": sequence * 100,
                "client_time_origin": "2026-09-06T00:00:00Z",
                "quality_flags": [],
                "app_version": "capture-v2",
                "payload": {
                    "capture_version": "capture-v2",
                    "operation_id": str(uuid.uuid4()),
                    "branch_id": str(branch_id),
                    "parent_event_id": str(parent_event_id) if parent_event_id else None,
                    "before_state": ResponseState(kind, before).state_json,
                    "after_state": ResponseState(kind, after).state_json,
                    "before_state_hash": ResponseState(kind, before).state_hash,
                    "after_state_hash": ResponseState(kind, after).state_hash,
                    "input_type": "insertText",
                    "source": "student",
                },
            }

        root = transition("", "A", parent_event_id=None, branch_id=primary_branch, sequence=2)
        continuation = transition("A", "AB", parent_event_id=root["event_id"], branch_id=primary_branch, sequence=3)
        alternate = transition("A", "AX", parent_event_id=root["event_id"], branch_id=alternate_branch, sequence=1)
        input_operation = {
            "event_id": str(uuid.uuid4()),
            "event_type": "input_operation",
            "event_version": 1,
            "session_id": attempt.session_id,
            "attempt_id": attempt.pk,
            "client_sequence_no": 1,
            "client_wall_time_utc": timezone.now().isoformat(),
            "client_monotonic_ms": 100,
            "client_time_origin": "2026-09-06T00:00:00Z",
            "quality_flags": [],
            "app_version": "capture-v2",
            "payload": {
                "capture_version": "capture-v2",
                "operation_id": root["payload"]["operation_id"],
                "branch_id": str(primary_branch),
                "parent_event_id": None,
                "input_type": "insertText",
                "composition": "none",
                "selection_start": 0,
                "selection_end": 0,
            },
        }
        for client_id, document in (
            (primary_branch, input_operation),
            (primary_branch, root),
            (primary_branch, continuation),
            (alternate_branch, alternate),
        ):
            ingest_capture_batch(
                student=attempt.session.student,
                document={"client_instance_id": str(client_id), "events": [document]},
            )

        resolution = append_assessment_event(
            event_type="response_branch_resolved",
            event_version=1,
            event_domain=AssessmentEvent.EventDomain.ASSESSMENT,
            actor_source=AssessmentEvent.ActorSource.SERVER,
            session=attempt.session,
            assignment=attempt.assignment,
            attempt=attempt,
            payload={
                "capture_version": "capture-v2",
                "operation_id": str(uuid.uuid4()),
                "candidate_event_ids": [continuation["event_id"], alternate["event_id"]],
                "selected_event_id": continuation["event_id"],
                "expected_version": 1,
            },
            capture_mode=AssessmentEvent.CaptureMode.NATIVE,
            app_version="capture-v2",
        )
        active_event = AssessmentEvent.objects.get(event_id=continuation["event_id"])
        attempt.active_response_event = active_event
        attempt.final_response_event = active_event
        attempt.response_version = 2
        attempt.current_response_snapshot_id = active_event.payload["after_snapshot_id"]
        attempt.final_response_snapshot_id = active_event.payload["after_snapshot_id"]
        attempt.formal_capture_boundary = {
            "capture_contract_id": str(capture_contract.pk),
            "active_response_event_id": str(active_event.pk),
            "response_version": 2,
            "server_ingest_sequence_no": resolution.server_ingest_sequence_no,
            "student_id": attempt.session.student_id,
        }
        attempt.save(update_fields=[
            "active_response_event", "final_response_event", "response_version",
            "current_response_snapshot", "final_response_snapshot", "formal_capture_boundary",
        ])
        conflict_observation = record_response_pointer_conflict(
            attempt=attempt,
            conflict=ResponseVersionConflict(
                attempt,
                AssessmentEvent.objects.get(event_id=alternate["event_id"]),
            ),
            requested_expected_version=1,
            request_client_instance_id=uuid.uuid4(),
        )

        snapshot = create_export_snapshot(
            course=self.course,
            experiment=self.experiment,
            created_by=self.teacher,
        )
        self.qs.text = "捕获后被篡改的题干"
        self.qs.correct_answer = {"answer": "tampered"}
        self.qs.save(update_fields=["text", "correct_answer"])
        self.version.configuration = {"research_tools": {"calculator": {"enabled": False}}}
        self.version.save(update_fields=["configuration"])

        job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="complete",
        )
        run_export_job(job)

        with zipfile.ZipFile(job.artifact_path) as archive:
            question_content = list(csv.DictReader(StringIO(
                archive.read("01_design_metadata/question_content.csv").decode("utf-8-sig")
            )))
            trajectory_bytes = archive.read("06_derived/response_trajectory.csv")
            trajectory = list(csv.DictReader(StringIO(trajectory_bytes.decode("utf-8-sig"))))
            response_events = list(csv.DictReader(StringIO(
                archive.read("03_raw_events/response_events.csv").decode("utf-8-sig")
            )))
            assessment_events = list(csv.DictReader(StringIO(
                archive.read("03_raw_events/assessment_events.csv").decode("utf-8-sig")
            )))
            replay = list(csv.DictReader(StringIO(
                archive.read("06_derived/response_replay.csv").decode("utf-8-sig")
            )))
            session_summary = list(csv.DictReader(StringIO(
                archive.read("06_derived/session_summary.csv").decode("utf-8-sig")
            )))
            attempt_rows = list(csv.DictReader(StringIO(
                archive.read("02_core/attempts.csv").decode("utf-8-sig")
            )))
            research_tools = list(csv.DictReader(StringIO(
                archive.read("01_design_metadata/research_tools.csv").decode("utf-8-sig")
            )))
            replay_script = archive.read("06_derived/replay_response_trajectory.py")

        content_row = next(row for row in question_content if row["question_snapshot_id"] == str(self.qs.pk))
        self.assertEqual(content_row["text"], "冻结后的完整题干")
        self.assertEqual(json.loads(content_row["correct_answer"]), {"answer": "AB"})
        self.assertEqual(json.loads(content_row["accepted_answers"]), ["AB", "A B"])
        self.assertEqual(content_row["explanation"], "冻结解析")
        self.assertEqual(content_row["release_fingerprint"], "question-release-fingerprint")
        tool_row = next(row for row in research_tools if row["version_id"] == str(self.version.pk))
        self.assertEqual(json.loads(tool_row["research_tools"]), frozen_research_tools)

        by_event_id = {row["event_id"]: row for row in trajectory}
        self.assertEqual(set(by_event_id), {root["event_id"], continuation["event_id"], alternate["event_id"]})
        self.assertEqual(by_event_id[continuation["event_id"]]["parent_event_id"], root["event_id"])
        response_event = next(
            row for row in response_events
            if row["event_id"] == continuation["event_id"]
        )
        self.assertEqual(json.loads(response_event["after_state"])["value"], "AB")
        self.assertEqual(
            response_event["response_snapshot_id"],
            active_event.payload["after_snapshot_id"],
        )
        conflict_raw_event = next(
            row for row in assessment_events
            if row["event_id"] == str(conflict_observation.pk)
        )
        conflict_payload = json.loads(conflict_raw_event["payload"])
        self.assertEqual(conflict_raw_event["actor_source"], "server")
        self.assertEqual(
            conflict_payload["candidate_event_id"], alternate["event_id"],
        )
        self.assertEqual(
            conflict_payload["current_event_id"], continuation["event_id"],
        )
        self.assertEqual(conflict_payload["requested_expected_version"], 1)
        self.assertEqual(conflict_payload["current_response_version"], 2)
        self.assertFalse({"candidate_answer", "current_answer", "final_answer", "draft_answer"} & set(conflict_payload))
        conflict_response_event = next(
            row for row in response_events
            if row["event_id"] == str(conflict_observation.pk)
        )
        self.assertEqual(
            conflict_response_event["event_type"],
            "response_pointer_conflict_observed",
        )
        input_operation_response_event = next(
            row for row in response_events
            if row["event_id"] == input_operation["event_id"]
        )
        self.assertEqual(input_operation_response_event["event_type"], "input_operation")
        self.assertEqual(
            json.loads(input_operation_response_event["payload"])["operation_id"],
            root["payload"]["operation_id"],
        )

        process_job = ExportJob.objects.create(
            course=self.course,
            experiment=self.experiment,
            snapshot=snapshot,
            created_by=self.teacher,
            package="process",
        )
        run_export_job(process_job)
        with zipfile.ZipFile(process_job.artifact_path) as archive:
            process_assessment_events = list(csv.DictReader(StringIO(
                archive.read("03_raw_events/assessment_events.csv").decode("utf-8-sig")
            )))
            process_response_events = list(csv.DictReader(StringIO(
                archive.read("03_raw_events/response_events.csv").decode("utf-8-sig")
            )))
        self.assertEqual(
            next(
                row for row in process_assessment_events
                if row["event_id"] == str(conflict_observation.pk)
            )["payload"],
            conflict_raw_event["payload"],
        )
        self.assertEqual(
            next(
                row for row in process_response_events
                if row["event_id"] == str(conflict_observation.pk)
            )["event_type"],
            "response_pointer_conflict_observed",
        )
        self.assertEqual(
            next(
                row for row in process_response_events
                if row["event_id"] == input_operation["event_id"]
            )["event_type"],
            "input_operation",
        )

        replay_chain = []
        current_event_id = continuation["event_id"]
        while current_event_id:
            row = by_event_id[current_event_id]
            replay_chain.append(row)
            current_event_id = row["parent_event_id"]
        replay_chain.reverse()
        for index, row in enumerate(replay_chain):
            before_state = json.loads(row["before_state"])
            after_state = json.loads(row["after_state"])
            self.assertEqual(
                row["before_state_hash"],
                ResponseState(self.qs.question_type, before_state["value"]).state_hash,
            )
            self.assertEqual(
                row["after_state_hash"],
                ResponseState(self.qs.question_type, after_state["value"]).state_hash,
            )
            if index:
                self.assertEqual(replay_chain[index - 1]["after_state_hash"], row["before_state_hash"])
        self.assertEqual(json.loads(replay_chain[-1]["after_state"])["value"], "AB")

        replay_row = next(row for row in replay if row["attempt_id"] == str(attempt.pk))
        self.assertEqual(replay_row["active_response_event_id"], continuation["event_id"])
        self.assertEqual(replay_row["final_response_event_id"], continuation["event_id"])
        self.assertEqual(json.loads(replay_row["active_state"])["value"], "AB")
        resolutions = json.loads(replay_row["branch_resolutions"])
        self.assertEqual(resolutions[0]["event_id"], str(resolution.pk))
        self.assertEqual(resolutions[0]["selected_event_id"], continuation["event_id"])
        self.assertEqual(set(resolutions[0]["candidate_event_ids"]), {continuation["event_id"], alternate["event_id"]})

        summary_row = next(row for row in session_summary if row["session_id"] == str(attempt.session_id))
        self.assertEqual(summary_row["capture_contract_id"], str(capture_contract.pk))
        self.assertEqual(summary_row["response_transition_count"], "3")
        self.assertEqual(summary_row["response_branch_resolved_count"], "1")

        attempt_row = next(row for row in attempt_rows if row["attempt_id"] == str(attempt.pk))
        self.assertEqual(attempt_row["active_response_event_id"], continuation["event_id"])
        self.assertEqual(attempt_row["final_response_event_id"], continuation["event_id"])
        self.assertEqual(json.loads(attempt_row["formal_capture_boundary"])["response_version"], 2)
        self.assertEqual(
            hashlib.sha256(trajectory_bytes).hexdigest(),
            snapshot.source_bundle["resource_manifest"]["table_sha256"]["response_trajectory"],
        )
        with tempfile.TemporaryDirectory() as directory:
            script_path = os.path.join(directory, "replay_response_trajectory.py")
            with open(script_path, "wb") as script_file:
                script_file.write(replay_script)
            valid = subprocess.run(
                [sys.executable, script_path, job.artifact_path],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(valid.returncode, 0, valid.stderr)
            self.assertIn("validated", valid.stdout)

            with zipfile.ZipFile(job.artifact_path) as source_archive:
                files = {
                    name: source_archive.read(name)
                    for name in source_archive.namelist()
                }
            trajectory_rows = list(csv.DictReader(StringIO(
                files["06_derived/response_trajectory.csv"].decode("utf-8-sig")
            )))
            trajectory_headers = list(trajectory_rows[0])
            next(
                row for row in trajectory_rows
                if row["event_id"] == continuation["event_id"]
            )["parent_event_id"] = alternate["event_id"]
            trajectory_stream = StringIO(newline="")
            trajectory_writer = csv.DictWriter(trajectory_stream, fieldnames=trajectory_headers)
            trajectory_writer.writeheader()
            trajectory_writer.writerows(trajectory_rows)
            tampered_trajectory = ("\ufeff" + trajectory_stream.getvalue()).encode("utf-8")
            files["06_derived/response_trajectory.csv"] = tampered_trajectory
            raw_event_rows = list(csv.DictReader(StringIO(
                files["03_raw_events/assessment_events.csv"].decode("utf-8-sig")
            )))
            raw_event_headers = list(raw_event_rows[0])
            raw_event = next(
                row for row in raw_event_rows
                if row["event_id"] == continuation["event_id"]
            )
            raw_payload = json.loads(raw_event["payload"])
            raw_payload["parent_event_id"] = alternate["event_id"]
            raw_event["payload"] = json.dumps(raw_payload, ensure_ascii=False)
            raw_event_stream = StringIO(newline="")
            raw_event_writer = csv.DictWriter(raw_event_stream, fieldnames=raw_event_headers)
            raw_event_writer.writeheader()
            raw_event_writer.writerows(raw_event_rows)
            tampered_raw_events = ("\ufeff" + raw_event_stream.getvalue()).encode("utf-8")
            files["03_raw_events/assessment_events.csv"] = tampered_raw_events
            manifest = json.loads(files["00_manifest/manifest.json"])
            manifest["resource_manifest"]["table_sha256"]["response_trajectory"] = hashlib.sha256(
                tampered_trajectory
            ).hexdigest()
            manifest["resource_manifest"]["table_sha256"]["assessment_events"] = hashlib.sha256(
                tampered_raw_events
            ).hexdigest()
            files["00_manifest/manifest.json"] = json.dumps(
                manifest,
                ensure_ascii=False,
                sort_keys=True,
            ).encode("utf-8")
            tampered_archive_path = os.path.join(directory, "tampered.zip")
            with zipfile.ZipFile(tampered_archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for name, body in files.items():
                    archive.writestr(name, body)
            invalid = subprocess.run(
                [sys.executable, script_path, tampered_archive_path],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertNotEqual(invalid.returncode, 0)
            self.assertIn("parent after-state hash", invalid.stderr)

    def test_purge_invalidates_snapshot_and_job_with_reason(self):
        snapshot = create_export_snapshot(course=self.course, experiment=self.experiment, created_by=self.teacher)
        job = ExportJob.objects.create(course=self.course, experiment=self.experiment, snapshot=snapshot, created_by=self.teacher)
        request = DataPurgeRequest.objects.create(scope_type="course", scope_definition={"course_id": self.course.pk}, reason_code="collection_error", requested_by=self.teacher, preview_counts={})
        execute_purge_request(request, executed_by=self.teacher)
        snapshot.refresh_from_db(); job.refresh_from_db()
        self.assertIsNotNone(snapshot.invalidated_at)
        self.assertTrue(snapshot.invalidated_reason)
        self.assertEqual(job.status, ExportJob.Status.INVALIDATED)
