import csv
from io import BytesIO, StringIO

from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import ClassGroup, Course, CourseTeacher, User
from apps.assessment.models import Attempt, ExperimentSession, QuestionVisitEvent, StudentQuestionAssignment
from apps.audit.models import AuditLog
from apps.experiments.models import Experiment, ExperimentVersion, QuestionSnapshot
from apps.question_bank.models import Question, QuestionBank

class ExportPermissionTests(TestCase):
    def test_other_teacher_cannot_export_experiment(self):
        owner = User.objects.create_user(username="owner", role=User.Role.TEACHER)
        other = User.objects.create_user(username="other", role=User.Role.TEACHER)
        bank = QuestionBank.objects.create(owner=owner, name="bank")
        experiment = Experiment.objects.create(owner=owner, question_bank=bank, name="exp")
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse("analytics:kt_export", args=[experiment.id])).status_code, 404)

    def test_course_collaborator_can_export_course_experiment(self):
        owner = User.objects.create_user(username="course-owner", role=User.Role.TEACHER)
        collaborator = User.objects.create_user(username="collaborator", role=User.Role.TEACHER)
        course = Course.objects.create(owner=owner, name="课程", code="C1", term="2026春")
        CourseTeacher.objects.create(course=course, teacher=collaborator)
        bank = QuestionBank.objects.create(course=course, created_by=owner, name="bank")
        experiment = Experiment.objects.create(
            course=course,
            created_by=owner,
            question_bank=bank,
            name="exp",
        )
        self.client.force_login(collaborator)

        response = self.client.get(reverse("analytics:kt_export", args=[experiment.id]))

        self.assertEqual(response.status_code, 200)


class IdentifiedExportTests(TestCase):
    all_fields = [
        "student_number", "student_name", "class_name",
        "course_name", "course_code", "course_term",
        "experiment_id", "experiment_version", "session_id", "session_status",
        "course_sequence", "stable_code", "source_bank_id", "source_bank_name", "bank_sequence",
        "question_code", "question_type", "primary_kc", "all_kcs",
        "teacher_difficulty", "empirical_difficulty",
        "assigned_sequence_index", "first_open_sequence_index",
        "submission_sequence_index", "visit_count",
        "first_opened_at", "last_opened_at", "submitted_at",
        "first_to_submit_duration_ms", "last_visit_to_submission_ms",
        "final_answer", "is_correct", "score", "attempt_status",
    ]

    def setUp(self):
        self.teacher = User.objects.create_user(username="teacher", password="x", role=User.Role.TEACHER)
        self.other_teacher = User.objects.create_user(username="other", password="x", role=User.Role.TEACHER)
        class_group = ClassGroup.objects.create(owner=self.teacher, academic_year="2026", name="中文实验班")
        self.student = User.objects.create_user(
            username="00123",
            student_number="00123",
            display_name="林晓雨",
            class_group=class_group,
            role=User.Role.STUDENT,
        )
        bank = QuestionBank.objects.create(owner=self.teacher, name="题库")
        question = Question.objects.create(
            bank=bank,
            sequence=1,
            question_type=Question.Type.SHORT,
            text="中文题目",
            correct_answer="答案",
        )
        self.experiment = Experiment.objects.create(owner=self.teacher, question_bank=bank, name="中文实验")
        version = ExperimentVersion.objects.create(
            experiment=self.experiment,
            version_number=2,
            published_by=self.teacher,
            course_name_snapshot="认知诊断导论",
            course_code_snapshot="CD101",
            course_term_snapshot="2026春",
        )
        self.snapshot = QuestionSnapshot.objects.create(
            version=version,
            source_question=question,
            source_sequence=1,
            course_sequence=42,
            stable_code="C000001-Q000042",
            source_bank_id=7,
            source_bank_name="题库名称快照",
            bank_sequence=1,
            code="Q1",
            question_type=Question.Type.SHORT,
            text="中文题目",
            correct_answer="答案",
            primary_kc="中文知识点",
            all_kcs=["中文知识点", "次级知识点"],
        )
        self.session = ExperimentSession.objects.create(
            version=version,
            student=self.student,
            class_snapshot_name="2026 中文实验班",
            status=ExperimentSession.Status.COMPLETED,
        )
        assignment = StudentQuestionAssignment.objects.create(
            session=self.session,
            snapshot=self.snapshot,
            assigned_sequence_index=1,
        )
        self.attempt = Attempt.objects.create(
            assignment=assignment,
            session=self.session,
            status=Attempt.Status.SUBMITTED,
            first_open_sequence_index=1,
            submission_sequence_index=1,
            visit_count=2,
            final_answer="=1+1",
            is_correct=True,
            score=1,
        )
        QuestionVisitEvent.objects.create(
            session=self.session,
            attempt=self.attempt,
            visit_index=1,
            event_sequence=1,
            entry_reason="navigation_click",
        )
        self.export_url = f"/analytics/experiments/{self.experiment.pk}/identified.xlsx"
        self.navigation_url = f"/analytics/experiments/{self.experiment.pk}/navigation-identified.csv"
        self.client.force_login(self.teacher)

    def test_exports_selected_columns_in_allowlist_order_with_literal_unicode_text(self):
        response = self.client.post(
            self.export_url,
            {"fields": ["final_answer", "student_name", "student_number"]},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("filename*=UTF-8''", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["KT数据"])
        sheet = workbook["KT数据"]
        self.assertEqual([cell.value for cell in sheet[1]], ["学号", "姓名", "最终答案"])
        self.assertEqual([cell.value for cell in sheet[2]], ["00123", "林晓雨", "=1+1"])
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertEqual(sheet["C2"].data_type, "s")
        log = AuditLog.objects.get(action="experiment.identified_export")
        self.assertEqual(
            log.details,
            {"fields": ["student_number", "student_name", "final_answer"], "row_count": 1},
        )

    def test_course_fields_come_from_publish_snapshot(self):
        response = self.client.post(
            self.export_url,
            {"fields": ["course_term", "course_name", "course_code"]},
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        rows = list(workbook["KT数据"].iter_rows(values_only=True))
        self.assertEqual(rows[0], ("课程名称", "课程代码", "学期"))
        self.assertEqual(rows[1], ("认知诊断导论", "CD101", "2026春"))

    def test_course_question_identity_fields_use_the_fixed_allowlist_order_and_literal_stable_code(self):
        response = self.client.post(
            self.export_url,
            {
                "fields": [
                    "source_bank_name",
                    "stable_code",
                    "bank_sequence",
                    "course_sequence",
                    "source_bank_id",
                ]
            },
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        sheet = workbook["KT数据"]
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            ["课程永久题号", "稳定课程代码", "题库 ID", "题库名称快照", "题库内序号"],
        )
        self.assertEqual(
            [cell.value for cell in sheet[2]],
            [42, "C000001-Q000042", 7, "题库名称快照", 1],
        )
        self.assertEqual(sheet["B2"].data_type, "s")

    def test_identified_rows_sort_by_student_then_course_sequence_with_legacy_nulls_last(self):
        source_question = self.snapshot.source_question
        version = self.snapshot.version
        early_snapshot = QuestionSnapshot.objects.create(
            version=version,
            source_question=source_question,
            source_sequence=2,
            course_sequence=1,
            stable_code="C000001-Q000001",
            code="Q-early",
            question_type=Question.Type.SHORT,
            text="较早课程题号",
            correct_answer="答案",
            primary_kc="中文知识点",
            all_kcs=["中文知识点"],
        )
        legacy_snapshot = QuestionSnapshot.objects.create(
            version=version,
            source_question=source_question,
            source_sequence=3,
            code="Q-legacy",
            question_type=Question.Type.SHORT,
            text="遗留题目",
            correct_answer="答案",
            primary_kc="中文知识点",
            all_kcs=["中文知识点"],
        )
        for snapshot, assigned_index in ((early_snapshot, 2), (legacy_snapshot, 3)):
            assignment = StudentQuestionAssignment.objects.create(
                session=self.session,
                snapshot=snapshot,
                assigned_sequence_index=assigned_index,
            )
            Attempt.objects.create(session=self.session, assignment=assignment)

        earlier_student = User.objects.create_user(
            username="00002",
            student_number="00002",
            display_name="前置学生",
            role=User.Role.STUDENT,
        )
        earlier_session = ExperimentSession.objects.create(
            version=version,
            student=earlier_student,
            class_snapshot_name="前置班级",
        )
        earlier_assignment = StudentQuestionAssignment.objects.create(
            session=earlier_session,
            snapshot=self.snapshot,
            assigned_sequence_index=1,
        )
        Attempt.objects.create(session=earlier_session, assignment=earlier_assignment)

        response = self.client.post(
            self.export_url,
            {"fields": ["student_number", "course_sequence", "question_code"]},
        )

        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(BytesIO(response.content), data_only=True)["KT数据"].iter_rows(values_only=True))
        self.assertEqual(
            rows[1:],
            [
                ("00002", 42, "Q1"),
                ("00123", 1, "Q-early"),
                ("00123", 42, "Q1"),
                ("00123", None, "Q-legacy"),
            ],
        )

    def test_rejects_empty_or_unknown_fields_and_requires_post(self):
        self.assertEqual(self.client.get(self.export_url).status_code, 405)
        self.assertEqual(self.client.post(self.export_url, {}).status_code, 400)
        self.assertEqual(
            self.client.post(self.export_url, {"fields": ["student_number", "password"]}).status_code,
            400,
        )
        self.assertFalse(AuditLog.objects.filter(action="experiment.identified_export").exists())

    def test_enforces_owner_and_csrf(self):
        self.client.force_login(self.other_teacher)
        self.assertEqual(
            self.client.post(self.export_url, {"fields": ["student_number"]}).status_code,
            404,
        )

        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.teacher)
        self.assertEqual(
            csrf_client.post(self.export_url, {"fields": ["student_number"]}).status_code,
            403,
        )

    def test_course_collaborator_cannot_download_identified_exports(self):
        course = Course.objects.create(
            owner=self.teacher,
            name="协作课程",
            code="COLLAB",
            term="2026春",
        )
        CourseTeacher.objects.create(course=course, teacher=self.other_teacher)
        bank = QuestionBank.objects.create(
            course=course,
            created_by=self.teacher,
            name="协作课程题库",
        )
        experiment = Experiment.objects.create(
            course=course,
            created_by=self.teacher,
            question_bank=bank,
            name="协作课程实验",
        )
        self.client.force_login(self.other_teacher)

        identified = self.client.post(
            reverse("analytics:identified_export", args=[experiment.pk]),
            {"fields": ["student_number"]},
        )
        navigation = self.client.post(
            reverse("analytics:identified_navigation_export", args=[experiment.pk]),
        )

        self.assertEqual(identified.status_code, 404)
        self.assertEqual(navigation.status_code, 404)
        self.assertFalse(
            AuditLog.objects.filter(
                action__in=[
                    "experiment.identified_export",
                    "experiment.identified_navigation_export",
                ]
            ).exists()
        )

    def test_identified_navigation_contains_original_identity_and_audit(self):
        response = self.client.post(self.navigation_url)

        self.assertEqual(response.status_code, 200)
        text = response.content.decode("utf-8-sig")
        rows = list(csv.reader(StringIO(text)))
        self.assertEqual(rows[0][:3], ["student_number", "student_name", "class_name"])
        self.assertEqual(rows[1][:3], ["00123", "林晓雨", "2026 中文实验班"])
        log = AuditLog.objects.get(action="experiment.identified_navigation_export")
        self.assertEqual(log.details, {"row_count": 1})

    def test_legacy_anonymous_exports_keep_their_identity_boundary(self):
        kt = self.client.get(reverse("analytics:kt_export", args=[self.experiment.pk]))
        navigation = self.client.get(reverse("analytics:navigation_export", args=[self.experiment.pk]))

        kt_rows = list(csv.reader(StringIO(kt.content.decode("utf-8-sig"))))
        navigation_rows = list(csv.reader(StringIO(navigation.content.decode("utf-8-sig"))))
        self.assertEqual(kt_rows[0][0], "anonymous_student_id")
        self.assertNotIn("student_number", kt_rows[0])
        self.assertNotIn("student_name", kt_rows[0])
        self.assertEqual(navigation_rows[0][0], "anonymous_student_id")
        self.assertNotIn("student_number", navigation_rows[0])
        self.assertNotIn("student_name", navigation_rows[0])

    def test_legacy_anonymous_csv_urls_headers_and_question_code_remain_exactly_compatible(self):
        kt_url = reverse("analytics:kt_export", args=[self.experiment.pk])
        navigation_url = reverse("analytics:navigation_export", args=[self.experiment.pk])
        identity_url = reverse("analytics:identity_export", args=[self.experiment.pk])
        self.assertEqual(kt_url, f"/analytics/experiments/{self.experiment.pk}/kt.csv")
        self.assertEqual(navigation_url, f"/analytics/experiments/{self.experiment.pk}/navigation.csv")
        self.assertEqual(identity_url, f"/analytics/experiments/{self.experiment.pk}/identity.csv")

        kt_rows = list(csv.reader(StringIO(self.client.get(kt_url).content.decode("utf-8-sig"))))
        navigation_rows = list(
            csv.reader(StringIO(self.client.get(navigation_url).content.decode("utf-8-sig")))
        )
        identity_rows = list(
            csv.reader(StringIO(self.client.get(identity_url).content.decode("utf-8-sig")))
        )
        self.assertEqual(
            kt_rows[0],
            [
                "anonymous_student_id", "experiment_id", "experiment_version", "session_id",
                "question_code", "primary_kc", "all_kcs", "question_type",
                "assigned_sequence_index", "first_open_sequence_index",
                "submission_sequence_index", "first_opened_at", "last_opened_at",
                "submitted_at", "visit_count", "first_to_submit_duration_ms",
                "last_visit_to_submission_ms", "final_answer", "is_correct", "score",
                "status", "teacher_difficulty", "empirical_difficulty", "class_name",
            ],
        )
        self.assertEqual(
            navigation_rows[0],
            [
                "anonymous_student_id", "experiment_id", "session_id", "event_sequence",
                "from_question", "to_question", "visit_index", "entry_reason",
                "client_occurred_at", "server_received_at",
            ],
        )
        self.assertEqual(
            identity_rows[0],
            ["anonymous_student_id", "student_number", "student_name", "class_name"],
        )
        self.assertEqual(kt_rows[1][4], "Q1")
        self.assertEqual(navigation_rows[1][5], "Q1")

    def test_all_allowed_fields_export_as_one_question_level_sheet(self):
        response = self.client.post(self.export_url, {"fields": self.all_fields})

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
        rows = list(workbook["KT数据"].iter_rows(values_only=True))
        self.assertEqual(len(rows[0]), 34)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][:7], ("学号", "姓名", "班级", "课程名称", "课程代码", "学期", "实验 ID"))
        self.assertEqual(rows[1][:3], ("00123", "林晓雨", "2026 中文实验班"))
        self.assertEqual(rows[1][9], ExperimentSession.Status.COMPLETED)
        self.assertEqual(rows[1][-1], Attempt.Status.SUBMITTED)

    def test_exports_large_attempt_set_without_changing_row_grain(self):
        snapshot = self.attempt.assignment.snapshot
        version = snapshot.version
        source_question = snapshot.source_question
        snapshots = QuestionSnapshot.objects.bulk_create(
            [
                QuestionSnapshot(
                    version=version,
                    source_question=source_question,
                    source_sequence=index,
                    code=f"Q{index}",
                    question_type=Question.Type.SHORT,
                    text=f"题目 {index}",
                    correct_answer="答案",
                    primary_kc="中文知识点",
                    all_kcs=["中文知识点"],
                )
                for index in range(2, 1002)
            ]
        )
        assignments = StudentQuestionAssignment.objects.bulk_create(
            [
                StudentQuestionAssignment(
                    session=self.session,
                    snapshot=item,
                    assigned_sequence_index=index,
                )
                for index, item in enumerate(snapshots, start=2)
            ]
        )
        Attempt.objects.bulk_create(
            [
                Attempt(
                    assignment=assignment,
                    session=self.session,
                    status=Attempt.Status.UNVISITED,
                )
                for assignment in assignments
            ]
        )

        response = self.client.post(self.export_url, {"fields": ["student_number", "question_code"]})

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        self.assertEqual(sum(1 for _row in workbook["KT数据"].iter_rows(values_only=True)), 1002)
        log = AuditLog.objects.filter(action="experiment.identified_export").latest("created_at")
        self.assertEqual(log.details["row_count"], 1001)
