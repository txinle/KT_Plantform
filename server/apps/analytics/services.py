import base64
import hashlib
import json
import csv
import io
import os
import subprocess
import zipfile
import threading
import uuid
from contextlib import contextmanager
from datetime import timedelta
from types import SimpleNamespace

from django.conf import settings
from django.db import connection, models, transaction
from django.db.models.functions import Concat, Trim
from django.utils import timezone
from openpyxl import Workbook

from apps.assessment.models import (
    AssessmentEvent, Attempt, ExperimentSession, EventSchemaArtifact,
    StudentQuestionAssignment, ConfidenceAssignment, SimulationSnapshot,
    GradeRecord, ResponseInputObservation, ResponseSnapshot, SubmissionEvent,
    CaptureContractArtifact, ClientCaptureCheckpoint, ClientTimeSync,
)
from apps.experiments.models import (
    ClassExperimentTask,
    ConditionAssignment,
    Experiment,
    ExperimentClassSnapshot,
    ExperimentVersion,
    ExperimentVersionParticipantSnapshot,
    ItemInstance,
    PublishPreviewClassResult,
    PublishPreviewRun,
    QuestionSnapshot,
    ReleaseSnapshotQuestionMapping,
    ScheduleRevision,
)
from .export_catalog import (
    build_field_dictionary,
    EXPOSURE_EVENT_TYPES,
    PACKAGE_TABLES,
    PLACEHOLDER_TABLES,
    RESPONSE_EVENT_TYPES,
    SIMULATION_EVENT_TYPES,
    TABLE_PATHS,
)

from .models import (
    ExportJob, ExportSnapshot, RecordingCoverageAttempt,
    RecordingCoverageSession,
)
from .data_scope import DataScope


class ExportContractError(ValueError):
    pass


UNRECORDED_CLASS_FILTER = "__unrecorded__"
EXPORT_DERIVATION_VERSION = "v2"
LEGACY_EXPORT_DERIVATION_VERSION = "v1"
FIELD_DICTIONARY_PATH = "00_manifest/field_dictionary.json"
REPLAY_SCRIPT_PATH = "06_derived/replay_response_trajectory.py"
IDENTITY_ROSTER_PATH = "01_identity/student_roster.csv"
FROZEN_RESOURCE_PATHS = (FIELD_DICTIONARY_PATH, REPLAY_SCRIPT_PATH)


def _class_snapshot_filter_value(class_name):
    return "" if class_name == UNRECORDED_CLASS_FILTER else class_name


def _validate_export_scope(scope, package):
    if package == "experiment" and not scope.is_complete:
        raise ExportContractError("请先选择实验和班级")
    if package in {"teaching", "kt", "cd"} and not scope.is_single_version:
        raise ExportContractError(f"{package.upper()} 导出必须选择单一实验版本")


def _set_export_freeze_isolation():
    """Start a PostgreSQL source capture on one repeatable-read snapshot."""
    if connection.vendor == "postgresql":
        with connection.cursor() as cursor:
            cursor.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ")


def _export_freeze_strategy(*, entered_with_outer_transaction):
    if connection.vendor != "postgresql":
        return "source_bundle"
    if not entered_with_outer_transaction:
        return "repeatable_read_source_bundle"
    if any(
        getattr(block, "_from_testcase", False)
        for block in getattr(connection, "atomic_blocks", ())
    ):
        # Django TestCase owns the surrounding transaction.  It cannot change
        # isolation after setup queries, and it does not permit concurrent
        # writers, so retain a truthful test-only provenance label.
        return "testcase_source_bundle"
    with connection.cursor() as cursor:
        cursor.execute("SHOW transaction_isolation")
        isolation = cursor.fetchone()[0].lower()
    if isolation != "repeatable read":
        raise ExportContractError(
            "导出源冻结需要自动提交调用，或外层 PostgreSQL REPEATABLE READ 事务。"
        )
    return "repeatable_read_source_bundle"


def create_export_snapshot(*, course=None, experiment=None, research_study=None,
                           created_by, filter_definition=None, scope=None,
                           package=None):
    data_scope = scope
    if scope is not None:
        if scope.course.pk != course.pk:
            raise ExportContractError("导出范围不属于当前课程")
        if experiment is not None and scope.experiment_id not in {None, experiment.pk}:
            raise ExportContractError("导出范围与实验不一致")
        _validate_export_scope(scope, package)
        filter_definition = scope.serialize()
    canonical_filter_definition = filter_definition or {}
    scope_hash = hashlib.sha256(
        json.dumps(canonical_filter_definition, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()
    entered_with_outer_transaction = connection.in_atomic_block
    database_snapshot_strategy = _export_freeze_strategy(
        entered_with_outer_transaction=entered_with_outer_transaction,
    )
    with transaction.atomic():
        # A production request normally enters from autocommit, so this is the
        # first statement of the source transaction. Django's TestCase wrapper
        # owns an already-started transaction and cannot change its isolation.
        if not entered_with_outer_transaction:
            _set_export_freeze_isolation()

        from apps.accounts.course_access import lock_course
        if research_study is not None:
            for locked_course_id in sorted(research_study.course_links.values_list("course_id", flat=True)):
                lock_course(locked_course_id)
        elif course is not None:
            lock_course(course.pk)
        requested_experiments = set((scope.experiment_ids if scope is not None else ()) or ())
        if experiment is not None:
            requested_experiments.add(experiment.pk)
        if requested_experiments and set(Experiment.objects.filter(
            pk__in=requested_experiments, course=course).values_list("pk", flat=True)) != requested_experiments:
            raise ExportContractError("导出范围内的实验已删除或变更，请重新选择范围。")

        version_queryset = ExperimentVersion.objects.none()
        if experiment is not None:
            version_queryset = experiment.versions.order_by("version_number")
        elif course is not None:
            version_queryset = ExperimentVersion.objects.filter(
                experiment__course=course
            ).order_by("experiment_id", "version_number")
        if scope is not None and scope.version_id is not None:
            version_queryset = version_queryset.filter(pk=scope.version_id)
        if scope is not None and scope.experiment_id is not None:
            version_queryset = version_queryset.filter(experiment_id=scope.experiment_id)
        if scope is not None and scope.experiment_ids:
            version_queryset = version_queryset.filter(experiment_id__in=scope.experiment_ids)
        versions = list(version_queryset.values_list("pk", "version_number"))
        if package in {"teaching", "kt", "cd"} and len(versions) != 1:
            raise ExportContractError(f"{package.upper()} 导出必须选择单一实验版本")
        cutoff = timezone.now()
        event_scope = (
            {"session__version__experiment_id": experiment.pk}
            if experiment is not None
            else {"session__version__experiment__course_id": course.pk}
        )
        if research_study is not None:
            course_ids = list(
                research_study.course_links.values_list("course_id", flat=True)
            )
            event_scope = {"session__version__experiment__course_id__in": course_ids}
        watermark = (
            AssessmentEvent.objects.filter(
                **event_scope, server_received_at__lte=cutoff
            ).aggregate(max_seq=models.Max("server_ingest_sequence_no"))["max_seq"]
            or 0
        )
        schema_hashes = {
            f"{row.event_type}:v{row.event_version}": row.schema_sha256
            for row in EventSchemaArtifact.objects.all().order_by(
                "event_type", "event_version"
            )
        }
        schema_manifest_hash = hashlib.sha256(
            json.dumps(schema_hashes, sort_keys=True).encode()
        ).hexdigest()
        coverage_run = None
        if course is not None and research_study is None and package != "experiment":
            coverage_mapping = dict(filter_definition or {})
            if experiment is not None and not coverage_mapping.get("experiment_id"):
                coverage_mapping["experiment_id"] = experiment.pk
            coverage_scope = data_scope or DataScope.from_mapping(course, coverage_mapping)
            finalized_sessions = list(
                coverage_scope.sessions()
                .filter(
                    status__in=("completed", "expired", "invalidated"),
                    terminal_at__lte=cutoff - timedelta(seconds=300),
                )
                .select_related("capture_contract")
                .order_by("pk")
            )
            if finalized_sessions:
                from .recording_coverage import generate_scope_coverage

                coverage_run = generate_scope_coverage(
                    finalized_sessions,
                    scope_type="export",
                    scope_definition={
                        "scope_hash": scope_hash,
                        "session_ids": [session.pk for session in finalized_sessions],
                    },
                    source_cutoff_at=cutoff,
                )

        snapshot = ExportSnapshot(
            course=course,
            experiment=experiment,
            research_study=research_study,
            created_by=created_by,
            cutoff_at=cutoff,
            source_cutoff_at=cutoff,
            event_high_watermark=watermark,
            schema_manifest_hash=schema_manifest_hash,
            database_snapshot_strategy=database_snapshot_strategy,
            selected_versions=[{"id": pk, "version": version} for pk, version in versions],
            filter_definition=canonical_filter_definition,
            scope_hash=scope_hash,
            coverage_run=coverage_run,
        )
        captured = _capture_export_source(
            SimpleNamespace(
                snapshot=snapshot,
                course_id=course.pk,
                experiment=experiment,
                experiment_id=experiment.pk if experiment is not None else None,
                research_study=research_study,
                research_study_id=research_study.pk if research_study is not None else None,
            )
        )
        snapshot.source_bundle = _build_source_bundle(snapshot, captured)
        if package == "experiment":
            from .research_dataset import build_research_bundle
            from .research_contract import FILES

            research_scope = data_scope or DataScope.from_mapping(course, canonical_filter_definition)
            _validate_export_scope(research_scope, package)
            dataset = build_research_bundle(captured, scope=research_scope)["files"]
            snapshot.source_bundle["research_files"] = {
                table["path"]: base64.b64encode(
                    json.dumps(dataset[table["key"]], ensure_ascii=False, indent=2).encode("utf-8")
                    if table["key"] == "question_library" else _csv_bytes(table["headers"], dataset[table["key"]])
                ).decode("ascii")
                for table in FILES
            }
        snapshot.source_bundle_hash = _source_bundle_hash(snapshot.source_bundle)
        snapshot.save(force_insert=True)
        return snapshot


def create_export_job(*, course, snapshot, created_by, experiment=None,
                      research_study=None, kind=ExportJob.Kind.RESEARCH,
                      package="complete", selected_tables=None,
                      selected_fields=None):
    """Reuse only this requester's identical, still valid frozen export."""
    from apps.accounts.course_access import lock_course
    frozen = snapshot.source_bundle or {}
    v2 = frozen.get('research_export_v2')
    revision = ({key: v2.get(key) for key in (
        'contract_version', 'course_asset_hash', 'experiment_session_asset_hash', 'materializations',
        'class_mode', 'class_id')} if isinstance(v2, dict) else {
            key: frozen.get(key) for key in ('resource_manifest', 'teaching_timing', 'research_files')})
    fingerprint = hashlib.sha256(json.dumps({
        'requester': created_by.pk, 'course': course.pk, 'experiment': getattr(experiment, 'pk', None),
        'study': getattr(research_study, 'pk', None), 'scope': snapshot.filter_definition,
        'versions': snapshot.selected_versions, 'revision': revision, 'kind': kind, 'package': package,
        'tables': selected_tables or [], 'fields': selected_fields or [],
    }, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    candidates = ExportJob.objects.filter(
        course=course, created_by=created_by, request_fingerprint=fingerprint,
        snapshot__invalidated_at=None, snapshot__redacted_at=None,
        status__in=(ExportJob.Status.QUEUED, ExportJob.Status.RUNNING, ExportJob.Status.SUCCEEDED),
    ).filter(models.Q(expires_at__isnull=True) | models.Q(expires_at__gt=timezone.now())).select_related('snapshot').only(
        'job_id', 'status', 'artifact_path', 'artifact_sha256', 'expires_at', 'snapshot_id', 'course_id', 'package', 'snapshot__filter_definition', 'snapshot__source_bundle_hash', 'snapshot__invalidated_at', 'snapshot__redacted_at')
    while True:
        inspected = list(candidates.all())
        checked = {candidate.pk: _artifact_signature(candidate) for candidate in inspected
                   if candidate.status == ExportJob.Status.SUCCEEDED and export_artifact_available(candidate)}
        inspected_status = {candidate.pk: candidate.status for candidate in inspected}
        retry_validation = False
        with transaction.atomic():
            locked_course = lock_course(course.pk)
            if not locked_course.can_manage(created_by):
                raise ExportContractError('当前用户已无权导出该课程数据。')
            if not ExportSnapshot.objects.filter(pk=snapshot.pk, invalidated_at=None, redacted_at=None).exists():
                raise ExportContractError('导出源已删除或失效，请重新选择范围。')
            for candidate in candidates.all().filter(
                    models.Q(expires_at__isnull=True) | models.Q(expires_at__gt=timezone.now())):
                if candidate.status != ExportJob.Status.SUCCEEDED or (
                        checked.get(candidate.pk) is not None and checked[candidate.pk] == _artifact_signature(candidate)):
                    candidate._reused = True
                    return candidate
                if inspected_status.get(candidate.pk) != ExportJob.Status.SUCCEEDED:
                    retry_validation = True
                    break
            if not retry_validation:
                return ExportJob.objects.create(
                    course=course,
                    experiment=experiment,
                    research_study=research_study,
                    snapshot=snapshot,
                    created_by=created_by,
                    kind=kind,
                    package=package,
                    selected_tables=selected_tables or [],
                    selected_fields=selected_fields or [],
                    status=ExportJob.Status.QUEUED,
                    progress_stage="queued",
                    request_fingerprint=fingerprint,
                )


def _artifact_signature(job):
    try:
        stat = os.stat(job.artifact_path)
        return job.artifact_path, job.artifact_sha256, stat.st_size, stat.st_mtime_ns
    except OSError:
        return None


def export_artifact_available(job):
    if not job.artifact_path or not job.artifact_sha256 or not os.path.isfile(job.artifact_path):
        return False
    try:
        before = _artifact_signature(job)
        return before is not None and _file_sha256(job.artifact_path) == job.artifact_sha256 and before == _artifact_signature(job)
    except OSError:
        return False


def _file_sha256(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def remove_export_artifact(job, *, reason, invalidated=False):
    """Remove a materialized artifact while retaining its audit identity."""
    if job.artifact_path and os.path.exists(job.artifact_path):
        artifact_dir = os.path.dirname(job.artifact_path)
        os.remove(job.artifact_path)
        if artifact_dir and os.path.isdir(artifact_dir) and not os.listdir(artifact_dir):
            os.rmdir(artifact_dir)
    job.status = ExportJob.Status.INVALIDATED if invalidated else ExportJob.Status.EXPIRED
    job.artifact_path = ""
    job.artifact_removed_at = timezone.now()
    job.artifact_removal_reason = reason
    job.save(update_fields=[
        "status", "artifact_path", "artifact_removed_at", "artifact_removal_reason"
    ])
    return job


def _git_commit():
    try:
        return subprocess.check_output(["git", "rev-parse", "HEAD"], text=True, stderr=subprocess.DEVNULL).strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def build_export_manifest(snapshot, *, selected_tables, row_counts=None, data_status=None,
                          experiment_name=None, version_number=None,
                          session_ids=None, attempt_ids=None,
                          capture_contracts=()):
    schemas = {
        f"{row.event_type}:v{row.event_version}": row.schema_sha256
        for row in EventSchemaArtifact.objects.all()
    }
    config = {
        "selected_tables": list(selected_tables),
        "selected_versions": snapshot.selected_versions,
        "filter_definition": snapshot.filter_definition,
    }
    filter_display = dict(snapshot.filter_definition)
    if filter_display.get("class_name") == UNRECORDED_CLASS_FILTER:
        filter_display["class_name"] = "未记录班级"
    provenance = {
        "export_snapshot": {"id": str(snapshot.snapshot_id), "source_cutoff_at": snapshot.source_cutoff_at.isoformat(), "event_high_watermark": snapshot.event_high_watermark},
        "schema": {"manifest_hash": snapshot.schema_manifest_hash, "event_schema_hashes": schemas},
        "derivation": {
            "version": EXPORT_DERIVATION_VERSION,
            "deterministic": True,
            "code_sha": _git_commit(),
        },
        "source_records": {"sessions": "assessment.ExperimentSession", "attempts": "assessment.Attempt", "events": "assessment.AssessmentEvent", "grades": "assessment.GradeRecord", "confidence": "assessment.ConfidenceAssignment", "coverage_runs": "analytics.RecordingCoverageRun"},
    }
    coverage_filter = {
        "session__version__experiment_id": snapshot.experiment_id
    } if snapshot.experiment_id else {
        "session__version__experiment__course_id": snapshot.course_id
    }
    coverage_query = RecordingCoverageSession.objects.filter(
            **coverage_filter,
            run__run_status="completed",
            run__source_cutoff_at__lte=snapshot.source_cutoff_at,
        )
    if getattr(snapshot, "coverage_run_id", None):
        coverage_query = coverage_query.filter(run_id=snapshot.coverage_run_id)
    if session_ids is not None:
        coverage_query = coverage_query.filter(session_id__in=session_ids)
    all_coverage_results = list(
        coverage_query.select_related("run", "session__capture_contract").order_by(
            "session_id", "-run__generated_at"
        )
    )
    latest_by_session = {}
    for item in all_coverage_results:
        latest_by_session.setdefault(item.session_id, item)
    coverage_results = list(latest_by_session.values())
    coverage_runs = {str(item.run_id): item.run for item in coverage_results}
    contract_rows = {
        str(item.session.capture_contract_id): item.session.capture_contract
        for item in coverage_results
        if item.session.capture_contract_id
    }
    contract_rows.update({
        str(contract.pk): contract
        for contract in capture_contracts
        if contract is not None
    })
    status_counts = {}
    for item in coverage_results:
        status_counts[item.overall_status] = status_counts.get(item.overall_status, 0) + 1
    run_ids = sorted(coverage_runs)
    coverage_versions = sorted(
        {run.coverage_version for run in coverage_runs.values()}
    )
    parameter_hashes = sorted({run.parameter_hash for run in coverage_runs.values()})
    coverage_attempt_query = RecordingCoverageAttempt.objects.filter(
        run_id__in=run_ids
    )
    if attempt_ids is not None:
        coverage_attempt_query = coverage_attempt_query.filter(
            attempt_id__in=attempt_ids
        )
    time_quality_counts = {}
    for quality in coverage_attempt_query.values_list("time_quality", flat=True):
        time_quality_counts[quality] = time_quality_counts.get(quality, 0) + 1
    checkpoint_query = ClientCaptureCheckpoint.objects.filter(
        server_received_at__lte=snapshot.source_cutoff_at
    )
    if session_ids is not None:
        checkpoint_query = checkpoint_query.filter(session_id__in=session_ids)
    elif snapshot.experiment_id:
        checkpoint_query = checkpoint_query.filter(
            session__version__experiment_id=snapshot.experiment_id
        )
    else:
        checkpoint_query = checkpoint_query.filter(
            session__version__experiment__course_id=snapshot.course_id
        )
    latest_checkpoint_by_instance = {}
    for checkpoint in checkpoint_query.order_by(
        "session_id", "client_instance_id", "-server_received_at", "-checkpoint_id"
    ):
        latest_checkpoint_by_instance.setdefault(
            (checkpoint.session_id, checkpoint.client_instance_id), checkpoint
        )
    return {
        "manifest_version": 2,
        "contains_direct_identifiers": True,
        "identity_columns": ["student_number", "student_name"],
        "primary_question_identifier": "question_code",
        "export_snapshot_id": str(snapshot.snapshot_id),
        "course_id": snapshot.course_id,
        "experiment_id": snapshot.experiment_id,
        "experiment_name": experiment_name,
        "version_number": version_number,
        "research_study_id": snapshot.research_study_id,
        "cutoff_at": snapshot.cutoff_at.isoformat(),
        "source_cutoff_at": snapshot.source_cutoff_at.isoformat(),
        "event_high_watermark": snapshot.event_high_watermark,
        "schema_manifest_hash": snapshot.schema_manifest_hash,
        "scope_hash": snapshot.scope_hash,
        "database_snapshot_strategy": snapshot.database_snapshot_strategy,
        "filter_definition": snapshot.filter_definition,
        "filter_display": filter_display,
        "selected_tables": list(selected_tables),
        "event_schema_hashes": schemas,
        "application_git_commit_sha": _git_commit(),
        "export_config_hash": hashlib.sha256(json.dumps(config, sort_keys=True).encode()).hexdigest(),
        "row_counts": row_counts or {},
        "data_status": data_status or {},
        "provenance": provenance,
        "capture_contract_ids": sorted(contract_rows),
        "capture_contract_hashes": sorted({row.schema_sha256 for row in contract_rows.values()}),
        "coverage_run_id": run_ids[0] if len(run_ids) == 1 else None,
        "coverage_run_ids": run_ids,
        "coverage_version": coverage_versions[0] if len(coverage_versions) == 1 else None,
        "coverage_parameter_hash": parameter_hashes[0] if len(parameter_hashes) == 1 else None,
        "coverage_source_cutoff": max((run.source_cutoff_at for run in coverage_runs.values()), default=None).isoformat() if run_ids else None,
        "coverage_source_high_water_marks": {run_id: coverage_runs[run_id].source_high_water_marks for run_id in run_ids},
        "coverage_status_counts": status_counts,
        "sequence_gap_count": sum(item.sequence_gap_count for item in coverage_results),
        "tail_gap_count": sum(bool(item.tail_gap_detected) for item in coverage_results),
        "dead_letter_count": sum(
            checkpoint.failed_event_count
            for checkpoint in latest_checkpoint_by_instance.values()
        ),
        "time_quality_counts": time_quality_counts,
    }


def manifest_bytes(manifest):
    return json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8")


def sha256_bytes(value):
    return hashlib.sha256(value).hexdigest()


def claim_export_job(*, worker_id, now=None):
    from .runtime_pressure import runtime_pressure

    if runtime_pressure()['mode'] in ('capture_priority', 'capture_emergency'):
        return None
    now = now or timezone.now()
    with transaction.atomic():
        job = (
            ExportJob.objects.select_for_update()
            .filter(
                models.Q(status=ExportJob.Status.QUEUED)
                | models.Q(status=ExportJob.Status.RUNNING, lease_expires_at__lt=now)
            )
            .order_by("created_at")
            .first()
        )
        if not job:
            return None
        job.status = ExportJob.Status.RUNNING
        job.locked_by = worker_id
        job.started_at = now
        job.heartbeat_at = now
        job.lease_expires_at = now + timedelta(seconds=90)
        job.progress_stage = "snapshot"
        job.save(update_fields=["status", "locked_by", "started_at", "heartbeat_at", "lease_expires_at", "progress_stage"])
        return job


def renew_export_job_lease(job, *, worker_id, now=None):
    now = now or timezone.now()
    updated = ExportJob.objects.filter(
        pk=job.pk,
        status=ExportJob.Status.RUNNING,
        locked_by=worker_id,
    ).update(
        heartbeat_at=now,
        lease_expires_at=now + timedelta(seconds=90),
    )
    if not updated:
        return None
    job.refresh_from_db(fields=["heartbeat_at", "lease_expires_at"])
    return job


def _csv_bytes(headers, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _capture_export_source(job, *, session_projection=False):
    """Read the live source only while ``create_export_snapshot`` is open."""
    snapshot = job.snapshot
    study_course_ids = list(job.research_study.course_links.values_list("course_id", flat=True)) if job.research_study_id else []
    study_student_ids = list(job.research_study.participants.values_list("student_id", flat=True)) if job.research_study_id else []
    filter_definition = snapshot.filter_definition or {}
    scope = None
    if filter_definition.get("scope_version"):
        try:
            scope = DataScope.from_mapping(snapshot.course, filter_definition)
        except (LookupError, ValueError) as error:
            raise ExportContractError(f"冻结的数据范围无效：{error}") from error
        session_queryset = scope.sessions()
    else:
        session_filter = {"version__experiment_id": job.experiment_id} if job.experiment_id else ({"version__experiment__course_id__in": study_course_ids, "student_id__in": study_student_ids} if study_course_ids else {"version__experiment__course_id": job.course_id})
        if filter_definition.get("version_id"):
            session_filter["version_id"] = int(filter_definition["version_id"])
        if filter_definition.get("class_name"):
            session_filter["class_snapshot_name"] = _class_snapshot_filter_value(filter_definition["class_name"])
        if filter_definition.get("student_number"):
            session_filter["student__student_number"] = filter_definition["student_number"]
        if filter_definition.get("session_id"):
            session_filter["pk"] = int(filter_definition["session_id"])
        if filter_definition.get("status") in {choice for choice, _label in ExperimentSession.Status.choices}:
            session_filter["status"] = filter_definition["status"]
        session_queryset = ExperimentSession.objects.filter(**session_filter)
    with transaction.atomic():
        sessions = list(
            session_queryset.filter(started_at__lte=snapshot.source_cutoff_at)
            .select_related("student", "version__experiment", "capture_contract")
        )
        session_ids = [s.pk for s in sessions]
        attempt_queryset = (
            scope.attempts() if scope is not None else Attempt.objects.filter(session_id__in=session_ids)
        )
        attempt_queryset = attempt_queryset.filter(
            session_id__in=session_ids,
            session__started_at__lte=snapshot.source_cutoff_at,
        ).filter(
            models.Q(submitted_at__isnull=True)
            | models.Q(submitted_at__lte=snapshot.source_cutoff_at)
        )
        legacy_attempt_status = filter_definition.get("status") in {
            choice for choice, _label in Attempt.Status.choices
        }
        if scope is None and legacy_attempt_status:
            attempt_queryset = attempt_queryset.filter(status=filter_definition["status"])
        attempts = list(attempt_queryset.select_related(
            "session", "assignment__snapshot", "assignment__item_instance", "submission_event"
        ))
        if (scope is not None and scope.attempt_status is not None) or legacy_attempt_status:
            selected_session_ids = {attempt.session_id for attempt in attempts}
            sessions = [session for session in sessions if session.pk in selected_session_ids]
            session_ids = [session.pk for session in sessions]
        event_queryset = scope.events() if scope is not None else AssessmentEvent.objects.all()
        events = list(
            event_queryset.filter(
                session_id__in=session_ids,
                server_received_at__lte=snapshot.source_cutoff_at,
                server_ingest_sequence_no__lte=snapshot.event_high_watermark,
            ).order_by("session_id", "server_ingest_sequence_no")
        )
        assignments_queryset = StudentQuestionAssignment.objects.filter(session_id__in=session_ids).select_related("snapshot", "item_instance")
        if (scope is not None and scope.attempt_status is not None) or legacy_attempt_status:
            assignments_queryset = assignments_queryset.filter(attempt__in=attempts)
        assignments = list(assignments_queryset)
        _validate_question_codes(assignments)
        confidence_assignments = list(
            ConfidenceAssignment.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                assigned_at__lte=snapshot.source_cutoff_at,
            )
        )
        simulations = list(
            SimulationSnapshot.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                created_at__lte=snapshot.source_cutoff_at,
            )
        )
        response_snapshots = list(
            ResponseSnapshot.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                created_at__lte=snapshot.source_cutoff_at,
            ).order_by("attempt_id", "created_at", "snapshot_id")
        )
        input_observations = list(
            ResponseInputObservation.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                observed_at__lte=snapshot.source_cutoff_at,
            ).select_related("source_event").order_by(
                "attempt_id", "client_instance_id", "client_sequence_no"
            )
        )
        submissions = list(
            SubmissionEvent.objects.filter(
                attempt_id__in=[a.pk for a in attempts],
                server_received_at__lte=snapshot.source_cutoff_at,
            )
        )
        grade_records = list(
            GradeRecord.objects.filter(
                submission__attempt_id__in=[a.pk for a in attempts],
                created_at__lte=snapshot.source_cutoff_at,
            ).select_related("submission")
        )
        version_ids = {
            int(item["id"] if isinstance(item, dict) else item[0])
            for item in snapshot.selected_versions
        } or {session.version_id for session in sessions}
        class_snapshot_queryset = ExperimentClassSnapshot.objects.filter(
            version_id__in=version_ids,
            created_at__lte=snapshot.source_cutoff_at,
        )
        participant_snapshot_queryset = (
            ExperimentVersionParticipantSnapshot.objects.filter(
                version_id__in=version_ids,
                created_at__lte=snapshot.source_cutoff_at,
            )
        )
        if scope is not None and scope.class_snapshot_id is not None:
            class_snapshot_queryset = class_snapshot_queryset.filter(
                pk=scope.class_snapshot_id
            )
            participant_snapshot_queryset = participant_snapshot_queryset.filter(
                class_snapshot_id=scope.class_snapshot_id
            )
        elif scope is not None and scope.class_id is not None:
            class_snapshot_queryset = class_snapshot_queryset.filter(
                source_class_group_id=scope.class_id
            )
            participant_snapshot_queryset = participant_snapshot_queryset.filter(
                class_snapshot__source_class_group_id=scope.class_id
            )
        if scope is not None and scope.student_id is not None:
            participant_snapshot_queryset = participant_snapshot_queryset.filter(
                student_id=scope.student_id
            )
            class_snapshot_queryset = class_snapshot_queryset.filter(
                pk__in=participant_snapshot_queryset.values("class_snapshot_id")
            )
        narrow_to_started_sessions = False
        if scope is not None:
            narrow_to_started_sessions = any(
                value is not None
                for value in (
                    scope.session_id,
                    scope.condition_id,
                    scope.session_status,
                    scope.attempt_status,
                    scope.started_from_utc,
                    scope.started_before_utc,
                )
            )
        else:
            if filter_definition.get("student_number"):
                participant_snapshot_queryset = participant_snapshot_queryset.filter(
                    student__student_number=filter_definition["student_number"]
                )
                class_snapshot_queryset = class_snapshot_queryset.filter(
                    pk__in=participant_snapshot_queryset.values("class_snapshot_id")
                )
            if filter_definition.get("class_name"):
                class_name = _class_snapshot_filter_value(
                    filter_definition["class_name"]
                )
                class_snapshot_queryset = class_snapshot_queryset.annotate(
                    display_name_filter=Trim(
                        Concat(
                            "academic_year_snapshot",
                            models.Value(" "),
                            "name_snapshot",
                        )
                    )
                ).filter(display_name_filter=class_name)
                participant_snapshot_queryset = participant_snapshot_queryset.filter(
                    class_snapshot_id__in=class_snapshot_queryset.values("pk")
                )
            narrow_to_started_sessions = bool(
                filter_definition.get("session_id")
                or filter_definition.get("status")
            )
        if narrow_to_started_sessions:
            participant_snapshot_ids = [
                session.participant_snapshot_id
                for session in sessions
                if session.participant_snapshot_id is not None
            ]
            participant_snapshot_queryset = participant_snapshot_queryset.filter(
                pk__in=participant_snapshot_ids
            )
            class_snapshot_queryset = class_snapshot_queryset.filter(
                pk__in=participant_snapshot_queryset.values("class_snapshot_id")
            )
        class_snapshots = list(
            class_snapshot_queryset.order_by("version_id", "created_at")
        )
        participant_snapshots = list(
            participant_snapshot_queryset.order_by("version_id", "student_id")
        )
        condition_assignments = list(
            ConditionAssignment.objects.filter(
                models.Q(session_id__in=session_ids)
                | models.Q(
                    subject_id__in=[session.student_id for session in sessions],
                    condition__experiment_id__in=[
                        session.version.experiment_id for session in sessions
                    ],
                ),
                assigned_at__lte=snapshot.source_cutoff_at,
            ).select_related("condition")
        )
    roster_rows = []
    if job.experiment_id:
        students = list(
            job.experiment.students.select_related("class_group").order_by("pk")
        )
        roster_rows = [
            [student.student_number, student.shown_name, str(student.class_group or "")]
            for student in students
        ]
    rows = _table_rows(
        job,
        sessions,
        assignments,
        attempts,
        events,
        confidence_assignments,
        simulations,
        response_snapshots=response_snapshots,
        input_observations=input_observations,
        submissions=submissions,
        grade_records=grade_records,
        condition_assignments=condition_assignments,
        class_snapshots=class_snapshots,
        participant_snapshots=participant_snapshots,
        snapshot=snapshot,
        session_projection=session_projection,
    )
    selected_version_number = None
    if filter_definition.get("version_id"):
        selected_version_number = next(
            (item["version"] for item in snapshot.selected_versions if str(item["id"]) == str(filter_definition["version_id"])),
            None,
        )
    elif len(snapshot.selected_versions) == 1:
        selected_version_number = snapshot.selected_versions[0]["version"]
    return {
        "attempt_ids": [attempt.pk for attempt in attempts],
        "attempts": attempts,
        "events": events,
        "experiment_name": job.experiment.name if job.experiment_id else None,
        "roster_rows": roster_rows,
        "rows": rows,
        "sessions": sessions,
        "session_ids": session_ids,
        "submissions": submissions,
        "version_number": selected_version_number,
    }


def _source_bundle_hash(source_bundle):
    return sha256_bytes(
        json.dumps(
            source_bundle,
            ensure_ascii=False,
            separators=(",", ":"),
            sort_keys=True,
        ).encode("utf-8")
    )


def _build_source_bundle(snapshot, captured):
    table_bytes = {}
    row_counts = {}
    table_sha256 = {}
    for table, path in TABLE_PATHS.items():
        if table not in captured["rows"]:
            continue
        headers, values = captured["rows"][table]
        body = _csv_bytes(headers, values)
        table_bytes[table] = base64.b64encode(body).decode("ascii")
        row_counts[table] = len(values)
        table_sha256[table] = sha256_bytes(body)

    source_tables = list(table_bytes)
    frozen_table_paths = {
        table: TABLE_PATHS[table]
        for table in source_tables
    }
    frozen_table_paths["student_roster"] = IDENTITY_ROSTER_PATH
    frozen_table_headers = {
        table: captured["rows"][table][0]
        for table in source_tables
    }
    frozen_table_headers["student_roster"] = (
        "student_number",
        "student_name",
        "class_name",
    )
    field_dictionary_body = manifest_bytes(build_field_dictionary(
        table_headers=frozen_table_headers,
        table_paths=frozen_table_paths,
    ))
    from .replay_artifact import standalone_response_replay_script

    replay_script_body = standalone_response_replay_script().encode("utf-8")
    resource_bodies = {
        FIELD_DICTIONARY_PATH: field_dictionary_body,
        REPLAY_SCRIPT_PATH: replay_script_body,
    }
    data_status = {
        table: "no_observation" if row_counts[table] == 0 else "observed"
        for table in source_tables
    }
    source_manifest = build_export_manifest(
        snapshot,
        selected_tables=source_tables,
        row_counts=row_counts,
        data_status=data_status,
        experiment_name=captured["experiment_name"],
        version_number=captured["version_number"],
        session_ids=captured["session_ids"],
        attempt_ids=captured["attempt_ids"],
        capture_contracts=(
            session.capture_contract
            for session in captured["sessions"]
            if session.capture_contract_id
        ),
    )
    roster_csv = _csv_bytes(
        ("student_number", "student_name", "class_name"),
        captured["roster_rows"],
    )
    events_by_attempt = {}
    for event in captured["events"]:
        events_by_attempt.setdefault(event.attempt_id, []).append(event)
    submissions_by_attempt = {
        submission.attempt_id: submission
        for submission in captured["submissions"]
    }
    from .teaching_data import summarize_attempt_timing

    teaching_timing = {}
    for attempt in captured["attempts"]:
        timing = summarize_attempt_timing(
            attempt,
            events=events_by_attempt.get(attempt.pk, []),
            submission=submissions_by_attempt.get(attempt.pk),
        )
        teaching_timing[str(attempt.pk)] = {
            "display_duration_ms": timing.display_duration_ms,
            "display_kind": str(timing.display_kind),
            "interval_quality": timing.interval_quality,
        }
    return {
        "version": 1,
        "manifest": source_manifest,
        "resource_manifest": {
            "identity_roster_sha256": sha256_bytes(roster_csv),
            "row_counts": row_counts,
            "table_paths": frozen_table_paths,
            "table_sha256": table_sha256,
            "resource_freeze_status": "frozen",
            "resource_sha256": {
                path: sha256_bytes(body)
                for path, body in resource_bodies.items()
            },
        },
        "student_roster_csv": base64.b64encode(roster_csv).decode("ascii"),
        "tables": table_bytes,
        "resources": {
            path: base64.b64encode(body).decode("ascii")
            for path, body in resource_bodies.items()
        },
        "teaching_timing": teaching_timing,
    }


def _load_source_bundle(snapshot):
    source_bundle = getattr(snapshot, "source_bundle", None)
    if not isinstance(source_bundle, dict) or source_bundle.get("version") != 1:
        raise ExportContractError(
            "历史导出快照没有不可变源包，不能声明精确重跑；请创建新的导出快照。"
        )
    if not snapshot.source_bundle_hash or (
        _source_bundle_hash(source_bundle) != snapshot.source_bundle_hash
    ):
        raise ExportContractError("导出源包校验失败，不能声明精确重跑。")
    return source_bundle


def _source_table_bytes(source_bundle, table):
    encoded = source_bundle["tables"].get(table)
    if not encoded:
        raise ExportContractError(f"冻结源包缺少 {table} 表，不能声明精确重跑。")
    return base64.b64decode(encoded.encode("ascii"))


def _frozen_resource_status(source_bundle):
    resource_manifest = source_bundle.get("resource_manifest")
    if not isinstance(resource_manifest, dict):
        raise ExportContractError("导出源包缺少资源清单，不能声明精确重跑。")
    has_resources = "resources" in source_bundle
    has_resource_hashes = "resource_sha256" in resource_manifest
    if not has_resources and not has_resource_hashes:
        return "legacy_unfrozen"
    resources = source_bundle.get("resources")
    resource_hashes = resource_manifest.get("resource_sha256")
    if (
        resource_manifest.get("resource_freeze_status") != "frozen"
        or not isinstance(resources, dict)
        or not isinstance(resource_hashes, dict)
    ):
        raise ExportContractError("导出源包冻结资源不完整，不能声明精确重跑。")
    for path in FROZEN_RESOURCE_PATHS:
        if not isinstance(resources.get(path), str) or not isinstance(
            resource_hashes.get(path), str
        ):
            raise ExportContractError("导出源包冻结资源不完整，不能声明精确重跑。")
    return "frozen"


def _source_resource_bytes(source_bundle, path):
    if _frozen_resource_status(source_bundle) != "frozen":
        raise ExportContractError("历史导出源包没有冻结资源，不能写入当前资源。")
    try:
        body = base64.b64decode(
            source_bundle["resources"][path].encode("ascii"),
            validate=True,
        )
    except (ValueError, UnicodeEncodeError) as error:
        raise ExportContractError("导出源包冻结资源编码无效，不能声明精确重跑。") from error
    expected_hash = source_bundle["resource_manifest"]["resource_sha256"][path]
    if sha256_bytes(body) != expected_hash:
        raise ExportContractError("导出源包冻结资源校验失败，不能声明精确重跑。")
    return body


def _source_table_path(source_bundle, table):
    paths = source_bundle["resource_manifest"].get("table_paths")
    path = paths.get(table) if isinstance(paths, dict) else None
    if isinstance(path, str) and path:
        return path
    if table == "student_roster" and _frozen_resource_status(source_bundle) == "legacy_unfrozen":
        return IDENTITY_ROSTER_PATH
    raise ExportContractError(f"冻结源包缺少 {table} 表路径，不能声明精确重跑。")


def _frozen_export_manifest(snapshot, source_bundle, *, selected_tables, row_counts,
                            data_status):
    manifest = json.loads(json.dumps(source_bundle["manifest"], ensure_ascii=False))
    manifest["selected_tables"] = list(selected_tables)
    manifest["row_counts"] = row_counts
    manifest["data_status"] = data_status
    config = {
        "selected_tables": list(selected_tables),
        "selected_versions": snapshot.selected_versions,
        "filter_definition": snapshot.filter_definition,
    }
    manifest["export_config_hash"] = hashlib.sha256(
        json.dumps(config, sort_keys=True).encode()
    ).hexdigest()
    manifest["source_bundle_hash"] = snapshot.source_bundle_hash
    manifest["source_bundle_version"] = source_bundle["version"]
    resource_manifest = json.loads(
        json.dumps(source_bundle["resource_manifest"], ensure_ascii=False)
    )
    resource_status = _frozen_resource_status(source_bundle)
    resource_manifest["resource_freeze_status"] = resource_status
    if resource_status == "legacy_unfrozen":
        resource_manifest["unfrozen_resource_paths"] = list(
            FROZEN_RESOURCE_PATHS
        )
    manifest["resource_manifest"] = resource_manifest
    return manifest


def _frozen_derivation_version(manifest):
    """Use the version captured with the source; old source bundles are v1."""
    provenance = manifest.get("provenance") if isinstance(manifest, dict) else None
    derivation = provenance.get("derivation") if isinstance(provenance, dict) else None
    version = derivation.get("version") if isinstance(derivation, dict) else None
    return version or LEGACY_EXPORT_DERIVATION_VERSION


def _selected_export_tables(job):
    if job.kind == ExportJob.Kind.IDENTITY:
        return ["student_roster"]
    if job.selected_tables:
        return list(job.selected_tables)
    return list(PACKAGE_TABLES.get(job.package, PACKAGE_TABLES["complete"]))


def _csv_table_from_bytes(body):
    rows = list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))
    return tuple(rows[0]), rows[1:]


@contextmanager
def _export_file_lock(job_id):
    root = os.path.join(settings.MEDIA_ROOT, 'exports', str(job_id))
    os.makedirs(root, exist_ok=True)
    with open(os.path.join(root, '.generation.lock'), 'a+b') as handle:
        if os.name == 'nt':
            import msvcrt
            if handle.tell() == 0:
                handle.write(b'0')
                handle.flush()
            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_LOCK, 1)
        else:
            import fcntl
            fcntl.flock(handle, fcntl.LOCK_EX)
        try:
            yield
        finally:
            if os.name == 'nt':
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle, fcntl.LOCK_UN)


@contextmanager
def _export_heartbeat(job):
    from django.db import close_old_connections, connections
    stopped, lost = threading.Event(), threading.Event()
    job._lease_lost = lost
    def renew():
        close_old_connections()
        try:
            while not stopped.wait(20):
                try:
                    renewed = renew_export_job_lease(job, worker_id=job.locked_by)
                except Exception:
                    renewed = None
                if renewed is None:
                    lost.set()
                    return
        finally:
            connections.close_all()
    thread = threading.Thread(target=renew, daemon=True)
    thread.start()
    try:
        yield
    finally:
        stopped.set()
        thread.join(timeout=25)


def run_export_job(job):
    from apps.accounts.course_access import lock_course
    inspected = ExportJob.objects.filter(pk=job.pk).first()
    if inspected is None or inspected.status in {ExportJob.Status.INVALIDATED, ExportJob.Status.EXPIRED}:
        raise ExportContractError("导出任务已删除或失效，不能重新生成。")
    with _export_file_lock(job.pk):
        inspected = ExportJob.objects.filter(pk=job.pk).first()
        if inspected is None:
            raise ExportContractError("导出任务已删除或失效，不能重新生成。")
        verified_artifact = (_artifact_signature(inspected)
            if inspected.status == ExportJob.Status.SUCCEEDED and export_artifact_available(inspected) else None)
        with transaction.atomic():
            lock_course(job.course_id)
            current = ExportJob.objects.select_for_update().select_related('snapshot').filter(pk=job.pk).exclude(
                status__in=[ExportJob.Status.INVALIDATED, ExportJob.Status.EXPIRED]).first()
            if current is None or current.snapshot.invalidated_at or current.snapshot.redacted_at:
                raise ExportContractError("导出任务已删除或失效，不能重新生成。")
            if current.expires_at and current.expires_at <= timezone.now():
                raise ExportContractError('导出任务已过期，不能重新生成。')
            if current.status == ExportJob.Status.SUCCEEDED and verified_artifact is not None and verified_artifact == _artifact_signature(current):
                job.refresh_from_db()
                return job
            if current.status == ExportJob.Status.RUNNING and current.locked_by != job.locked_by:
                raise ExportContractError('导出任务租约已转交，不能继续生成。')
            current.status = ExportJob.Status.RUNNING
            current.locked_by = current.locked_by or 'export-' + str(uuid.uuid4())
            current.lease_expires_at = timezone.now() + timedelta(seconds=90)
            current.heartbeat_at = timezone.now()
            current.save(update_fields=['status', 'locked_by', 'lease_expires_at', 'heartbeat_at'])
            job.locked_by = current.locked_by
        with _export_heartbeat(current):
            _materialize_export_job(current)
        job.refresh_from_db()
        return job


def _export_candidate_path(job, root, extension):
    path = os.path.join(root, f"{job.package}-{job.pk}-{uuid.uuid4()}.{extension}")
    job._candidate_path = path
    return path


def _materialize_export_job(job):
    """Materialize an export from an immutable source bundle only."""
    snapshot = job.snapshot
    if snapshot.invalidated_at is not None:
        raise ExportContractError("ExportSnapshot has been invalidated and cannot be rerun.")
    if job.package in {"teaching", "kt", "cd"} and len(snapshot.selected_versions) != 1:
        raise ExportContractError(f"{job.package.upper()} 导出必须选择单一实验版本")
    if job.selected_fields:
        raise ExportContractError("固定研究数据包不支持字段级选择，以免破坏冻结表的引用和回放契约。")
    raw_source_bundle = getattr(snapshot, "source_bundle", None)
    if (
        job.package == "experiment"
        and isinstance(raw_source_bundle, dict)
        and raw_source_bundle.get("version") == 2
        and isinstance(raw_source_bundle.get("research_export_v2"), dict)
    ):
        if _source_bundle_hash(raw_source_bundle) != snapshot.source_bundle_hash:
            raise ExportContractError("导出源包校验失败，不能声明精确重跑。")
        from .research_export_v2 import materialize_research_export

        return _finalize_export_job(
            job, materialize_research_export(job), {}
        )
    source_bundle = _load_source_bundle(snapshot)
    if job.package == "experiment":
        files = source_bundle.get("research_files")
        if not files:
            raise ExportContractError("冻结源中没有实验研究数据包内容")
        root = os.path.join(settings.MEDIA_ROOT, "exports", str(job.job_id))
        os.makedirs(root, exist_ok=True)
        artifact_path = _export_candidate_path(job, root, "zip")
        with zipfile.ZipFile(artifact_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(files):
                archive.writestr(path, base64.b64decode(files[path].encode("ascii")))
        return _finalize_export_job(job, artifact_path, {})
    resource_freeze_status = _frozen_resource_status(source_bundle)
    selected_tables = _selected_export_tables(job)
    if job.package == "teaching" and job.selected_tables and selected_tables != ["attempts"]:
        raise ExportContractError("教学包只支持固定的 attempts 成绩投影，不能选择其他表。")
    resource_manifest = source_bundle["resource_manifest"]
    row_counts = {
        table: (
            0
            if table == "student_roster"
            else resource_manifest["row_counts"].get(table, 0)
        )
        for table in selected_tables
        if table == "student_roster" or table in source_bundle["tables"]
    }
    if job.kind == ExportJob.Kind.IDENTITY:
        row_counts["student_roster"] = len(
            _csv_table_from_bytes(
                base64.b64decode(source_bundle["student_roster_csv"].encode("ascii"))
            )[1]
        )
    data_status = {
        table: "no_observation" if row_counts.get(table, 0) == 0 else "observed"
        for table in selected_tables
    }
    manifest = _frozen_export_manifest(
        snapshot,
        source_bundle,
        selected_tables=selected_tables,
        row_counts=row_counts,
        data_status=data_status,
    )
    derivation_version = _frozen_derivation_version(manifest)
    root = os.path.join(settings.MEDIA_ROOT, "exports", str(job.job_id))
    os.makedirs(root, exist_ok=True)
    artifact_extension = "xlsx" if job.package == "teaching" else "zip"
    artifact_path = _export_candidate_path(job, root, artifact_extension)
    if job.package == "teaching":
        _write_teaching_workbook(
            artifact_path,
            _csv_table_from_bytes(_source_table_bytes(source_bundle, "attempts")),
            manifest,
            (),
            events=(),
            submissions=(),
            timing_by_attempt=source_bundle["teaching_timing"],
        )
        return _finalize_export_job(job, artifact_path, manifest)
    with zipfile.ZipFile(artifact_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("00_manifest/manifest.json", manifest_bytes(manifest))
        archive.writestr("00_manifest/derivation_runs.json", manifest_bytes({
            "derivation_version": derivation_version,
            "deterministic": True,
            "source_snapshot_id": str(snapshot.snapshot_id),
            "source_cutoff_at": snapshot.source_cutoff_at.isoformat(),
            "event_high_watermark": snapshot.event_high_watermark,
            "schema_manifest_hash": snapshot.schema_manifest_hash,
            "application_git_commit_sha": manifest["application_git_commit_sha"],
            "source_bundle_hash": snapshot.source_bundle_hash,
        }))
        if resource_freeze_status == "frozen":
            archive.writestr(
                FIELD_DICTIONARY_PATH,
                _source_resource_bytes(source_bundle, FIELD_DICTIONARY_PATH),
            )
        archive.writestr("01_design_metadata/README.txt", "This package contains direct student identifiers and must be handled as identified research data.\n")
        if job.kind == ExportJob.Kind.IDENTITY:
            archive.writestr(
                _source_table_path(source_bundle, "student_roster"),
                base64.b64decode(source_bundle["student_roster_csv"].encode("ascii")),
            )
            return _finalize_export_job(job, artifact_path, manifest)
        for table in selected_tables:
            if table == "identity_mapping" or table not in source_bundle["tables"]:
                continue
            archive.writestr(
                _source_table_path(source_bundle, table),
                _source_table_bytes(source_bundle, table),
            )
        replay_validator_tables = {
            "assessment_events", "attempts", "response_snapshots",
            "response_trajectory", "response_replay", "session_summary",
        }
        if (
            resource_freeze_status == "frozen"
            and job.package == "complete"
            and replay_validator_tables.issubset(selected_tables)
        ):
            archive.writestr(
                REPLAY_SCRIPT_PATH,
                _source_resource_bytes(source_bundle, REPLAY_SCRIPT_PATH),
            )
    return _finalize_export_job(job, artifact_path, manifest)


def _write_teaching_workbook(
    artifact_path, attempt_table, manifest, attempts, *, events, submissions,
    timing_by_attempt=None,
):
    """Write the teaching-facing grade workbook from the same frozen rows.

    The source is the canonical attempts export table, whose grade columns are
    already restricted to the effective GradeRecord.  The second sheet keeps
    the snapshot provenance with the downloadable workbook rather than
    requiring a separate, mutable dashboard lookup.
    """
    headers, rows = attempt_table
    requested = (
        "student_number", "student_name", "question_code", "session_id",
        "status", "grade_score", "grade_max_score", "grade_is_correct",
        "submission_grade_status", "graded_at",
    )
    positions = [headers.index(column) for column in requested if column in headers]
    attempt_id_position = headers.index("attempt_id")
    if timing_by_attempt is None:
        events_by_attempt = {}
        for event in events:
            events_by_attempt.setdefault(event.attempt_id, []).append(event)
        submissions_by_attempt = {
            submission.attempt_id: submission
            for submission in submissions
        }
        from .teaching_data import summarize_attempt_timing

        timing_by_attempt = {
            attempt.pk: summarize_attempt_timing(
                attempt,
                events=events_by_attempt.get(attempt.pk, []),
                submission=submissions_by_attempt.get(attempt.pk),
            )
            for attempt in attempts
        }
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("教学成绩")
    sheet.append([
        *[headers[position] for position in positions],
        "display_duration_ms", "display_kind", "interval_quality",
    ])
    for row in rows:
        timing = (
            timing_by_attempt.get(row[attempt_id_position])
            or timing_by_attempt.get(str(row[attempt_id_position]))
        )
        if isinstance(timing, dict):
            duration = timing.get("display_duration_ms")
            kind = timing.get("display_kind", "unavailable")
            quality = timing.get("interval_quality", "unknown")
        else:
            duration = timing.display_duration_ms if timing else None
            kind = timing.display_kind if timing else "unavailable"
            quality = timing.interval_quality if timing else "unknown"
        sheet.append([
            *[row[position] for position in positions],
            duration,
            kind,
            quality,
        ])
    lineage = workbook.create_sheet("导出血缘")
    lineage.append(["field", "value"])
    for key in (
        "export_snapshot_id", "source_cutoff_at", "event_high_watermark",
        "schema_manifest_hash", "scope_hash", "contains_direct_identifiers",
    ):
        lineage.append([key, json.dumps(manifest.get(key), ensure_ascii=False)])
    workbook.save(artifact_path)


def _validate_question_codes(assignments):
    seen = {}
    for assignment in assignments:
        snapshot = assignment.snapshot
        code = snapshot.code.strip()
        if not code:
            raise ExportContractError(f"题目快照 {snapshot.pk} 缺少 question_code")
        # Bank-local codes are legal in a multi-bank experiment. Frozen snapshot
        # IDs remain the internal joins; standard datasets expose course sequence.
        key = (snapshot.version_id, snapshot.source_bank_id, code)
        previous_snapshot_id = seen.setdefault(key, snapshot.pk)
        if previous_snapshot_id != snapshot.pk:
            raise ExportContractError(f"同一实验版本题目编码重复：{code}")


def preview_export_table(*, course, experiment, table, version_id=None,
                         class_name=None, student_number=None, session_id=None, status=None,
                         page=1, page_size=20, scope=None, student_search=""):
    """Return a paginated export preview.

    ``scope`` is the data-center path and is deliberately preferred over the
    legacy scalar parameters.  Keeping the preview on the same ``DataScope``
    as teaching, quality, and ZIP export prevents a table view from showing a
    wider population than the visible range bar claims.
    """
    student_search = str(student_search or "").strip()[:100]
    if scope is not None:
        if scope.course.pk != course.pk:
            raise ExportContractError("预览范围不属于当前课程")
        if experiment is not None and scope.experiment_id not in {None, experiment.pk}:
            raise ExportContractError("预览范围与实验不一致")
        filter_definition = scope.serialize()
        session_queryset = scope.sessions()
    else:
        filter_definition = {
            key: value
            for key, value in {
                "version_id": version_id,
                "class_name": class_name,
                "student_number": student_number,
                "session_id": session_id,
                "status": status,
            }.items()
            if value not in (None, "")
        }
        session_filter = {"version__experiment_id": experiment.pk}
        if version_id:
            session_filter["version_id"] = int(version_id)
        if class_name:
            session_filter["class_snapshot_name"] = _class_snapshot_filter_value(class_name)
        if student_number:
            session_filter["student__student_number"] = student_number
        if session_id:
            session_filter["pk"] = int(session_id)
        if status in {choice for choice, _label in ExperimentSession.Status.choices}:
            session_filter["status"] = status
        session_queryset = ExperimentSession.objects.filter(**session_filter)
    if student_search:
        filter_definition["student_search"] = student_search
        session_queryset = session_queryset.filter(
            models.Q(student__display_name__icontains=student_search)
            | models.Q(student__student_number__icontains=student_search)
            | models.Q(student__username__icontains=student_search)
        )
    sessions = list(
        session_queryset.select_related("student", "version").order_by("pk")
    )
    session_ids = [session.pk for session in sessions]
    attempt_queryset = (
        scope.attempts() if scope is not None else Attempt.objects.all()
    ).filter(session_id__in=session_ids).select_related(
        "session", "assignment__snapshot", "assignment__item_instance", "submission_event"
    )
    has_attempt_status = (
        scope is not None and scope.attempt_status is not None
    ) or (scope is None and status in {choice for choice, _label in Attempt.Status.choices})
    if scope is None and status in {choice for choice, _label in Attempt.Status.choices}:
        attempt_queryset = attempt_queryset.filter(status=status)
    attempts = list(attempt_queryset.order_by("pk"))
    if has_attempt_status:
        selected_session_ids = {attempt.session_id for attempt in attempts}
        sessions = [session for session in sessions if session.pk in selected_session_ids]
        session_ids = [session.pk for session in sessions]
    attempt_ids = {attempt.pk for attempt in attempts}
    assignments = list(
        StudentQuestionAssignment.objects.filter(session_id__in=session_ids)
        .select_related("snapshot", "item_instance")
        .order_by("session_id", "assigned_sequence_index")
    )
    if has_attempt_status:
        assignment_ids = {attempt.assignment_id for attempt in attempts}
        assignments = [assignment for assignment in assignments if assignment.pk in assignment_ids]
    events = list(
        (scope.events() if scope is not None else AssessmentEvent.objects.all()).filter(
            session_id__in=session_ids,
            server_received_at__lte=timezone.now(),
        ).order_by("session_id", "server_ingest_sequence_no")
    )
    confidence_assignments = list(ConfidenceAssignment.objects.filter(attempt_id__in=attempt_ids))
    simulations = list(SimulationSnapshot.objects.filter(attempt_id__in=attempt_ids))
    response_snapshots = list(
        ResponseSnapshot.objects.filter(attempt_id__in=attempt_ids).order_by(
            "attempt_id", "created_at", "snapshot_id"
        )
    )
    input_observations = list(
        ResponseInputObservation.objects.filter(
            attempt_id__in=attempt_ids
        ).select_related("source_event").order_by(
            "attempt_id", "client_instance_id", "client_sequence_no"
        )
    )
    submissions = list(SubmissionEvent.objects.filter(attempt_id__in=attempt_ids))
    grade_records = list(
        GradeRecord.objects.filter(submission__attempt_id__in=attempt_ids).select_related("submission")
    )
    version_ids = {session.version_id for session in sessions}
    class_snapshots = ExperimentClassSnapshot.objects.filter(version_id__in=version_ids)
    if scope is not None and scope.class_snapshot_id is not None:
        class_snapshots = class_snapshots.filter(pk=scope.class_snapshot_id)
    class_snapshots = list(class_snapshots.order_by("version_id", "created_at"))
    participant_snapshots = list(
        ExperimentVersionParticipantSnapshot.objects.filter(
            version_id__in=version_ids,
            student_id__in=[session.student_id for session in sessions],
        ).order_by("version_id", "student_id")
    )
    condition_assignments = list(
        ConditionAssignment.objects.filter(
            models.Q(session_id__in=session_ids)
            | models.Q(
                subject_id__in=[session.student_id for session in sessions],
                condition__experiment_id__in=[session.version.experiment_id for session in sessions],
            )
        ).select_related("condition")
    )
    fake_job = SimpleNamespace(course_id=course.pk, research_study_id=None)
    rows = _table_rows(
        fake_job,
        sessions,
        assignments,
        attempts,
        events,
        confidence_assignments,
        simulations,
        response_snapshots=response_snapshots,
        input_observations=input_observations,
        submissions=submissions,
        grade_records=grade_records,
        condition_assignments=condition_assignments,
        class_snapshots=class_snapshots,
        participant_snapshots=participant_snapshots,
    )
    headers, values = rows.get(table, ((), []))
    total = len(values)
    try:
        page = max(1, int(page or 1))
    except (TypeError, ValueError):
        page = 1
    try:
        page_size = max(1, min(100, int(page_size or 20)))
    except (TypeError, ValueError):
        page_size = 20
    start = (page - 1) * page_size
    return {
        "table": table,
        "headers": headers,
        "rows": values[start:start + page_size],
        "total": total,
        "page": page,
        "page_size": page_size,
        "has_previous": page > 1,
        "has_next": start + page_size < total,
        "previous_page": max(1, page - 1),
        "next_page": page + 1,
        "filter_definition": filter_definition,
        "data_status": "placeholder" if table in PLACEHOLDER_TABLES else ("no_observation" if total == 0 else "observed"),
    }


def _table_rows(job, sessions, assignments, attempts, events, confidence_assignments, simulations,
                *, response_snapshots=(), input_observations=(), submissions=(),
                grade_records=(), condition_assignments=(), class_snapshots=(),
                participant_snapshots=(), snapshot=None, session_projection=False):
    identity_headers = ("student_number", "student_name")
    session_by_id = {session.pk: session for session in sessions}
    attempt_by_assignment = {attempt.assignment_id: attempt for attempt in attempts}
    attempt_by_id = {attempt.pk: attempt for attempt in attempts}
    snapshot_code_by_id = {assignment.snapshot_id: assignment.snapshot.code for assignment in assignments}

    def identity_for_session(session_id):
        session = session_by_id[session_id]
        return [session.student.student_number or "", session.student.shown_name]

    def question_code_for_snapshot(snapshot_id):
        return snapshot_code_by_id.get(snapshot_id, "")

    subject_values = []
    seen_student_numbers = set()
    for session in sessions:
        student_number = session.student.student_number or ""
        if student_number not in seen_student_numbers:
            subject_values.append([student_number, session.student.shown_name])
            seen_student_numbers.add(student_number)

    selected_attempt_ids = set(attempt_by_id)
    selected_events = [event for event in events if event.attempt_id is None or event.attempt_id in selected_attempt_ids]
    selected_confidence = [assignment for assignment in confidence_assignments if assignment.attempt_id in selected_attempt_ids]
    selected_simulations = [simulation for simulation in simulations if simulation.attempt_id in selected_attempt_ids]

    cutoff = snapshot.source_cutoff_at if snapshot is not None else timezone.now()
    capture_contracts = list(
        CaptureContractArtifact.objects.filter(
            capture_contract_id__in={
                session.capture_contract_id
                for session in sessions
                if session.capture_contract_id
            }
        ).order_by("contract_name")
    )
    time_syncs = list(
        ClientTimeSync.objects.filter(
            session_id__in=session_by_id,
            created_at__lte=cutoff,
        ).order_by("session_id", "client_instance_id", "created_at")
    )
    valid_syncs_by_instance = {}
    for sync in time_syncs:
        if (
            sync.sync_quality == ClientTimeSync.SyncQuality.VALID
            and sync.estimated_clock_offset_ms is not None
        ):
            valid_syncs_by_instance.setdefault(
                (str(sync.client_instance_id), sync.client_time_origin), []
            ).append(sync)

    def derived_event_time(event):
        if event.estimated_occurred_at_utc is not None:
            return event.estimated_occurred_at_utc
        if (
            event.client_instance_id is None
            or event.client_time_origin is None
            or event.client_monotonic_ms is None
        ):
            return None
        candidates = valid_syncs_by_instance.get(
            (str(event.client_instance_id), event.client_time_origin), []
        )
        if not candidates:
            return None
        sample = min(
            candidates,
            key=lambda row: abs(row.client_monotonic_ms - event.client_monotonic_ms),
        )
        delta_ms = event.client_monotonic_ms - sample.client_monotonic_ms
        if abs(delta_ms) > 20 * 60 * 1000:
            return None
        return (
            sample.client_wall_time_utc
            + timedelta(milliseconds=sample.estimated_clock_offset_ms)
            + timedelta(milliseconds=float(delta_ms))
        )
    checkpoints = list(
        ClientCaptureCheckpoint.objects.filter(
            session_id__in=session_by_id,
            server_received_at__lte=cutoff,
        ).order_by("session_id", "client_instance_id", "server_received_at")
    )
    coverage_sessions = list(
        RecordingCoverageSession.objects.filter(
            session_id__in=session_by_id,
            run__run_status="completed",
            run__source_cutoff_at__lte=cutoff,
        ).filter(
            **({"run_id": snapshot.coverage_run_id} if getattr(snapshot, "coverage_run_id", None) else {})
        ).select_related("run").order_by("session_id", "-run__generated_at")
    )
    coverage_attempts = list(
        RecordingCoverageAttempt.objects.filter(
            attempt_id__in=selected_attempt_ids,
            run__run_status="completed",
            run__source_cutoff_at__lte=cutoff,
        ).filter(
            **({"run_id": snapshot.coverage_run_id} if getattr(snapshot, "coverage_run_id", None) else {})
        ).select_related("run", "attempt").order_by("attempt_id", "-run__generated_at")
    )
    latest_coverage_session = {}
    for result in coverage_sessions:
        latest_coverage_session.setdefault(result.session_id, result)
    latest_coverage_attempt = {}
    for result in coverage_attempts:
        latest_coverage_attempt.setdefault(result.attempt_id, result)
    coverage_runs = {
        str(result.run_id): result.run
        for result in [
            *latest_coverage_session.values(),
            *latest_coverage_attempt.values(),
        ]
    }

    event_headers = (
        *identity_headers, "student_id", "question_code", "event_id", "event_type", "event_version", "recovery_round", "session_id", "attempt_id",
        "client_instance_id", "client_sequence_no", "server_ingest_sequence_no", "correlation_id",
        "server_received_at", "capture_mode", "quality_flags", "payload", "event_domain",
        "actor_source", "envelope_version", "assignment_id", "question_snapshot_id", "item_instance_id",
        "client_wall_time_utc", "client_monotonic_ms", "client_time_origin", "time_sync_id",
        "estimated_clock_offset_ms", "estimated_rtt_ms", "estimated_occurred_at_utc", "derived_occurred_at_utc", "server_committed_at",
        "app_version", "server_clock_reversed", "server_processing_elapsed_ms",
    )

    def event_row(event):
        return [
            *identity_for_session(event.session_id), event.student_id,
            question_code_for_snapshot(event.question_snapshot_id),
            event.event_id, event.event_type, event.event_version, event.recovery_round, event.session_id, event.attempt_id,
            event.client_instance_id, event.client_sequence_no, event.server_ingest_sequence_no,
            event.correlation_id, event.server_received_at.isoformat(), event.capture_mode,
            json.dumps(event.quality_flags, ensure_ascii=False), json.dumps(event.payload, ensure_ascii=False),
            event.event_domain, event.actor_source, event.envelope_version, event.assignment_id,
            event.question_snapshot_id, event.item_instance_id,
            event.client_wall_time_utc.isoformat() if event.client_wall_time_utc else "",
            event.client_monotonic_ms, event.client_time_origin.isoformat() if event.client_time_origin else "",
            event.time_sync_id, event.estimated_clock_offset_ms, event.estimated_rtt_ms,
            event.estimated_occurred_at_utc.isoformat() if event.estimated_occurred_at_utc else "",
            derived_event_time(event).isoformat() if derived_event_time(event) else "",
            event.server_committed_at.isoformat() if event.server_committed_at else "", event.app_version,
            event.server_clock_reversed, event.server_processing_elapsed_ms,
        ]

    event_rows = [event_row(event) for event in selected_events]
    compact_event_headers = (*identity_headers, "question_code", "event_id", "event_type", "session_id", "attempt_id", "server_received_at", "payload")

    def compact_event_row(event):
        return [
            *identity_for_session(event.session_id), question_code_for_snapshot(event.question_snapshot_id),
            event.event_id, event.event_type, event.session_id, event.attempt_id,
            event.server_received_at.isoformat(), json.dumps(event.payload, ensure_ascii=False),
        ]

    projection_headers = (*event_headers, "time_quality", "coverage_run_id")

    def projection_row(event):
        coverage = latest_coverage_attempt.get(event.attempt_id) or latest_coverage_session.get(event.session_id)
        return [
            *event_row(event),
            getattr(coverage, "time_quality", None) or event.interval_quality,
            coverage.run_id if coverage else "",
        ]

    response_headers = (
        *projection_headers, "before_state", "after_state", "before_state_hash",
        "after_state_hash", "response_snapshot_id", "question_type", "change_source",
    )
    response_rows = [
        [
            *projection_row(event),
            json.dumps(
                event.payload.get("response_state_before")
                if "response_state_before" in event.payload
                else event.payload.get("before_state"),
                ensure_ascii=False,
            ),
            json.dumps(
                event.payload.get("response_state_after")
                if "response_state_after" in event.payload
                else event.payload.get("after_state"),
                ensure_ascii=False,
            ),
            event.payload.get("before_state_hash", ""),
            event.payload.get("after_state_hash") or event.payload.get("state_hash", ""),
            event.payload.get("response_snapshot_id")
            or event.payload.get("after_snapshot_id", ""),
            event.payload.get("question_type")
            or event.payload.get("after_state", {}).get("question_type", ""),
            event.payload.get("change_source") or event.payload.get("source", ""),
        ]
        for event in selected_events if event.event_type in RESPONSE_EVENT_TYPES
    ]
    exposure_rows = [projection_row(event) for event in selected_events if event.event_type in EXPOSURE_EVENT_TYPES]
    navigation_headers = (
        *projection_headers, "from_assignment_id", "to_assignment_id",
        "entry_reason", "leave_reason",
    )
    navigation_rows = [
        [
            *projection_row(event), event.payload.get("from_assignment_id", ""),
            event.payload.get("to_assignment_id", ""), event.payload.get("entry_reason", ""),
            event.payload.get("leave_reason", ""),
        ]
        for event in selected_events
        if event.event_type in ("item_entered", "item_rendered", "item_left")
    ]
    telemetry_rows = [projection_row(event) for event in selected_events if event.event_domain == AssessmentEvent.EventDomain.TELEMETRY]
    simulation_event_rows = [compact_event_row(event) for event in selected_events if event.event_type in SIMULATION_EVENT_TYPES]

    def identity_and_question_for_attempt(attempt_id):
        attempt = attempt_by_id[attempt_id]
        return [*identity_for_session(attempt.session_id), attempt.assignment.snapshot.code]

    def export_time(value):
        return value.isoformat() if value else ""

    transition_events = [
        event for event in selected_events
        if event.event_type == "response_transition" and event.attempt_id in attempt_by_id
    ]
    transition_by_event_id = {
        str(event.event_id): event
        for event in transition_events
    }

    def transition_record(event):
        payload = event.payload or {}
        return {
            "event_id": str(event.event_id),
            "parent_event_id": str(payload["parent_event_id"])
            if payload.get("parent_event_id") else None,
            "branch_id": payload.get("branch_id"),
            "operation_id": payload.get("operation_id"),
            "server_ingest_sequence_no": event.server_ingest_sequence_no,
            "server_received_at": export_time(event.server_received_at),
            "before_snapshot_id": payload.get("before_snapshot_id"),
            "after_snapshot_id": payload.get("after_snapshot_id"),
            "before_state": payload.get("before_state"),
            "after_state": payload.get("after_state"),
            "before_state_hash": payload.get("before_state_hash"),
            "after_state_hash": payload.get("after_state_hash"),
        }

    def replay_chain(event_id):
        chain = []
        seen = set()
        current_event_id = str(event_id) if event_id else None
        if current_event_id is None:
            return chain, "not_selected"
        while current_event_id:
            if current_event_id in seen:
                return list(reversed(chain)), "cycle_detected"
            seen.add(current_event_id)
            event = transition_by_event_id.get(current_event_id)
            if event is None:
                return list(reversed(chain)), "missing_transition"
            record = transition_record(event)
            chain.append(record)
            current_event_id = record["parent_event_id"]
        return list(reversed(chain)), "complete"

    response_snapshot_by_id = {
        str(item.snapshot_id): item
        for item in response_snapshots
        if item.attempt_id in attempt_by_id
    }

    def response_state_for(event_id, snapshot_id):
        event = transition_by_event_id.get(str(event_id)) if event_id else None
        if event is not None:
            payload = event.payload or {}
            return payload.get("after_state"), payload.get("after_state_hash") or ""
        response_snapshot = response_snapshot_by_id.get(str(snapshot_id)) if snapshot_id else None
        if response_snapshot is not None:
            return response_snapshot.state_json, response_snapshot.state_hash or ""
        return None, ""

    branch_resolutions_by_attempt = {}
    for event in selected_events:
        if event.event_type != "response_branch_resolved" or event.attempt_id not in attempt_by_id:
            continue
        payload = event.payload or {}
        branch_resolutions_by_attempt.setdefault(event.attempt_id, []).append({
            "event_id": str(event.event_id),
            "operation_id": payload.get("operation_id"),
            "candidate_event_ids": [str(item) for item in payload.get("candidate_event_ids", [])],
            "selected_event_id": str(payload["selected_event_id"])
            if payload.get("selected_event_id") else None,
            "expected_version": payload.get("expected_version"),
            "server_ingest_sequence_no": event.server_ingest_sequence_no,
            "server_received_at": export_time(event.server_received_at),
        })

    response_trajectory_headers = (
        *identity_headers, "student_id", "question_code", "session_id", "attempt_id", "event_id",
        "event_version", "server_ingest_sequence_no", "server_received_at",
        "server_committed_at", "client_instance_id", "client_sequence_no",
        "capture_version", "operation_id", "branch_id", "parent_event_id",
        "before_snapshot_id", "after_snapshot_id", "before_state", "after_state",
        "before_state_hash", "after_state_hash", "input_type", "source",
        "option_key", "option_action", "capture_mode", "quality_flags",
    )
    response_trajectory_rows = [
        [
            *identity_for_session(event.session_id), event.student_id,
            question_code_for_snapshot(event.question_snapshot_id),
            event.session_id, event.attempt_id, event.event_id, event.event_version,
            event.server_ingest_sequence_no, export_time(event.server_received_at),
            export_time(event.server_committed_at), event.client_instance_id or "",
            event.client_sequence_no or "", event.payload.get("capture_version", ""),
            event.payload.get("operation_id", ""), event.payload.get("branch_id", ""),
            event.payload.get("parent_event_id") or "",
            event.payload.get("before_snapshot_id", ""),
            event.payload.get("after_snapshot_id", ""),
            json.dumps(event.payload.get("before_state"), ensure_ascii=False),
            json.dumps(event.payload.get("after_state"), ensure_ascii=False),
            event.payload.get("before_state_hash", ""),
            event.payload.get("after_state_hash", ""),
            event.payload.get("input_type", ""), event.payload.get("source", ""),
            event.payload.get("option_key", ""), event.payload.get("option_action", ""),
            event.capture_mode, json.dumps(event.quality_flags, ensure_ascii=False),
        ]
        for event in transition_events
    ]

    response_replay_headers = (
        *identity_headers, "student_id", "question_code", "session_id", "attempt_id",
        "capture_contract_id", "active_response_event_id", "final_response_event_id",
        "response_version", "formal_capture_boundary", "current_response_snapshot_id",
        "final_response_snapshot_id", "active_state", "active_state_hash",
        "final_state", "final_state_hash", "active_replay_status",
        "active_replay_chain", "final_replay_status", "final_replay_chain",
        "branch_resolutions",
    )
    response_replay_rows = []
    for attempt in attempts:
        active_chain, active_replay_status = replay_chain(attempt.active_response_event_id)
        final_chain, final_replay_status = replay_chain(attempt.final_response_event_id)
        active_state, active_state_hash = response_state_for(
            attempt.active_response_event_id,
            attempt.current_response_snapshot_id,
        )
        final_state, final_state_hash = response_state_for(
            attempt.final_response_event_id,
            attempt.final_response_snapshot_id,
        )
        response_replay_rows.append([
            *identity_for_session(attempt.session_id), attempt.session.student_id,
            attempt.assignment.snapshot.code, attempt.session_id, attempt.pk,
            attempt.session.capture_contract_id or "", attempt.active_response_event_id or "",
            attempt.final_response_event_id or "", attempt.response_version,
            json.dumps(attempt.formal_capture_boundary or {}, ensure_ascii=False),
            attempt.current_response_snapshot_id or "", attempt.final_response_snapshot_id or "",
            json.dumps(active_state, ensure_ascii=False) if active_state is not None else "",
            active_state_hash,
            json.dumps(final_state, ensure_ascii=False) if final_state is not None else "",
            final_state_hash, active_replay_status,
            json.dumps(active_chain, ensure_ascii=False), final_replay_status,
            json.dumps(final_chain, ensure_ascii=False),
            json.dumps(branch_resolutions_by_attempt.get(attempt.pk, []), ensure_ascii=False),
        ])

    attempts_by_session = {}
    for attempt in attempts:
        attempts_by_session.setdefault(attempt.session_id, []).append(attempt)
    event_counts_by_session = {}
    for event in selected_events:
        counts = event_counts_by_session.setdefault(event.session_id, {
            "response_transition": 0,
            "input_operation": 0,
            "response_branch_resolved": 0,
            "maximum_server_ingest_sequence_no": 0,
        })
        if event.event_type in counts:
            counts[event.event_type] += 1
        counts["maximum_server_ingest_sequence_no"] = max(
            counts["maximum_server_ingest_sequence_no"], event.server_ingest_sequence_no,
        )
    session_summary_headers = (
        *identity_headers, "session_id", "experiment_id", "version", "status",
        "capture_contract_id", "attempt_count", "response_transition_count",
        "input_operation_count", "response_branch_resolved_count",
        "active_response_attempt_count", "final_response_attempt_count",
        "formal_capture_boundary_count", "maximum_server_ingest_sequence_no",
    )
    session_summary_rows = []
    for session in sessions:
        session_attempts = attempts_by_session.get(session.pk, [])
        event_counts = event_counts_by_session.get(session.pk, {})
        session_summary_rows.append([
            *identity_for_session(session.pk), session.pk, session.version.experiment_id,
            session.version.version_number, session.status, session.capture_contract_id or "",
            len(session_attempts), event_counts.get("response_transition", 0),
            event_counts.get("input_operation", 0),
            event_counts.get("response_branch_resolved", 0),
            sum(bool(item.active_response_event_id) for item in session_attempts),
            sum(bool(item.final_response_event_id) for item in session_attempts),
            sum(bool(item.formal_capture_boundary) for item in session_attempts),
            event_counts.get("maximum_server_ingest_sequence_no", 0),
        ])

    def input_observation_row(item):
        event = item.source_event if item.source_event_id else None
        coverage = latest_coverage_attempt.get(item.attempt_id)
        derived_time = derived_event_time(event) if event is not None else None
        return [
            *identity_and_question_for_attempt(item.attempt_id),
            item.observation_id,
            item.source_event_id or "",
            event.event_type if event else "",
            event.event_version if event else "",
            event.envelope_version if event else "",
            item.attempt_id,
            item.client_instance_id,
            item.client_sequence_no,
            item.client_monotonic_ms,
            item.batch_start_client_monotonic_ms,
            item.batch_end_client_monotonic_ms,
            event.client_wall_time_utc.isoformat()
            if event and event.client_wall_time_utc else "",
            event.client_time_origin.isoformat()
            if event and event.client_time_origin else "",
            event.time_sync_id if event else "",
            event.estimated_occurred_at_utc.isoformat()
            if event and event.estimated_occurred_at_utc else "",
            derived_time.isoformat() if derived_time else "",
            event.server_received_at.isoformat() if event else "",
            event.server_committed_at.isoformat()
            if event and event.server_committed_at else "",
            event.server_ingest_sequence_no if event else "",
            event.actor_source if event else "",
            event.capture_mode if event else "",
            json.dumps(event.quality_flags, ensure_ascii=False) if event else "",
            (coverage.time_quality if coverage else event.interval_quality)
            if event else "unknown",
            coverage.run_id if coverage else "",
            event.app_version if event else "",
            item.edit_operation,
            item.content_hash_before,
            item.content_hash_after,
            item.content_length_before,
            item.content_length_after,
            item.input_method,
            item.is_ime,
            item.observed_at.isoformat(),
        ]

    grade_by_attempt = {}
    if not grade_records and selected_attempt_ids:
        grade_records = GradeRecord.objects.filter(
            submission__attempt_id__in=selected_attempt_ids,
            created_at__lte=(snapshot.source_cutoff_at if snapshot is not None else timezone.now()),
        ).select_related("submission")
    for grade in grade_records:
        if grade.validity_status == GradeRecord.ValidityStatus.EFFECTIVE:
            grade_by_attempt[grade.submission.attempt_id] = grade
    submission_by_attempt = {submission.attempt_id: submission for submission in submissions}
    task_ids = set() if session_projection else {
        session.class_task_id for session in sessions if session.class_task_id
    }
    if not session_projection:
        task_ids.update(
            participant.class_task_id
            for participant in participant_snapshots
            if participant.class_task_id
        )
    class_tasks = [] if session_projection else list(
        ClassExperimentTask.objects.filter(pk__in=task_ids)
        .select_related(
            "class_group",
            "release_snapshot",
            "randomization_identity",
            "preview_class_result__preview_run",
        )
        .prefetch_related("manual_closures")
        .order_by("pk")
    )
    schedule_revisions = [] if session_projection else list(
        ScheduleRevision.objects.filter(class_task_id__in=task_ids)
        .select_related("class_task")
        .order_by("class_task_id", "revision_no")
    )
    release_ids = {session.version_id for session in sessions} | {
        task.release_snapshot_id for task in class_tasks
    } | {
        participant.version_id for participant in participant_snapshots
    } | {
        class_snapshot.version_id for class_snapshot in class_snapshots
    }
    releases = [] if session_projection else list(
        ExperimentVersion.objects.filter(pk__in=release_ids)
        .select_related("experiment")
        .order_by("experiment_id", "version_number")
    )
    release_mappings = [] if session_projection else list(
        ReleaseSnapshotQuestionMapping.objects.filter(release_snapshot_id__in=release_ids)
        .select_related("question_snapshot")
        .order_by("release_snapshot_id", "source_question_id")
    )
    metadata_snapshots_by_id = {
        mapping.question_snapshot_id: mapping.question_snapshot
        for mapping in release_mappings
    }
    metadata_snapshots_by_id.update(
        {assignment.snapshot_id: assignment.snapshot for assignment in assignments}
    )
    participant_snapshot_question_ids = set() if session_projection else {
        item.get("question_snapshot_id")
        for participant in participant_snapshots
        for item in participant.assignment_snapshot
        if item.get("question_snapshot_id")
    }
    missing_snapshot_ids = (
        participant_snapshot_question_ids - metadata_snapshots_by_id.keys()
    )
    if missing_snapshot_ids:
        metadata_snapshots_by_id.update(
            {
                item.pk: item
                for item in QuestionSnapshot.objects.filter(
                    pk__in=missing_snapshot_ids
                )
            }
        )
    metadata_snapshots = [
        metadata_snapshots_by_id[snapshot_id]
        for snapshot_id in sorted(metadata_snapshots_by_id)
    ]
    kc_dictionary_by_id = {}
    for item in metadata_snapshots:
        frozen_concepts = item.all_kc_snapshots or []
        if not frozen_concepts and item.primary_kc:
            frozen_concepts = [
                {
                    "concept_id": item.primary_kc_id_snapshot,
                    "name": item.primary_kc,
                    "path": item.primary_kc_path,
                }
            ]
        for concept in frozen_concepts:
            concept_id = concept.get("concept_id")
            if concept_id is not None:
                kc_dictionary_by_id[concept_id] = [
                    concept_id,
                    concept.get("name", ""),
                    concept.get("path", ""),
                ]
    kc_dictionary_rows = [
        kc_dictionary_by_id[concept_id]
        for concept_id in sorted(kc_dictionary_by_id)
    ]

    def primary_level_name(item, index):
        levels = item.primary_kc_levels or []
        return levels[index].get("name", "") if len(levels) > index else ""

    v2_snapshots = [
        item
        for item in metadata_snapshots
        if item.kc_mapping_schema_version == "question-kc-tree-v2"
    ]
    frozen_kcs = {}
    structure_v2_rows = []
    for item in v2_snapshots:
        for node in item.kc_mapping_snapshot.get("nodes", []):
            frozen_kcs.setdefault(
                node["kc_code"],
                [
                    node["kc_code"],
                    node["name"],
                    node.get("description", ""),
                    node.get("status", "active"),
                ],
            )
            structure_v2_rows.append(
                [
                    item.code,
                    item.stable_code,
                    node["kc_code"],
                    node["depth"],
                    node.get("parent_kc_code") or "",
                    node["order"],
                    item.kc_mapping_hash,
                    item.kc_mapping_schema_version,
                ]
            )
    q_matrix_codes = sorted(frozen_kcs)
    q_matrix_v2_rows = []
    for item in v2_snapshots:
        mapped = {
            node["kc_code"]
            for node in item.kc_mapping_snapshot.get("nodes", [])
        }
        q_matrix_v2_rows.append(
            [item.code, item.stable_code, item.kc_mapping_hash]
            + [1 if code in mapped else 0 for code in q_matrix_codes]
        )
    snapshot_by_id = {item.pk: item for item in metadata_snapshots}
    attribution_v2_rows = []
    for participant in participant_snapshots:
        for assignment in participant.assignment_snapshot:
            source = assignment.get("sampling_source", {})
            item = snapshot_by_id.get(assignment.get("question_snapshot_id"))
            attribution_v2_rows.append(
                [
                    participant.student_id,
                    participant.version_id,
                    item.code if item else "",
                    item.stable_code if item else "",
                    assignment.get("assigned_sequence_index", ""),
                    source.get("quota_kc_code") or "",
                    source.get("quota_question_type") or "",
                    source.get("match_scope") or "",
                    hashlib.sha256(participant.sampling_seed.encode()).hexdigest()
                    if participant.sampling_seed
                    else "",
                    source.get("algorithm_version")
                    or participant.sampling_algorithm_version,
                    source.get("history_role") or "",
                    source.get("cross_experiment_item_policy") or "",
                ]
            )
    preview_results = [
        task.preview_class_result
        for task in class_tasks
        if task.preview_class_result_id
    ]
    preview_runs = {
        result.preview_run_id: result.preview_run
        for result in preview_results
    }

    rows = {
        "subjects": (identity_headers, subject_values),
        "sessions": ((*identity_headers, "session_id", "experiment_id", "version", "class_task_id", "participant_snapshot_id", "schedule_revision_id", "status", "started_at", "completed_at", "task_opens_at", "task_closes_at", "total_time_limit_seconds", "initial_effective_deadline_at", "timer_semantics_version", "terminal_reason", "authoritative_terminal_at", "terminal_recorded_at", "recovery_round", "current_deadline_at", "personal_deadline_at"), [[*identity_for_session(session.pk), session.pk, session.version.experiment_id, session.version.version_number, session.class_task_id or "", session.participant_snapshot_id or "", session.schedule_revision_id or "", session.status, session.started_at.isoformat(), session.completed_at.isoformat() if session.completed_at else "", export_time(session.task_opens_at), export_time(session.task_closes_at), session.total_time_limit_seconds or "", export_time(session.initial_effective_deadline_at), session.timer_semantics_version, session.terminal_reason, export_time(session.authoritative_terminal_at), export_time(session.terminal_recorded_at), session.recovery_round, export_time(session.deadline_at), export_time(session.personal_deadline_at)] for session in sessions]),
        "assignments": ((*identity_headers, "question_code", "assignment_id", "session_id", "question_snapshot_id", "item_instance_id", "assigned_sequence_index"), [[*identity_for_session(assignment.session_id), assignment.snapshot.code, assignment.pk, assignment.session_id, assignment.snapshot_id, assignment.item_instance_id, assignment.assigned_sequence_index] for assignment in assignments]),
        "item_instances": (("question_code", "item_instance_id", "question_snapshot_id", "presentation_variant", "option_order", "parameter_values", "random_seed", "presentation_hash"), [[assignment.snapshot.code, assignment.item_instance_id, assignment.snapshot_id, assignment.item_instance.presentation_variant, json.dumps(assignment.item_instance.option_order), json.dumps(assignment.item_instance.parameter_values), assignment.item_instance.random_seed, assignment.item_instance.presentation_hash] for assignment in _unique_assignments(assignments, "item_instance_id")]),
        "attempts": ((*identity_headers, "question_code", "attempt_id", "session_id", "question_snapshot_id", "item_instance_id", "status", "first_opened_at", "last_opened_at", "visit_count", "draft_answer", "draft_saved_at", "final_answer", "submitted_at", "first_to_submit_duration_ms", "last_visit_to_submission_ms", "is_correct", "score", "first_open_sequence_index", "submission_sequence_index", "active_response_event_id", "response_version", "final_response_event_id", "formal_capture_boundary", "current_response_snapshot_id", "final_response_snapshot_id", "submission_event_id", "submission_grade_status", "grade_version", "grade_is_correct", "grade_score", "grade_max_score", "normalized_response", "grading_method", "grading_version", "grader_version", "graded_at", "item_started_at", "per_item_time_limit_seconds", "item_deadline_at", "initial_effective_deadline_at", "current_effective_deadline_at", "active_seconds", "active_seconds_at_recovery", "close_reason", "authoritative_terminal_at", "terminal_recorded_at"), [_attempt_export_row(assignment, attempt_by_assignment.get(assignment.pk), session_by_id, grade_by_attempt) for assignment in assignments if assignment.session_id in session_by_id]),
        "assessment_events": (event_headers, event_rows),
        "response_events": (response_headers, response_rows),
        "exposure_events": (projection_headers, exposure_rows),
        "navigation_events": (navigation_headers, navigation_rows),
        "telemetry_events": (projection_headers, telemetry_rows),
        "simulation_events": (compact_event_headers, simulation_event_rows),
        "simulation_snapshots": ((*identity_headers, "question_code", "snapshot_id", "attempt_id", "simulation_version", "state_metadata", "content_hash", "content_size", "created_at"), [[*identity_and_question_for_attempt(simulation.attempt_id), simulation.snapshot_id, simulation.attempt_id, simulation.simulation_version, json.dumps(simulation.state_metadata), simulation.content_hash, simulation.content_size, simulation.created_at.isoformat()] for simulation in selected_simulations]),
        "confidence_assignments": ((*identity_headers, "question_code", "confidence_assignment_id", "attempt_id", "condition", "assignment_unit", "assignment_version", "response_status", "probability", "assigned_at", "committed_at"), [[*identity_and_question_for_attempt(confidence.attempt_id), confidence.confidence_assignment_id, confidence.attempt_id, confidence.condition, confidence.assignment_unit, confidence.assignment_version, confidence.response_status, confidence.probability, confidence.assigned_at.isoformat(), confidence.committed_at.isoformat() if confidence.committed_at else ""] for confidence in selected_confidence]),
        "item_snapshots": (("question_code", "question_snapshot_id", "version_id", "stable_code", "question_type", "primary_kc", "primary_kc_id", "primary_kc_path", "kc_level_1", "kc_level_2", "kc_level_3", "all_kcs", "all_kc_paths", "all_kc_snapshots"), [[item.code, item.pk, item.version_id, item.stable_code, item.question_type, item.primary_kc, item.primary_kc_id_snapshot or "", item.primary_kc_path, primary_level_name(item, 0), primary_level_name(item, 1), primary_level_name(item, 2), json.dumps(item.all_kcs, ensure_ascii=False), json.dumps(item.all_kc_paths, ensure_ascii=False), json.dumps(item.all_kc_snapshots, ensure_ascii=False)] for item in metadata_snapshots]),
        "question_content": (
            ("question_code", "question_snapshot_id", "version_id", "source_question_id", "source_sequence", "course_sequence", "stable_code", "source_bank_id", "source_bank_name", "bank_sequence", "question_type", "text", "options", "correct_answer", "accepted_answers", "absolute_tolerance", "relative_tolerance", "explanation", "teacher_difficulty", "empirical_difficulty", "primary_kc", "all_kcs", "primary_kc_id", "primary_kc_path", "primary_kc_levels", "all_kc_paths", "all_kc_snapshots", "root_kc_code", "kc_mapping_snapshot", "kc_mapping_schema_version", "kc_mapping_hash", "release_fingerprint"),
            [
                [
                    item.code, item.pk, item.version_id, item.source_question_id,
                    item.source_sequence, item.course_sequence or "", item.stable_code,
                    item.source_bank_id or "", item.source_bank_name,
                    item.bank_sequence or "", item.question_type, item.text,
                    json.dumps(item.options, ensure_ascii=False),
                    json.dumps(item.correct_answer, ensure_ascii=False),
                    json.dumps(item.accepted_answers, ensure_ascii=False),
                    item.absolute_tolerance if item.absolute_tolerance is not None else "",
                    item.relative_tolerance if item.relative_tolerance is not None else "",
                    item.explanation,
                    item.teacher_difficulty if item.teacher_difficulty is not None else "",
                    item.empirical_difficulty if item.empirical_difficulty is not None else "",
                    item.primary_kc, json.dumps(item.all_kcs, ensure_ascii=False),
                    item.primary_kc_id_snapshot or "", item.primary_kc_path,
                    json.dumps(item.primary_kc_levels, ensure_ascii=False),
                    json.dumps(item.all_kc_paths, ensure_ascii=False),
                    json.dumps(item.all_kc_snapshots, ensure_ascii=False),
                    item.root_kc_code,
                    json.dumps(item.kc_mapping_snapshot, ensure_ascii=False),
                    item.kc_mapping_schema_version, item.kc_mapping_hash,
                    item.release_fingerprint,
                ]
                for item in metadata_snapshots
            ],
        ),
        "item_kc_mapping": (("question_code", "question_snapshot_id", "primary_kc", "primary_kc_id", "primary_kc_path", "kc_level_1", "kc_level_2", "kc_level_3", "all_kcs", "all_kc_paths", "all_kc_snapshots"), [[item.code, item.pk, item.primary_kc, item.primary_kc_id_snapshot or "", item.primary_kc_path, primary_level_name(item, 0), primary_level_name(item, 1), primary_level_name(item, 2), json.dumps(item.all_kcs, ensure_ascii=False), json.dumps(item.all_kc_paths, ensure_ascii=False), json.dumps(item.all_kc_snapshots, ensure_ascii=False)] for item in metadata_snapshots]),
        "q_matrix": (("question_code", "question_snapshot_id", "primary_kc", "primary_kc_id", "primary_kc_path", "kc_level_1", "kc_level_2", "kc_level_3", "all_kcs", "all_kc_paths", "all_kc_snapshots"), [[item.code, item.pk, item.primary_kc, item.primary_kc_id_snapshot or "", item.primary_kc_path, primary_level_name(item, 0), primary_level_name(item, 1), primary_level_name(item, 2), json.dumps(item.all_kcs, ensure_ascii=False), json.dumps(item.all_kc_paths, ensure_ascii=False), json.dumps(item.all_kc_snapshots, ensure_ascii=False)] for item in metadata_snapshots]),
        "kc_dictionary": (("kc_id", "kc_name", "kc_full_path"), kc_dictionary_rows),
        "kc_dictionary_v2": (
            ("kc_code", "kc_name", "description", "status"),
            [frozen_kcs[code] for code in sorted(frozen_kcs)],
        ),
        "question_kc_structure_v2": (
            ("question_code", "stable_question_code", "kc_code", "depth", "parent_kc_code", "order", "mapping_hash", "schema_version"),
            structure_v2_rows,
        ),
        "q_matrix_v2": (
            ("question_code", "stable_question_code", "mapping_hash", *q_matrix_codes),
            q_matrix_v2_rows,
        ),
        "participant_question_attribution_v2": (
            ("participant_id", "version_id", "question_code", "stable_question_code", "assigned_sequence_index", "quota_kc_code", "quota_question_type", "match_scope", "sampling_seed_sha256", "algorithm_version", "history_role", "cross_experiment_item_policy"),
            attribution_v2_rows,
        ),
        "response_snapshots": (
            (*identity_headers, "question_code", "snapshot_id", "attempt_id", "content", "content_hash", "content_length", "state_json", "state_hash", "schema_version", "created_at"),
            [
                [
                    *identity_and_question_for_attempt(item.attempt_id),
                    item.snapshot_id,
                    item.attempt_id,
                    item.content,
                    item.content_hash,
                    item.content_length,
                    json.dumps(item.state_json, ensure_ascii=False),
                    item.state_hash,
                    item.schema_version,
                    item.created_at.isoformat(),
                ]
                for item in response_snapshots
                if item.attempt_id in attempt_by_id
            ],
        ),
        "response_trajectory": (response_trajectory_headers, response_trajectory_rows),
        "response_replay": (response_replay_headers, response_replay_rows),
        "session_summary": (session_summary_headers, session_summary_rows),
        "response_input_observations": (
            (*identity_headers, "question_code", "observation_id", "source_event_id", "source_event_type", "event_version", "envelope_version", "attempt_id", "client_instance_id", "client_sequence_no", "client_monotonic_ms", "batch_start_client_monotonic_ms", "batch_end_client_monotonic_ms", "client_wall_time_utc", "client_time_origin", "time_sync_id", "estimated_occurred_at_utc", "derived_occurred_at_utc", "server_received_at", "server_committed_at", "server_ingest_sequence_no", "actor_source", "capture_mode", "quality_flags", "time_quality", "coverage_run_id", "app_version", "edit_operation", "content_hash_before", "content_hash_after", "content_length_before", "content_length_after", "input_method", "is_ime", "observed_at"),
            [
                input_observation_row(item)
                for item in input_observations
                if item.attempt_id in attempt_by_id
            ],
        ),
        "submissions": (
            (*identity_headers, "question_code", "submission_id", "attempt_id", "submission_token", "final_response_snapshot_id", "grading_status", "score", "is_correct", "client_clicked_at", "server_received_at", "committed_at"),
            [
                [
                    *identity_and_question_for_attempt(item.attempt_id),
                    item.pk,
                    item.attempt_id,
                    item.submission_token,
                    item.final_response_snapshot_id,
                    item.grading_status,
                    item.score,
                    item.is_correct,
                    item.client_clicked_at.isoformat() if item.client_clicked_at else "",
                    item.server_received_at.isoformat(),
                    item.committed_at.isoformat(),
                ]
                for item in submissions
                if item.attempt_id in attempt_by_id
            ],
        ),
        "grade_records": (
            (*identity_headers, "question_code", "grade_id", "submission_id", "attempt_id", "version", "status", "validity_status", "score", "max_score", "is_correct", "grading_method", "grading_version", "graded_at", "supersedes_grade_id", "voided_at", "void_reason", "grader_id", "comment", "grader_version", "voided_by_id", "created_at"),
            [
                [
                    *identity_and_question_for_attempt(item.submission.attempt_id),
                    item.pk,
                    item.submission_id,
                    item.submission.attempt_id,
                    item.version,
                    item.status,
                    item.validity_status,
                    item.score,
                    item.max_score,
                    item.is_correct,
                    item.grading_method,
                    item.grading_version,
                    item.graded_at.isoformat() if item.graded_at else "",
                    item.supersedes_grade_id or "",
                    item.voided_at.isoformat() if item.voided_at else "",
                    item.void_reason,
                    item.grader_id or "",
                    item.comment,
                    item.grader_version,
                    item.voided_by_id or "",
                    item.created_at.isoformat(),
                ]
                for item in grade_records
                if item.submission.attempt_id in attempt_by_id
            ],
        ),
        "condition_assignments": (
            ("condition_assignment_id", "condition_id", "condition_key", "condition_label", "subject_id", "session_id", "assignment_unit_type", "assignment_algorithm", "assignment_version", "randomization_seed", "randomization_draw", "stratum", "assignment_probability", "assigned_at"),
            [
                [
                    item.pk,
                    item.condition_id,
                    item.condition.key,
                    item.condition.label,
                    item.subject_id or "",
                    item.session_id or "",
                    "subject" if item.subject_id else "session",
                    item.algorithm,
                    item.algorithm_version,
                    item.seed,
                    item.draw,
                    item.stratum,
                    item.probability,
                    item.assigned_at.isoformat(),
                ]
                for item in condition_assignments
            ],
        ),
        "class_snapshots": (
            ("class_snapshot_id", "version_id", "source_class_group_id", "name_snapshot", "academic_year_snapshot", "capture_mode", "created_at"),
            [
                [
                    item.class_snapshot_id,
                    item.version_id,
                    item.source_class_group_id or "",
                    item.name_snapshot,
                    item.academic_year_snapshot,
                    item.capture_mode,
                    item.created_at.isoformat(),
                ]
                for item in class_snapshots
            ],
        ),
        "participant_snapshots": (
            (
                "participant_snapshot_id",
                "version_id",
                "student_id",
                "class_snapshot_id",
                "sampling_seed_sha256",
                "ordering_seed_sha256",
                "sampling_algorithm_version",
                "assignment_snapshot",
                "created_at",
            ),
            [
                [
                    item.pk,
                    item.version_id,
                    item.student_id,
                    item.class_snapshot_id or "",
                    hashlib.sha256(item.sampling_seed.encode()).hexdigest()
                    if item.sampling_seed
                    else "",
                    hashlib.sha256(item.ordering_seed.encode()).hexdigest()
                    if item.ordering_seed
                    else "",
                    item.sampling_algorithm_version,
                    json.dumps(item.assignment_snapshot, ensure_ascii=False),
                    item.created_at.isoformat(),
                ]
                for item in participant_snapshots
            ],
        ),
        "capture_contracts": (
            ("capture_contract_id", "contract_name", "contract_version", "schema_document", "schema_sha256", "created_at"),
            [[item.pk, item.contract_name, item.contract_version, json.dumps(item.schema_document, ensure_ascii=False), item.schema_sha256, item.created_at.isoformat()] for item in capture_contracts],
        ),
        "client_time_syncs": (
            ("time_sync_id", "session_id", "client_instance_id", "client_wall_time_utc", "client_monotonic_ms", "client_time_origin", "server_received_at", "server_responded_at", "client_received_wall_time_utc", "client_received_monotonic_ms", "estimated_clock_offset_ms", "estimated_rtt_ms", "sync_quality", "finalized_at"),
            [[item.pk, item.session_id, item.client_instance_id, item.client_wall_time_utc.isoformat(), item.client_monotonic_ms, item.client_time_origin.isoformat(), item.server_received_at.isoformat(), item.server_responded_at.isoformat(), item.client_received_wall_time_utc.isoformat() if item.client_received_wall_time_utc else "", item.client_received_monotonic_ms if item.client_received_monotonic_ms is not None else "", item.estimated_clock_offset_ms if item.estimated_clock_offset_ms is not None else "", item.estimated_rtt_ms if item.estimated_rtt_ms is not None else "", item.sync_quality, item.finalized_at.isoformat() if item.finalized_at else ""] for item in time_syncs],
        ),
        "client_capture_checkpoints": (
            ("checkpoint_id", "session_id", "student_id", "client_instance_id", "highest_emitted_sequence_no", "highest_acked_sequence_no", "highest_contiguous_acked_sequence_no", "pending_event_count", "failed_event_count", "missing_local_sequence_ranges", "client_monotonic_ms", "client_wall_time_utc", "client_time_origin", "time_sync_id", "trigger", "quality_flags", "app_version", "server_received_at"),
            [[item.pk, item.session_id, item.student_id, item.client_instance_id, item.highest_emitted_sequence_no, item.highest_acked_sequence_no, item.highest_contiguous_acked_sequence_no, item.pending_event_count, item.failed_event_count, json.dumps(item.missing_local_sequence_ranges), item.client_monotonic_ms, item.client_wall_time_utc.isoformat(), item.client_time_origin.isoformat(), item.time_sync_id or "", item.trigger, json.dumps(item.quality_flags), item.app_version, item.server_received_at.isoformat()] for item in checkpoints],
        ),
        "recording_coverage_runs": (
            ("coverage_run_id", "scope_type", "scope_definition", "scope_hash", "coverage_version", "algorithm_version", "code_version", "parameter_document", "parameter_hash", "source_cutoff_at", "source_high_water_marks", "capture_contract_hashes", "source_hash", "run_status", "generated_at"),
            [
                [
                    run.run_id,
                    run.scope_type,
                    json.dumps(run.scope_definition, ensure_ascii=False, sort_keys=True),
                    run.scope_hash,
                    run.coverage_version,
                    run.algorithm_version,
                    run.code_version,
                    json.dumps(run.parameter_document, ensure_ascii=False, sort_keys=True),
                    run.parameter_hash,
                    run.source_cutoff_at.isoformat(),
                    json.dumps(run.source_high_water_marks, ensure_ascii=False, sort_keys=True),
                    json.dumps(run.capture_contract_hashes, ensure_ascii=False, sort_keys=True),
                    run.source_hash,
                    run.run_status,
                    run.generated_at.isoformat(),
                ]
                for _run_id, run in sorted(coverage_runs.items())
            ],
        ),
        "recording_coverage_sessions": (
            ("coverage_run_id", "coverage_version", "coverage_parameter_hash", "coverage_source_cutoff", "session_id", "overall_status", "lifecycle_status", "navigation_status", "client_sequence_status", "time_sync_status", "telemetry_status", "client_instance_count", "sequence_gap_count", "missing_sequence_ranges", "tail_gap_detected", "missing_reasons", "analysis_eligibility"),
            [[item.run_id, item.run.coverage_version, item.run.parameter_hash, item.run.source_cutoff_at.isoformat(), item.session_id, item.overall_status, item.lifecycle_status, item.navigation_status, item.client_sequence_status, item.time_sync_status, item.telemetry_status, item.client_instance_count, item.sequence_gap_count, json.dumps(item.missing_sequence_ranges), item.tail_gap_detected, json.dumps(item.missing_reasons), json.dumps(item.analysis_eligibility, ensure_ascii=False)] for item in latest_coverage_session.values()],
        ),
        "recording_coverage_attempts": (
            ("coverage_run_id", "coverage_version", "coverage_parameter_hash", "coverage_source_cutoff", "attempt_id", "overall_status", "enter_status", "render_status", "leave_status", "response_chain_status", "final_state_status", "time_quality", "response_state_count", "sequence_gap_count", "missing_reasons", "analysis_eligibility", "derived_features"),
            [[item.run_id, item.run.coverage_version, item.run.parameter_hash, item.run.source_cutoff_at.isoformat(), item.attempt_id, item.overall_status, item.enter_status, item.render_status, item.leave_status, item.response_chain_status, item.final_state_status, item.time_quality, item.response_state_count, item.sequence_gap_count, json.dumps(item.missing_reasons), json.dumps(item.analysis_eligibility, ensure_ascii=False), json.dumps(item.derived_features, ensure_ascii=False)] for item in latest_coverage_attempt.values()],
        ),
    }
    rows.update(
        {
            "release_metadata": (
                (
                    "release_snapshot_id", "experiment_id", "version_number",
                    "published_at", "timer_semantics_version",
                    "randomization_semantics_version", "definition_fingerprint",
                    "canonicalization_version", "feedback_mode",
                    "feedback_release_policy", "short_answer_grading_policy",
                    "sampling_mode",
                    "sampling_total_count", "sampling_kcs", "sampling_inventory",
                    "cross_experiment_item_policy", "review_count",
                    "taxonomy_snapshot", "candidate_resolution",
                    "sampling_definition", "timezone",
                ),
                [
                    [
                        release.pk, release.experiment_id, release.version_number,
                        export_time(release.published_at), release.timer_semantics_version,
                        release.randomization_semantics_version,
                        release.definition_fingerprint, release.canonicalization_version,
                        release.configuration.get("feedback_mode", ""),
                        release.configuration.get("feedback_release_policy", ""),
                        release.configuration.get("short_answer_grading_policy", ""),
                        release.configuration.get("sampling_mode")
                        or (
                            "type_quota_legacy"
                            if release.configuration.get("selection_mode") == "quota"
                            else "all_legacy"
                        ),
                        release.configuration.get("sampling_total_count", ""),
                        json.dumps(release.configuration.get("sampling_kcs", []), ensure_ascii=False),
                        json.dumps(release.configuration.get("sampling_inventory", []), ensure_ascii=False),
                        release.configuration.get("cross_experiment_item_policy", "unrestricted"),
                        release.configuration.get("review_count", 0),
                        json.dumps(release.configuration.get("taxonomy_snapshot", []), ensure_ascii=False),
                        json.dumps(release.configuration.get("candidate_resolution", []), ensure_ascii=False),
                        json.dumps(release.configuration.get("sampling", {}), ensure_ascii=False),
                        release.configuration.get("timezone", settings.TIME_ZONE),
                    ]
                    for release in releases
                ],
            ),
            "research_tools": (
                (
                    "experiment_id", "version_id", "version_number",
                    "research_tools", "research_tools_sha256",
                ),
                [
                    [
                        release.experiment_id, release.pk, release.version_number,
                        json.dumps(
                            release.configuration.get("research_tools", {}),
                            ensure_ascii=False,
                            sort_keys=True,
                        ),
                        sha256_bytes(
                            json.dumps(
                                release.configuration.get("research_tools", {}),
                                ensure_ascii=False,
                                separators=(",", ":"),
                                sort_keys=True,
                            ).encode("utf-8")
                        ),
                    ]
                    for release in releases
                ],
            ),
            "class_experiment_tasks": (
                (
                    "class_task_id", "experiment_id", "release_snapshot_id",
                    "class_group_id", "class_snapshot_id", "published_at",
                    "class_published_at", "randomization_algorithm_version",
                    "root_seed_fingerprint", "preview_class_result_id",
                    "cancelled_at", "manual_closed_at", "invalidated_at",
                    "feedback_released_at", "feedback_released_by_id",
                    "terminal_reason", "provenance", "manual_closure_history",
                ),
                [
                    [
                        task.pk, task.experiment_id, task.release_snapshot_id,
                        task.class_group_id, task.class_snapshot_id,
                        export_time(task.published_at), export_time(task.class_published_at),
                        task.randomization_identity.algorithm_version
                        if task.randomization_identity_id else "",
                        task.randomization_identity.root_seed_fingerprint
                        if task.randomization_identity_id else "",
                        task.preview_class_result_id or "", export_time(task.cancelled_at),
                        export_time(task.manual_closed_at), export_time(task.invalidated_at),
                        export_time(task.feedback_released_at),
                        task.feedback_released_by_id or "",
                        task.terminal_reason or "", task.provenance or "",
                        json.dumps([
                            {"closed_at": export_time(closure.closed_at), "schedule_revision_no": closure.schedule_revision_no, "created_by_id": closure.created_by_id, "reason": closure.reason}
                            for closure in task.manual_closures.all()
                        ], ensure_ascii=False),
                    ]
                    for task in class_tasks
                ],
            ),
            "schedule_revisions": (
                (
                    "class_task_id", "revision_no", "opens_at", "closes_at",
                    "timezone", "reason", "provenance", "created_at", "total_time_limit_seconds", "per_question_time_limit_seconds", "reopens_manual_close", "created_by_id",
                ),
                [
                    [
                        revision.class_task_id, revision.revision_no,
                        export_time(revision.opens_at), export_time(revision.closes_at),
                        revision.timezone_name, revision.reason, revision.provenance,
                        export_time(revision.created_at), revision.total_time_limit_seconds or "", revision.per_question_time_limit_seconds or "", revision.reopens_manual_close, revision.created_by_id or "",
                    ]
                    for revision in schedule_revisions
                ],
            ),
            "release_question_mappings": (
                (
                    "release_snapshot_id", "source_question_id",
                    "question_release_fingerprint", "question_snapshot_id",
                    "created_at",
                ),
                [
                    [
                        mapping.release_snapshot_id, mapping.source_question_id,
                        mapping.question_release_fingerprint,
                        mapping.question_snapshot_id, export_time(mapping.created_at),
                    ]
                    for mapping in release_mappings
                ],
            ),
            "publish_preview_runs": (
                (
                    "preview_run_id", "plan_revision", "status",
                    "definition_fingerprint", "input_hash", "output_hash",
                    "algorithm_version", "created_at", "invalidated_at",
                    "invalidation_reason",
                ),
                [
                    [
                        run.pk, run.plan_revision, run.status,
                        run.definition_fingerprint, run.input_hash, run.output_hash,
                        run.algorithm_version, export_time(run.created_at),
                        export_time(run.invalidated_at), run.invalidation_reason,
                    ]
                    for run in preview_runs.values()
                ],
            ),
            "publish_preview_class_results": (
                (
                    "preview_class_result_id", "preview_run_id",
                    "class_task_id", "participant_count", "quota_valid",
                    "input_hash", "output_hash", "algorithm_version",
                    "root_seed_fingerprint", "created_at",
                ),
                [
                    [
                        result.pk, result.preview_run_id,
                        next(
                            task.pk for task in class_tasks
                            if task.preview_class_result_id == result.pk
                        ),
                        result.participant_count, result.quota_valid,
                        result.input_hash, result.output_hash,
                        result.randomization_identity.algorithm_version,
                        result.randomization_identity.root_seed_fingerprint,
                        export_time(result.created_at),
                    ]
                    for result in preview_results
                ],
            ),
        }
    )
    derivation_id = str(getattr(snapshot, "snapshot_id", "preview"))
    derivation_version = EXPORT_DERIVATION_VERSION
    base = (*identity_headers, "question_code", "attempt_id", "derivation_run_id", "derivation_version")
    feature_rows, kt_rows, cd_rows, cal_rows = [], [], [], []
    for attempt in attempts:
        ident = [
            *identity_and_question_for_attempt(attempt.pk),
            attempt.pk,
            derivation_id,
            derivation_version,
        ]
        feature_rows.append([*ident, attempt.is_correct, attempt.score, attempt.visit_count, attempt.first_opened_at.isoformat() if attempt.first_opened_at else "", attempt.submitted_at.isoformat() if attempt.submitted_at else ""])
        item_snapshot = attempt.assignment.snapshot
        if (
            attempt.status == Attempt.Status.SUBMITTED
            and attempt.submission_sequence_index is not None
            and attempt.is_correct is not None
        ):
            kt_rows.append([
                *ident,
                "formal_submission",
                attempt.submission_sequence_index,
                int(bool(attempt.is_correct)),
                item_snapshot.root_kc_code
                if item_snapshot.kc_mapping_schema_version == "question-kc-tree-v2"
                else item_snapshot.primary_kc or "",
            ])
        response_value = attempt.final_answer if attempt.final_answer is not None else attempt.draft_answer
        cd_kcs = (
            [
                node["kc_code"]
                for node in item_snapshot.kc_mapping_snapshot.get("nodes", [])
            ]
            if item_snapshot.kc_mapping_schema_version == "question-kc-tree-v2"
            else item_snapshot.all_kcs
        )
        cd_rows.append([*ident, json.dumps(response_value, ensure_ascii=False) if response_value is not None else "", json.dumps(cd_kcs, ensure_ascii=False)])
        confidence = next((c for c in selected_confidence if c.attempt_id == attempt.pk), None)
        probability = confidence.probability if confidence and confidence.probability is not None else ""
        signed_prediction_error = (
            float(probability) / 100.0 - float(bool(attempt.is_correct))
            if probability != "" and attempt.is_correct is not None
            else ""
        )
        cal_rows.append([
            *ident,
            probability,
            attempt.is_correct,
            signed_prediction_error,
            signed_prediction_error,
            confidence.condition if confidence else "",
        ])
    rows.update({
        "attempt_features": (base + ("is_correct", "score", "visit_count", "first_opened_at", "submitted_at"), feature_rows),
        "kt_sequences": (base + ("sequence_type", "sequence_index", "outcome", "root_kc_code"), kt_rows),
        "cd_response_matrix": (base + ("response", "kc_codes"), cd_rows),
        "calibration_features": (
            base + (
                "confidence_probability", "is_correct", "signed_prediction_error",
                "calibration_error", "condition",
            ),
            cal_rows,
        ),
    })
    return rows


def _unique_assignments(assignments, field):
    seen = set()
    result = []
    for assignment in assignments:
        value = getattr(assignment, field)
        if value not in seen:
            result.append(assignment)
            seen.add(value)
    return result


def _attempt_export_row(assignment, attempt, session_by_id, grade_by_attempt=None):
    session = session_by_id[assignment.session_id]
    identity = [session.student.student_number or "", session.student.shown_name, assignment.snapshot.code]
    linkage = [assignment.session_id, assignment.snapshot.pk, assignment.item_instance_id]
    if attempt is None:
        return [*identity, "", *linkage] + [""] * 41
    grade = (grade_by_attempt or {}).get(attempt.pk)
    submission = getattr(attempt, "submission_event", None)
    return [
        *identity,
        attempt.pk,
        *linkage,
        attempt.status,
        attempt.first_opened_at.isoformat() if attempt.first_opened_at else "",
        attempt.last_opened_at.isoformat() if attempt.last_opened_at else "",
        attempt.visit_count,
        json.dumps(attempt.draft_answer, ensure_ascii=False) if attempt.draft_answer is not None else "",
        attempt.draft_saved_at.isoformat() if attempt.draft_saved_at else "",
        json.dumps(attempt.final_answer, ensure_ascii=False) if attempt.final_answer is not None else "",
        attempt.submitted_at.isoformat() if attempt.submitted_at else "",
        attempt.first_to_submit_duration_ms,
        attempt.last_visit_to_submission_ms,
        attempt.is_correct,
        attempt.score,
        attempt.first_open_sequence_index,
        attempt.submission_sequence_index,
        attempt.active_response_event_id or "",
        attempt.response_version,
        attempt.final_response_event_id or "",
        json.dumps(attempt.formal_capture_boundary or {}, ensure_ascii=False),
        attempt.current_response_snapshot_id,
        attempt.final_response_snapshot_id,
        submission.pk if submission else "",
        submission.grading_status if submission else "",
        grade.version if grade else "",
        grade.is_correct if grade else "",
        grade.score if grade else "",
        grade.max_score if grade else "",
        json.dumps(grade.normalized_response, ensure_ascii=False) if grade and grade.normalized_response is not None else "",
        grade.grading_method if grade else "",
        grade.grading_version if grade else "",
        grade.grader_version if grade else "",
        grade.graded_at.isoformat() if grade and grade.graded_at else "",
        attempt.item_started_at.isoformat() if attempt.item_started_at else "",
        attempt.per_item_time_limit_seconds or "",
        attempt.item_deadline_at.isoformat() if attempt.item_deadline_at else "",
        attempt.initial_effective_deadline_at.isoformat()
        if attempt.initial_effective_deadline_at else "",
        attempt.current_effective_deadline_at.isoformat()
        if attempt.current_effective_deadline_at else "",
        attempt.active_seconds,
        attempt.active_seconds_at_recovery,
        attempt.close_reason,
        attempt.authoritative_terminal_at.isoformat()
        if attempt.authoritative_terminal_at else "",
        attempt.terminal_recorded_at.isoformat()
        if attempt.terminal_recorded_at else "",
    ]


def _finalize_export_job(job, zip_path, manifest):
    from apps.accounts.course_access import lock_course
    digest = _file_sha256(zip_path)
    with transaction.atomic():
        lock_course(job.course_id)
        current = ExportJob.objects.select_for_update().filter(
            pk=job.pk, status=ExportJob.Status.RUNNING, locked_by=job.locked_by,
            lease_expires_at__gt=timezone.now(), snapshot__invalidated_at=None, snapshot__redacted_at=None,
        ).filter(models.Q(expires_at__isnull=True) | models.Q(expires_at__gt=timezone.now())).first()
        if current is None or getattr(job, '_lease_lost', threading.Event()).is_set():
            if zip_path == getattr(job, '_candidate_path', None) and os.path.exists(zip_path):
                os.unlink(zip_path)
            raise ExportContractError('导出任务已失效或租约丢失，生成文件未发布。')
        current.status = ExportJob.Status.SUCCEEDED
        current.finished_at = timezone.now()
        current.progress_stage = 'complete'
        current.artifact_path = zip_path
        current.artifact_sha256 = digest
        current.locked_by = ''
        current.lease_expires_at = None
        current.save(update_fields=['status', 'finished_at', 'progress_stage', 'artifact_path',
            'artifact_sha256', 'locked_by', 'lease_expires_at'])
    job.refresh_from_db()
    return job


def execute_purge_request(purge_request, *, executed_by):
    from apps.accounts.models import DataPurgeRequest
    from apps.assessment.models import (
        AssessmentEvent, Attempt, ClientCaptureCheckpoint, ClientTimeSync,
        ConfidenceAssignment, DraftSaveEvent, ExperimentSession,
        QuestionVisitEvent, ResponseInputObservation, ResponseSnapshot,
        SimulationSnapshot, StudentQuestionAssignment, SubmissionEvent,
        GradeRecord,
    )
    from apps.experiments.models import ConditionAssignment

    using = purge_request._state.db or "default"

    def remove_committed_artifacts(job_ids, reason):
        for job in ExportJob.objects.using(using).filter(pk__in=job_ids):
            remove_export_artifact(job, reason=reason, invalidated=True)

    with transaction.atomic(using=using):
        locked_request = (
            DataPurgeRequest.objects.using(using)
            .select_for_update()
            .get(pk=purge_request.pk)
        )
        course_id = locked_request.scope_definition.get("course_id")
        if locked_request.scope_type != "course" or not course_id:
            raise ValueError("purge scope missing course_id")

        reason = f"purge:{locked_request.reason_code}"
        if locked_request.status == DataPurgeRequest.Status.EXECUTED:
            pending_job_ids = list(
                ExportJob.objects.using(using)
                .filter(
                    course_id=course_id,
                    status=ExportJob.Status.INVALIDATED,
                )
                .exclude(artifact_path="")
                .values_list("pk", flat=True)
            )
            if pending_job_ids:
                transaction.on_commit(
                    lambda: remove_committed_artifacts(pending_job_ids, reason),
                    using=using,
                )
            return dict(locked_request.result_counts)

        session_ids = list(
            ExperimentSession.objects.using(using)
            .select_for_update()
            .filter(version__experiment__course_id=course_id)
            .order_by("pk")
            .values_list("pk", flat=True)
        )
        attempt_ids = list(
            Attempt.objects.using(using)
            .select_for_update()
            .filter(session_id__in=session_ids)
            .order_by("pk")
            .values_list("pk", flat=True)
        )
        counts = {
            "sessions": len(session_ids),
            "attempts": len(attempt_ids),
            "events": AssessmentEvent.objects.using(using).filter(
                session_id__in=session_ids
            ).count(),
        }

        def purge_rows(key, queryset):
            count = queryset.count()
            counts[key] = count
            if count:
                queryset._raw_delete(using=using)

        # Coverage and randomization rows protect their source runtime records.
        # They are derived session results, while their run/condition identities
        # remain available as provenance for unaffected courses.
        purge_rows(
            "recording_coverage_attempts",
            RecordingCoverageAttempt.objects.using(using).filter(
                attempt_id__in=attempt_ids
            ),
        )
        purge_rows(
            "recording_coverage_sessions",
            RecordingCoverageSession.objects.using(using).filter(
                session_id__in=session_ids
            ),
        )
        purge_rows(
            "condition_assignments",
            ConditionAssignment.objects.using(using).filter(
                session_id__in=session_ids
            ),
        )
        purge_rows(
            "client_capture_checkpoints",
            ClientCaptureCheckpoint.objects.using(using).filter(
                session_id__in=session_ids
            ),
        )
        purge_rows(
            "response_input_observations",
            ResponseInputObservation.objects.using(using).filter(
                attempt_id__in=attempt_ids
            ),
        )
        purge_rows(
            "grade_records",
            GradeRecord.objects.using(using).filter(
                submission__attempt_id__in=attempt_ids
            ),
        )
        purge_rows(
            "submission_events",
            SubmissionEvent.objects.using(using).filter(
                attempt_id__in=attempt_ids
            ),
        )
        purge_rows(
            "confidence_assignments",
            ConfidenceAssignment.objects.using(using).filter(
                attempt_id__in=attempt_ids
            ),
        )
        purge_rows(
            "simulation_snapshots",
            SimulationSnapshot.objects.using(using).filter(
                attempt_id__in=attempt_ids
            ),
        )
        purge_rows(
            "draft_save_events",
            DraftSaveEvent.objects.using(using).filter(attempt_id__in=attempt_ids),
        )
        purge_rows(
            "question_visit_events",
            QuestionVisitEvent.objects.using(using).filter(
                session_id__in=session_ids
            ),
        )

        # Clear the four capture-v2 response heads and the session assignment
        # head before erasing their immutable targets through this explicit
        # governance path.
        Attempt.objects.using(using).filter(pk__in=attempt_ids).update(
            active_response_event_id=None,
            final_response_event_id=None,
            current_response_snapshot_id=None,
            final_response_snapshot_id=None,
        )
        ExperimentSession.objects.using(using).filter(pk__in=session_ids).update(
            current_assignment_id=None
        )
        purge_rows(
            "events",
            AssessmentEvent.objects.using(using).filter(session_id__in=session_ids),
        )
        purge_rows(
            "response_snapshots",
            ResponseSnapshot.objects.using(using).filter(attempt_id__in=attempt_ids),
        )
        purge_rows(
            "attempts",
            Attempt.objects.using(using).filter(pk__in=attempt_ids),
        )
        purge_rows(
            "student_question_assignments",
            StudentQuestionAssignment.objects.using(using).filter(
                session_id__in=session_ids
            ),
        )
        purge_rows(
            "client_time_syncs",
            ClientTimeSync.objects.using(using).filter(session_id__in=session_ids),
        )
        purge_rows(
            "sessions",
            ExperimentSession.objects.using(using).filter(pk__in=session_ids),
        )

        invalidated_at = timezone.now()
        snapshots = ExportSnapshot.objects.using(using).filter(course_id=course_id)
        counts["export_snapshots_invalidated"] = snapshots.count()
        snapshots.update(
            invalidated_at=invalidated_at,
            invalidated_reason=reason,
            filter_definition={"scope_redacted": True},
            scope_hash="",
            source_bundle=None,
            redacted_at=invalidated_at,
        )
        jobs = ExportJob.objects.using(using).select_for_update().filter(
            course_id=course_id
        )
        job_ids = list(jobs.values_list("pk", flat=True))
        counts["export_jobs_invalidated"] = len(job_ids)
        jobs.update(
            status=ExportJob.Status.INVALIDATED,
            artifact_removal_reason=reason,
        )

        locked_request.status = DataPurgeRequest.Status.EXECUTED
        locked_request.executed_by = executed_by
        locked_request.executed_at = timezone.now()
        locked_request.result_counts = counts
        locked_request.save(
            update_fields=["status", "executed_by", "executed_at", "result_counts"]
        )
        if job_ids:
            transaction.on_commit(
                lambda: remove_committed_artifacts(job_ids, reason),
                using=using,
            )
    return counts


def cleanup_expired_exports(*, now=None):
    now = now or timezone.now()
    removed = 0
    for job in ExportJob.objects.filter(expires_at__lte=now).exclude(status=ExportJob.Status.EXPIRED):
        if job.artifact_path and os.path.exists(job.artifact_path):
            removed += 1
        remove_export_artifact(job, reason="retention_expired")
    return removed
