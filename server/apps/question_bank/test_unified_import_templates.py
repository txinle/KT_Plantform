import csv
from io import BytesIO, StringIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
from openpyxl import load_workbook

from django.conf import settings
from django.test import TransactionTestCase
from django.urls import reverse

from apps.accounts.models import User


QUESTION_TEMPLATE_HEADERS = [
    "序号",
    "题目",
    "题型",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "选项E",
    "选项F",
    "答案",
    "可接受答案",
    "解析",
    "难度",
    "一级知识点",
    "二级知识点",
    "三级知识点",
]

QUESTION_TYPES = ["单选题", "多选题", "判断题", "填空题", "简答题", "数值题"]


class UnifiedImportTemplateDownloadTests(TransactionTestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(
            username="unified-template-teacher", role=User.Role.TEACHER
        )
        self.student = User.objects.create_user(
            username="unified-template-student", role=User.Role.STUDENT
        )
        self.urls = {
            "xlsx": reverse("teacher-question-import-template-excel"),
            "docx": reverse("teacher-question-import-template-word"),
        }

    def _download(self, url):
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        try:
            return response, (
                b"".join(response.streaming_content)
                if response.streaming
                else response.content
            )
        finally:
            response.close()

    def test_teacher_can_download_excel_and_word_templates_with_expected_mime_types(self):
        self.client.force_login(self.teacher)

        for extension, content_type in (
            ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ):
            with self.subTest(extension=extension):
                response, content = self._download(self.urls[extension])
                self.assertEqual(response["Content-Type"], content_type)
                self.assertEqual(
                    response["Content-Disposition"],
                    f'attachment; filename="question_import_template.{extension}"',
                )
                self.assertTrue(content)

    def test_excel_template_uses_one_data_sheet_and_matches_csv_samples(self):
        self.client.force_login(self.teacher)
        response, content = self._download(self.urls["xlsx"])
        workbook = load_workbook(BytesIO(content), data_only=True)

        self.assertEqual(workbook.sheetnames, ["题目导入"])
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, QUESTION_TEMPLATE_HEADERS)
        self.assertIn("替换", sheet["N1"].comment.text)
        self.assertIn("不会创建", sheet["N1"].comment.text)

        xlsx_rows = [
            {
                header: "" if value is None else str(value)
                for header, value in zip(headers, row)
            }
            for row in sheet.iter_rows(min_row=2, values_only=True)
        ]
        csv_rows = list(
            csv.DictReader(
                StringIO(
                    (settings.BASE_DIR / "fixtures" / "question_import_template.csv")
                    .read_text(encoding="utf-8-sig")
                )
            )
        )
        self.assertEqual(xlsx_rows, csv_rows)
        self.assertEqual([row["题型"] for row in xlsx_rows], QUESTION_TYPES)
        self.assertTrue(all(isinstance(row["序号"], str) for row in xlsx_rows))
        self.assertTrue(all(isinstance(row["题型"], str) for row in xlsx_rows))
        self.assertTrue(all(isinstance(row["答案"], str) for row in xlsx_rows))
        self.assertIn("\n", xlsx_rows[0]["题目"])
        self.assertIn("\n", xlsx_rows[1]["选项A"])
        self.assertIn("\n", xlsx_rows[1]["解析"])

    def test_word_template_has_instruction_before_structured_question_samples(self):
        self.client.force_login(self.teacher)
        response, content = self._download(self.urls["docx"])
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertEqual(document.styles["Title"].font.color.rgb, RGBColor(0, 0, 0))
        self.assertIsNone(
            document.styles["Title"].element.pPr.find(qn("w:pBdr"))
        )
        self.assertEqual(document.styles["Heading 2"].font.color.rgb, RGBColor(0, 0, 0))
        first_question = next(
            index for index, paragraph in enumerate(paragraphs) if paragraph == "第1题"
        )

        self.assertTrue(any("替换为当前课程" in paragraph for paragraph in paragraphs[:first_question]))
        self.assertTrue(any("不会创建课程 KC" in paragraph for paragraph in paragraphs[:first_question]))
        self.assertEqual(
            [paragraph for paragraph in paragraphs if paragraph.startswith("第") and paragraph.endswith("题")],
            [f"第{index}题" for index in range(1, 7)],
        )
        self.assertEqual(
            [
                paragraph.removeprefix("【题型】")
                for paragraph in paragraphs
                if paragraph.startswith("【题型】")
            ],
            QUESTION_TYPES,
        )
        self.assertIn("【选项A】2\n（质数）", paragraphs)
        self.assertIn("【解析】2、5 是质数。\n4、6 是合数。", paragraphs)
        labeled_paragraphs = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("【")
        ]
        self.assertTrue(labeled_paragraphs)
        self.assertTrue(
            all(paragraph.runs[0].bold for paragraph in labeled_paragraphs)
        )
        question_indices = [
            index
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text.startswith("第") and paragraph.text.endswith("题")
        ]
        for start, end in zip(question_indices, question_indices[1:] + [len(document.paragraphs)]):
            with self.subTest(question=document.paragraphs[start].text):
                self.assertTrue(
                    all(
                        paragraph.paragraph_format.keep_with_next
                        for paragraph in document.paragraphs[start : end - 1]
                    )
                )

    def test_excel_and_word_templates_require_teacher_or_admin(self):
        for url in self.urls.values():
            with self.subTest(url=url, account="anonymous"):
                self.assertEqual(self.client.get(url).status_code, 401)

        self.client.force_login(self.student)
        for url in self.urls.values():
            with self.subTest(url=url, account="student"):
                self.assertEqual(self.client.get(url).status_code, 403)
