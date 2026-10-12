from __future__ import annotations
import csv
import io
from openpyxl import load_workbook


def _sheet_rows(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(item or "").strip() for item in rows[0]]
    return [dict(zip(headers, values)) for values in rows[1:]]


def read_uploaded_workbook(uploaded_file):
    name = uploaded_file.name.lower()
    if not name.endswith(".xlsx"):
        raise ValueError("KC v2 工作簿仅支持 .xlsx 文件")
    workbook = load_workbook(
        io.BytesIO(uploaded_file.read()),
        read_only=True,
        data_only=True,
    )
    return {sheet.title: _sheet_rows(sheet) for sheet in workbook.worksheets}


def read_uploaded_rows(uploaded_file):
    name = uploaded_file.name.lower()
    raw = uploaded_file.read()
    if name.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    if name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        return _sheet_rows(sheet)
    raise ValueError("仅支持 .csv 或 .xlsx 文件")
