from __future__ import annotations

import io
import re
from zipfile import BadZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError


CANONICAL_KEYS = (
    "序号",
    "题目",
    "题型",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "选项E",
    "选项F",
    "答案",
    "可接受答案",
    "解析",
    "难度",
    "一级知识点",
    "二级知识点",
    "三级知识点",
)

_FIELD_ALIASES = {
    "序号": "序号",
    "题目": "题目",
    "题干": "题目",
    "题型": "题型",
    "答案": "答案",
    "可接受答案": "可接受答案",
    "其他可接受答案": "可接受答案",
    "解析": "解析",
    "难度": "难度",
    "知识点": "一级知识点",
    "一级知识点": "一级知识点",
    "二级知识点": "二级知识点",
    "三级知识点": "三级知识点",
    "主要知识点": "一级知识点",
    "主知识点": "一级知识点",
    "主要": "一级知识点",
    "次要知识点": "二级知识点",
    "次知识点": "二级知识点",
    "次要": "二级知识点",
}
_BARE_LABELS = "|".join(
    re.escape(label) for label in sorted(_FIELD_ALIASES, key=len, reverse=True)
)
_BRACKETED_LABEL = re.compile(r"[【\[［](?P<label>[^】\]］]+)[】\]］]")
_BARE_LABEL = re.compile(rf"(?<!\S)(?P<label>{_BARE_LABELS})[ \t]*[:：]")
_OPTION_LABEL = re.compile(
    r"(?P<label>(?:(?<![^\n])[A-F]|(?<!\S)选项[ \t]*[A-F]))[ \t]*(?:[:：]|[.．、])"
)
_HEADING_BOUNDARY = re.compile(
    r"^\s*第\s*(?P<sequence>\d+)\s*题[ \t:：.．、]*(?P<rest>[\s\S]*)$"
)
_NUMBERED_BOUNDARY = re.compile(
    r"^\s*(?P<sequence>\d+)\s*[.．、]\s*(?P<rest>[\s\S]+)$"
)
_TEMPLATE_TITLE = "题目与 KC 联合导入模板"


def _canonical_label(label: str) -> str | None:
    compact = label.strip()
    if (compact[:1], compact[-1:]) in {
        ("【", "】"),
        ("[", "]"),
        ("［", "］"),
    }:
        compact = compact[1:-1].strip()
    if compact in _FIELD_ALIASES:
        return _FIELD_ALIASES[compact]
    option = re.fullmatch(r"(?:选项[ \t]*)?([A-F])", compact)
    if option:
        return f"选项{option.group(1)}"
    return None


def _tag_matches(text: str) -> list[tuple[int, int, str | None, str]]:
    matches: list[tuple[int, int, str | None, str]] = []
    for match in _BRACKETED_LABEL.finditer(text):
        end = match.end()
        while end < len(text) and text[end] in " \t":
            end += 1
        if end < len(text) and text[end] in ":：":
            end += 1
            while end < len(text) and text[end] in " \t":
                end += 1
        label = match.group("label").strip()
        field = _canonical_label(label)
        if field is None and not any("\u4e00" <= char <= "\u9fff" for char in label):
            continue
        matches.append((match.start(), end, field, label))
    for pattern in (_BARE_LABEL, _OPTION_LABEL):
        for match in pattern.finditer(text):
            if any(match.start() < existing_end and match.end() > existing_start for existing_start, existing_end, _, _ in matches):
                continue
            label = match.group("label").strip()
            end = match.end()
            while end < len(text) and text[end] in " \t":
                end += 1
            matches.append((match.start(), end, _canonical_label(label), label))
    return sorted(matches, key=lambda item: item[0])


def _starts_with_known_field(text: str) -> bool:
    stripped_offset = len(text) - len(text.lstrip())
    matches = _tag_matches(text)
    return bool(
        matches
        and matches[0][0] == stripped_offset
        and matches[0][2] is not None
    )


def _has_structured_marker(text: str) -> bool:
    return bool(
        _HEADING_BOUNDARY.match(text)
        or _NUMBERED_BOUNDARY.match(text)
        or _tag_matches(text)
    )


def _is_template_preamble(text: str) -> bool:
    stripped = text.strip()
    return stripped == _TEMPLATE_TITLE or stripped.startswith(("填写说明：", "填写说明:"))


class _QuestionRow:
    def __init__(self, location: str, sequence: str = "") -> None:
        self.row = {key: "" for key in CANONICAL_KEYS}
        self.row["_errors"] = []
        self.locations = [location]
        self.active_field: str | None = None
        self.seen_fields: set[str] = set()
        if sequence:
            self.row["序号"] = sequence
            self.seen_fields.add("序号")

    def _add_location(self, location: str) -> None:
        if location not in self.locations:
            self.locations.append(location)

    def add_error(self, message: str, location: str) -> None:
        self._add_location(location)
        if message not in self.row["_errors"]:
            self.row["_errors"].append(message)

    def assign(self, field: str, value: str, location: str) -> None:
        self._add_location(location)
        if field in self.seen_fields:
            self.add_error(f"重复字段：{field}", location)
            if value:
                self._append_to(field, value)
        else:
            self.row[field] = value
            self.seen_fields.add(field)
        self.active_field = field

    def append(self, value: str, location: str) -> None:
        self._add_location(location)
        if not value:
            return
        field = self.active_field or "题目"
        self._append_to(field, value)
        self.active_field = field

    def _append_to(self, field: str, value: str) -> None:
        if self.row[field]:
            self.row[field] = f"{self.row[field]}\n{value}"
        else:
            self.row[field] = value

    def as_dict(self, fallback_sequence: str) -> dict[str, object]:
        if not self.row["序号"].strip():
            self.row["序号"] = fallback_sequence
        self.row["_location"] = "；".join(self.locations)
        return self.row


def _process_text(question: _QuestionRow, text: str, location: str) -> None:
    matches = _tag_matches(text)
    if not matches:
        question.append(text, location)
        return

    prefix = text[: matches[0][0]].strip(" \t")
    if prefix:
        question.append(prefix, location)
    for index, (start, end, field, label) in enumerate(matches):
        next_start = matches[index + 1][0] if index + 1 < len(matches) else len(text)
        value = text[end:next_start].strip(" \t")
        if field is None:
            question.add_error(f"无法识别的字段：{label}", location)
            question.append(text[start:next_start].strip(" \t"), location)
        elif field == "一级知识点" and label == "知识点" and not value:
            continue
        else:
            question.assign(field, value, location)


def _body_blocks(document):
    paragraph_number = 0
    table_number = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph_number += 1
            yield "paragraph", Paragraph(child, document), paragraph_number
        elif child.tag == qn("w:tbl"):
            table_number += 1
            yield "table", Table(child, document), table_number


def _cell_text(cell) -> str:
    return "\n".join(paragraph.text for paragraph in cell.paragraphs)


def _header_fields(table) -> list[str | None] | None:
    if len(table.rows) < 2:
        return None
    fields = [_canonical_label(_cell_text(cell).strip()) for cell in table.rows[0].cells]
    non_option_fields = [
        field for field in fields if field is not None and not field.startswith("选项")
    ]
    return fields if len(non_option_fields) >= 2 else None


def _unsupported_object_errors(element) -> list[str]:
    errors: list[str] = []
    for child in element.iter():
        local_name = child.tag.rsplit("}", 1)[-1]
        if local_name in {"drawing", "pict", "imagedata"}:
            message = "包含不支持的图片对象"
        elif local_name in {"oMath", "oMathPara"}:
            message = "包含不支持的公式对象"
        elif local_name in {"txbxContent", "wsp", "textbox"}:
            message = "包含不支持的文本框对象"
        elif local_name in {"object", "OLEObject", "altChunk"}:
            message = "包含不支持的嵌入对象"
        else:
            continue
        if message not in errors:
            errors.append(message)
    return errors


def _add_unsupported_object_errors(question: _QuestionRow, element, location: str) -> None:
    for message in _unsupported_object_errors(element):
        question.add_error(message, location)


def read_docx_rows(uploaded_file) -> list[dict[str, object]]:
    name = getattr(uploaded_file, "name", "")
    if name and not name.lower().endswith(".docx"):
        raise ValueError("仅支持 .docx 文件")
    try:
        document = Document(io.BytesIO(uploaded_file.read()))
    except (BadZipFile, KeyError, OSError, PackageNotFoundError, XMLSyntaxError) as error:
        raise ValueError("无法读取 DOCX 文档") from error
    rows: list[dict[str, object]] = []
    question: _QuestionRow | None = None
    pending_preamble_errors: list[tuple[str, str]] = []

    def finish_question(current: _QuestionRow) -> None:
        rows.append(current.as_dict(str(len(rows) + 1)))

    def start_question(location: str, sequence: str = "") -> _QuestionRow:
        nonlocal question
        if question is not None:
            finish_question(question)
        question = _QuestionRow(location, sequence)
        for message, error_location in pending_preamble_errors:
            question.add_error(message, error_location)
        pending_preamble_errors.clear()
        return question

    def process_paragraph_block(
        text: str,
        element,
        location: str,
        *,
        allow_template_preamble: bool,
    ) -> None:
        nonlocal question
        object_errors = _unsupported_object_errors(element)
        if (
            allow_template_preamble
            and question is None
            and _is_template_preamble(text)
        ):
            pending_preamble_errors.extend((message, location) for message in object_errors)
            return
        if not text.strip() and not object_errors:
            return
        heading = _HEADING_BOUNDARY.match(text)
        if heading:
            question = start_question(location, heading.group("sequence"))
            rest = heading.group("rest")
            if rest:
                _process_text(question, rest, location)
            _add_unsupported_object_errors(question, element, location)
            return
        numbered = _NUMBERED_BOUNDARY.match(text)
        if numbered:
            question = start_question(location, numbered.group("sequence"))
            rest = numbered.group("rest")
            if _starts_with_known_field(rest):
                _process_text(question, rest, location)
            else:
                question.add_error("题目边界不明确", location)
                question.active_field = None
                _process_text(question, rest, location)
            _add_unsupported_object_errors(question, element, location)
            return
        if question is None:
            question = start_question(location)
        _process_text(question, text, location)
        _add_unsupported_object_errors(question, element, location)

    for block_type, block, block_number in _body_blocks(document):
        if block_type == "table":
            header_fields = _header_fields(block)
            if header_fields is not None:
                header_cells = block.rows[0].cells
                headers = [_cell_text(cell).strip() for cell in header_cells]
                for row_number, row in enumerate(block.rows[1:], start=2):
                    values = [_cell_text(cell) for cell in row.cells]
                    if not any(value.strip() for value in values):
                        continue
                    location = f"表格{block_number}第{row_number}行"
                    question = start_question(location)
                    for column_number, cell in enumerate(header_cells, start=1):
                        _add_unsupported_object_errors(
                            question,
                            cell._tc,
                            f"表格{block_number}第1行第{column_number}列",
                        )
                    for column_number, (field, header, value) in enumerate(
                        zip(header_fields, headers, values), start=1
                    ):
                        cell_location = (
                            f"表格{block_number}第{row_number}行第{column_number}列"
                        )
                        if field is None:
                            if value.strip():
                                question.add_error(
                                    f"无法识别的字段：{header}",
                                    cell_location,
                                )
                            _add_unsupported_object_errors(
                                question, row.cells[column_number - 1]._tc, cell_location
                            )
                            continue
                        question.assign(
                            field,
                            value,
                            cell_location,
                        )
                        _add_unsupported_object_errors(
                            question, row.cells[column_number - 1]._tc, cell_location
                        )
                continue
            field_value_table = any(
                _canonical_label(_cell_text(row.cells[column_number]).strip()) is not None
                for row in block.rows
                for column_number in range(0, len(row.cells) - 1, 2)
            )
            if not field_value_table:
                first_location = ""
                has_content = False
                has_structure = False
                for row_number, row in enumerate(block.rows, start=1):
                    for column_number, cell in enumerate(row.cells, start=1):
                        for paragraph_number, paragraph in enumerate(
                            cell.paragraphs, start=1
                        ):
                            location = (
                                f"表格{block_number}第{row_number}行"
                                f"第{column_number}列第{paragraph_number}段"
                            )
                            text = paragraph.text
                            object_errors = _unsupported_object_errors(paragraph._p)
                            if not text.strip() and not object_errors:
                                continue
                            first_location = first_location or location
                            has_content = True
                            has_structure = has_structure or _has_structured_marker(text)
                            process_paragraph_block(
                                text,
                                paragraph._p,
                                location,
                                allow_template_preamble=False,
                            )
                if has_content and not has_structure:
                    if question is None:
                        question = start_question(first_location)
                    question.add_error("无法识别的表格结构", first_location)
                continue
            for row_number, row in enumerate(block.rows, start=1):
                cells = row.cells
                for column_number in range(0, len(cells) - 1, 2):
                    label = _cell_text(cells[column_number]).strip()
                    field = _canonical_label(label)
                    location = (
                        f"表格{block_number}第{row_number}行第{column_number + 1}列"
                    )
                    if field is None:
                        if not label:
                            continue
                        if question is None:
                            question = start_question(location)
                        value = _cell_text(cells[column_number + 1])
                        question.add_error(f"无法识别的字段：{label}", location)
                        question.append(f"{label}：{value}", location)
                        _add_unsupported_object_errors(
                            question, cells[column_number]._tc, location
                        )
                        _add_unsupported_object_errors(
                            question,
                            cells[column_number + 1]._tc,
                            f"表格{block_number}第{row_number}行第{column_number + 2}列",
                        )
                        continue
                    if question is None:
                        question = start_question(location)
                    question.assign(field, _cell_text(cells[column_number + 1]), location)
                    _add_unsupported_object_errors(
                        question, cells[column_number]._tc, location
                    )
                    _add_unsupported_object_errors(
                        question,
                        cells[column_number + 1]._tc,
                        f"表格{block_number}第{row_number}行第{column_number + 2}列",
                    )
            continue

        process_paragraph_block(
            block.text,
            block._p,
            f"第{block_number}段",
            allow_template_preamble=True,
        )

    if question is not None:
        finish_question(question)
    if pending_preamble_errors:
        raise ValueError("模板前言包含不支持的对象")
    if not rows:
        raise ValueError("文档中没有可解析的题目")
    return rows
