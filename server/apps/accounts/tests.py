from apps.api.test_support import domain_payload
from io import BytesIO
import json
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.audit.models import AuditLog
from .models import ClassGroup, User
from .student_import import parse_student_workbook





class StudentBatchImportCoreTests(TestCase):
    template_url = "/api/v1/teacher/accounts/students/import-template.xlsx"
    preview_url = "/api/v1/teacher/accounts/students/import/preview/"
    confirm_url = "/api/v1/teacher/accounts/students/import/confirm/"

    def setUp(self):
        self.teacher = User.objects.create_user(
            username="teacher",
            password="x",
            role=User.Role.TEACHER,
        )
        self.class_group = ClassGroup.objects.create(
            owner=self.teacher,
            academic_year="2026",
            name="认知追踪实验班",
        )
        self.client.force_login(self.teacher)

    def _upload(self, headers, rows, name="学生名单.xlsx"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        stream = BytesIO()
        workbook.save(stream)
        return SimpleUploadedFile(
            name,
            stream.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_downloads_unicode_template_with_instructions_and_owned_classes(self):
        response = self.client.get(self.template_url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("filename*=UTF-8''", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        self.assertEqual(workbook.sheetnames, ["学生导入", "填写说明"])
        self.assertEqual(
            [cell.value for cell in next(workbook["学生导入"].iter_rows(max_row=1))],
            ["学号", "姓名", "班级"],
        )
        instructions = "\n".join(
            str(cell.value or "") for row in workbook["填写说明"].iter_rows() for cell in row
        )
        self.assertIn("2026 认知追踪实验班", instructions)
        self.assertIn("S + 学号", instructions)
        self.assertIn("填写示例", instructions)
        self.assertIn("20260001", instructions)
        self.assertIn("示例学生", instructions)
        self.assertEqual(workbook["学生导入"].column_dimensions["A"].number_format, "@")
        self.assertEqual(workbook["学生导入"].column_dimensions["B"].number_format, "@")
        self.assertEqual(workbook["学生导入"].column_dimensions["C"].number_format, "@")

    def test_template_excludes_inactive_classes_from_its_available_class_list(self):
        inactive_class = ClassGroup.objects.create(
            owner=self.teacher,
            academic_year="2026",
            name="停用导入班",
            is_active=False,
        )

        response = self.client.get(self.template_url)

        workbook = load_workbook(BytesIO(response.content), data_only=True)
        instructions = "\n".join(
            str(cell.value or "") for row in workbook["填写说明"].iter_rows() for cell in row
        )
        self.assertIn(str(self.class_group), instructions)
        self.assertNotIn(str(inactive_class), instructions)

    def test_previews_chinese_rows_with_columns_in_any_order(self):
        upload = self._upload(
            ["班级", "姓名", "学号"],
            [["2026 认知追踪实验班", "林晓雨", "00123"]],
        )

        response = self.client.post(self.preview_url, {"file": upload})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            domain_payload(response)["preview"],
            [
                {
                    "student_number": "00123",
                    "display_name": "林晓雨",
                    "class_id": self.class_group.pk,
                    "class_name": "2026 认知追踪实验班",
                    "account_action": "create",
                }
            ],
        )
        self.assertFalse(User.objects.filter(student_number="00123").exists())

    def test_confirm_creates_students_atomically_with_documented_password(self):
        upload = self._upload(
            ["学号", "姓名", "班级"],
            [
                ["00123", "林晓雨", "2026 认知追踪实验班"],
                ["A-02", "王强", "2026 认知追踪实验班"],
            ],
        )
        preview = self.client.post(self.preview_url, {"file": upload})
        self.assertEqual(preview.status_code, 200)

        response = self.client.post(
            self.confirm_url,
            data=json.dumps({}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(domain_payload(response), {"imported": 2})
        first = User.objects.get(student_number="00123")
        second = User.objects.get(student_number="A-02")
        self.assertEqual(first.display_name, "林晓雨")
        self.assertEqual(first.class_group, self.class_group)
        self.assertTrue(first.check_password("S00123"))
        self.assertTrue(second.check_password("SA-02"))
        log = AuditLog.objects.get(action="student.batch_import")
        self.assertEqual(log.details, {"count": 2})
        self.assertNotIn("password", str(log.details).lower())


class StudentBatchImportValidationTests(StudentBatchImportCoreTests):
    def _assert_upload_error(self, headers, rows, expected):
        before = User.objects.filter(role=User.Role.STUDENT).count()
        response = self.client.post(self.preview_url, {"file": self._upload(headers, rows)})
        self.assertEqual(response.status_code, 400)
        self.assertIn(expected, str(domain_payload(response)))
        self.assertEqual(User.objects.filter(role=User.Role.STUDENT).count(), before)

    def test_rejects_duplicate_existing_unknown_class_and_blank_fields(self):
        User.objects.create_user(
            username="EXISTING",
            student_number="EXISTING",
            display_name="已有学生",
            class_group=self.class_group,
            role=User.Role.STUDENT,
        )
        cases = [
            (
                ["学号", "姓名", "班级"],
                [["DUP", "甲", str(self.class_group)], ["DUP", "乙", str(self.class_group)]],
                "学号重复：DUP",
            ),
            (["学号", "姓名", "班级"], [["EXISTING", "甲", str(self.class_group)]], "学号已存在"),
            (["学号", "姓名", "班级"], [["NEW", "甲", "不存在的班级"]], "班级无法匹配"),
            (["学号", "姓名", "班级"], [["NEW", "", str(self.class_group)]], "存在空字段"),
        ]
        for headers, rows, expected in cases:
            with self.subTest(expected=expected):
                self._assert_upload_error(headers, rows, expected)

    def test_multi_class_legacy_import_still_requires_the_full_class_name(self):
        ClassGroup.objects.create(
            owner=self.teacher,
            academic_year="2026",
            name="第二实验班",
        )

        self._assert_upload_error(
            ["学号", "姓名", "班级"],
            [["STRICT-CLASS", "严格匹配学生", self.class_group.name]],
            "班级无法匹配",
        )

    def test_rejects_missing_duplicate_unknown_headers_and_non_xlsx_files(self):
        cases = [
            (["学号", "姓名"], [["S1", "甲"]]),
            (["学号", "姓名", "班级", "学号"], [["S1", "甲", str(self.class_group), "S1"]]),
            (["学号", "姓名", "班级", "备注"], [["S1", "甲", str(self.class_group), "x"]]),
        ]
        for headers, rows in cases:
            with self.subTest(headers=headers):
                self._assert_upload_error(headers, rows, "表头必须且只能包含")

        response = self.client.post(
            self.preview_url,
            {"file": SimpleUploadedFile("名单.csv", b"student,name,class")},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("仅支持 .xlsx 文件", str(domain_payload(response)))

    def test_rejects_values_longer_than_model_fields(self):
        self._assert_upload_error(
            ["学号", "姓名", "班级"],
            [["S" * 65, "甲", str(self.class_group)]],
            "学号不能超过 64 个字符",
        )
        self._assert_upload_error(
            ["学号", "姓名", "班级"],
            [["NEW", "甲" * 81, str(self.class_group)]],
            "姓名不能超过 80 个字符",
        )

    def test_confirm_revalidates_conflicts_without_partial_writes(self):
        upload = self._upload(
            ["学号", "姓名", "班级"],
            [
                ["RACE", "并发学生", str(self.class_group)],
                ["SAFE", "不应写入", str(self.class_group)],
            ],
        )
        self.client.post(self.preview_url, {"file": upload})
        User.objects.create_user(
            username="RACE",
            student_number="RACE",
            display_name="先创建",
            class_group=self.class_group,
            role=User.Role.STUDENT,
        )

        response = self.client.post(
            self.confirm_url,
            data=json.dumps({}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("部分学号已存在", str(domain_payload(response)))
        self.assertFalse(User.objects.filter(student_number="SAFE").exists())
        self.assertEqual(User.objects.filter(student_number="RACE").count(), 1)

    def test_direct_legacy_import_preview_rejects_an_inactive_class(self):
        self.class_group.is_active = False
        self.class_group.save(update_fields=["is_active"])

        response = self.client.post(
            self.preview_url,
            {
                "file": self._upload(
                    ["学号", "姓名", "班级"],
                    [["INACTIVE-PREVIEW", "停用预览学生", str(self.class_group)]],
                )
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("班级已停用，不能导入学生", str(domain_payload(response)))
        self.assertFalse(User.objects.filter(student_number="INACTIVE-PREVIEW").exists())

    def test_direct_legacy_import_confirm_rejects_a_class_deactivated_after_preview(self):
        preview = self.client.post(
            self.preview_url,
            {
                "file": self._upload(
                    ["学号", "姓名", "班级"],
                    [["INACTIVE-CONFIRM", "停用确认学生", str(self.class_group)]],
                )
            },
        )
        self.assertEqual(preview.status_code, 200)
        self.class_group.is_active = False
        self.class_group.save(update_fields=["is_active"])

        response = self.client.post(
            self.confirm_url,
            data=json.dumps({}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("班级已停用，不能导入学生", str(domain_payload(response)))
        self.assertFalse(User.objects.filter(student_number="INACTIVE-CONFIRM").exists())

    def test_student_role_cannot_download_template_or_open_import(self):
        student = User.objects.create_user(
            username="STUDENT",
            student_number="STUDENT",
            role=User.Role.STUDENT,
        )
        self.client.force_login(student)

        self.assertEqual(self.client.get(self.template_url).status_code, 403)
        self.assertEqual(
            self.client.post(
                self.preview_url,
                {"file": self._upload(["学号", "姓名", "班级"], [])},
            ).status_code,
            403,
        )

    def test_failed_new_upload_invalidates_previous_preview(self):
        valid = self._upload(
            ["学号", "姓名", "班级"],
            [["STALE", "旧预览学生", str(self.class_group)]],
        )
        self.client.post(self.preview_url, {"file": valid})
        invalid = SimpleUploadedFile("错误名单.csv", b"student,name,class")
        failed_upload = self.client.post(self.preview_url, {"file": invalid})
        self.assertEqual(failed_upload.status_code, 400)
        self.assertIn("仅支持 .xlsx 文件", str(domain_payload(failed_upload)))

        response = self.client.post(
            self.confirm_url,
            data=json.dumps({}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("导入预览已失效", str(domain_payload(response)))
        self.assertFalse(User.objects.filter(student_number="STALE").exists())

    def test_existing_account_check_uses_constant_query_count(self):
        upload = self._upload(
            ["学号", "姓名", "班级"],
            [
                ["B001", "甲", str(self.class_group)],
                ["B002", "乙", str(self.class_group)],
                ["B003", "丙", str(self.class_group)],
            ],
        )

        with self.assertNumQueries(2):
            parsed = parse_student_workbook(
                upload,
                ClassGroup.objects.filter(pk=self.class_group.pk),
            )

        self.assertEqual(len(parsed), 3)
