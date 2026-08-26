from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import ClassGroup, Course, CourseEnrollment, User
from apps.assessment.models import ExperimentSession
from apps.experiments.models import (
    Experiment,
    ExperimentQuestion,
    ExperimentVersion,
)
from apps.experiments.services import assign_students
from apps.question_bank.models import Concept, Question, QuestionBank, QuestionConcept
from apps.question_bank.services.course_question_identity import allocate_course_question_identity


class CourseWorkspaceUIParityTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(
            username="teacher",
            password="password123",
            role=User.Role.TEACHER,
            display_name="课程负责人",
        )
        self.course = Course.objects.create(
            owner=self.teacher,
            name="认知诊断与知识追踪研究方法课程",
            code="KT-RESEARCH-2026",
            term="2026 秋季学期",
            description="用于课程内实验与数据采集。",
        )
        self.class_group = ClassGroup.objects.create(
            course=self.course,
            owner=self.teacher,
            name="研究一班",
            academic_year="2026",
        )
        self.students = []
        for number, name in (("S001", "张晓明"), ("S002", "李华")):
            student = User.objects.create_user(
                username=number,
                student_number=number,
                display_name=name,
                role=User.Role.STUDENT,
            )
            CourseEnrollment.objects.create(
                course=self.course,
                class_group=self.class_group,
                student=student,
            )
            self.students.append(student)
        self.bank = QuestionBank.objects.create(
            course=self.course,
            created_by=self.teacher,
            owner=self.teacher,
            name="知识追踪题库",
            description="主线题库说明",
            visibility=QuestionBank.Visibility.SHARED,
        )
        self.question = Question.objects.create(
            bank=self.bank,
            sequence=1,
            question_type=Question.Type.SINGLE,
            text="2 + 2 等于多少？",
            correct_answer="B",
            teacher_difficulty="0.250",
        )
        self.question_identity = allocate_course_question_identity(self.course, self.question)
        concept = Concept.objects.create(course=self.course, name="整数加法")
        QuestionConcept.objects.create(
            question=self.question,
            concept=concept,
            is_primary=True,
        )
        self.draft_experiment = Experiment.objects.create(
            course=self.course,
            created_by=self.teacher,
            owner=self.teacher,
            question_bank=self.bank,
            name="诊断实验",
            description="实验说明",
            ordering_mode=Experiment.Ordering.RANDOM,
            feedback_mode=Experiment.Feedback.FULL,
            time_mode=Experiment.TimeMode.TOTAL,
            total_time_limit_minutes=30,
        )
        ExperimentQuestion.objects.create(
            experiment=self.draft_experiment,
            question=self.question,
            display_order=1,
        )
        assign_students(self.draft_experiment, [self.students[0]])
        self.published_experiment = Experiment.objects.create(
            course=self.course,
            created_by=self.teacher,
            owner=self.teacher,
            question_bank=self.bank,
            name="已发布实验",
        )
        assign_students(self.published_experiment, [self.students[0]])
        self.published_experiment.status = Experiment.Status.PUBLISHED
        self.published_experiment.save(update_fields=["status"])
        self.version = ExperimentVersion.objects.create(
            experiment=self.published_experiment,
            version_number=1,
            published_by=self.teacher,
            course_name_snapshot=self.course.name,
            course_code_snapshot=self.course.code,
            course_term_snapshot=self.course.term,
        )
        self.session = ExperimentSession.objects.create(
            version=self.version,
            student=self.students[0],
            class_snapshot_name=str(self.class_group),
            submission_count=3,
        )
        self.client.force_login(self.teacher)

    def test_course_header_keeps_identity_separate_and_marks_active_navigation(self):
        response = self.client.get(reverse("courses:banks", args=[self.course.pk]))

        self.assertContains(response, 'class="teacher-course-context"')
        self.assertContains(response, self.course.name)
        self.assertContains(response, 'aria-current="page" href="{}">题库</a>'.format(
            reverse("courses:banks", args=[self.course.pk])
        ))
        self.assertNotContains(response, 'class="nav-link disabled"')

    def test_course_overview_preserves_dashboard_metrics_and_recent_resources(self):
        response = self.client.get(reverse("courses:detail", args=[self.course.pk]))

        self.assertContains(response, "课程概况")
        self.assertContains(response, "题库数量")
        self.assertContains(response, "班级数量")
        self.assertContains(response, "实验数量")
        self.assertContains(response, "近期实验")
        self.assertContains(response, "数据中心")
        self.assertContains(response, reverse("courses:data", args=[self.course.pk]))
        self.assertContains(response, self.draft_experiment.name)
        self.assertContains(response, "最近题库")
        self.assertContains(response, self.bank.name)

    def test_course_bank_pages_preserve_visibility_and_question_metadata_controls(self):
        list_response = self.client.get(reverse("courses:banks", args=[self.course.pk]))
        detail_response = self.client.get(
            reverse("courses:bank_detail", args=[self.course.pk, self.bank.pk])
        )

        self.assertContains(list_response, "data-visibility-control")
        self.assertContains(
            list_response,
            reverse("courses:bank_visibility", args=[self.course.pk, self.bank.pk]),
        )
        self.assertContains(list_response, "最近更新")
        self.assertContains(detail_response, "两个导入入口用途不同")
        self.assertContains(detail_response, self.question.get_question_type_display())
        self.assertContains(detail_response, "0.250")
        self.assertContains(detail_response, "status-badge")

    def test_course_bank_detail_shows_course_identity_and_bank_sequence_separately(self):
        response = self.client.get(
            reverse("courses:bank_detail", args=[self.course.pk, self.bank.pk])
        )

        self.assertContains(response, "课程题号")
        self.assertContains(response, "稳定课程代码")
        self.assertContains(response, "题库内序号")
        self.assertContains(response, self.question_identity.stable_code)

    def test_course_experiment_pages_preserve_main_lifecycle_and_snapshot_detail(self):
        draft_response = self.client.get(
            reverse(
                "courses:experiment_detail",
                args=[self.course.pk, self.draft_experiment.pk],
            )
        )
        published_response = self.client.get(
            reverse(
                "courses:experiment_detail",
                args=[self.course.pk, self.published_experiment.pk],
            )
        )

        for label in ("实验生命周期", "题目顺序", "反馈", "时间", "断点续答", "发布前检查"):
            self.assertContains(draft_response, label)
        self.assertContains(draft_response, 'data-bs-target="#publish-confirm-dialog"')
        self.assertContains(published_response, "发布快照")
        self.assertContains(published_response, "版本")
        self.assertContains(published_response, self.students[0].student_number)
        self.assertContains(published_response, self.class_group.name)

    def test_course_results_keep_submission_count_and_status_badges(self):
        response = self.client.get(
            reverse(
                "courses:experiment_results",
                args=[self.course.pk, self.published_experiment.pk],
            )
        )

        self.assertContains(response, "提交题数")
        self.assertContains(response, '<td class="numeric">3</td>', html=True)
        self.assertContains(response, "status-badge")

    def test_course_question_editor_keeps_main_sections_and_cancel_action(self):
        response = self.client.get(
            reverse(
                "courses:question_edit",
                args=[self.course.pk, self.bank.pk, self.question.pk],
            )
        )

        for label in ("题干与题型", "选项", "评分规则", "难度、解析与状态"):
            self.assertContains(response, label)
        self.assertContains(
            response,
            reverse("courses:bank_detail", args=[self.course.pk, self.bank.pk]),
        )
        self.assertContains(response, "取消")

    def test_course_class_detail_filters_students_by_number_or_name(self):
        url = reverse(
            "courses:class_detail",
            args=[self.course.pk, self.class_group.pk],
        )

        response = self.client.get(url, {"q": "晓明"})

        self.assertContains(response, "张晓明")
        self.assertNotContains(response, "李华")
        self.assertContains(response, 'value="晓明"')
        self.assertContains(response, "清除筛选")

    def test_course_student_import_restores_scoped_excel_template(self):
        import_url = reverse(
            "courses:student_import",
            args=[self.course.pk, self.class_group.pk],
        )
        template_url = (
            f"/teacher/courses/{self.course.pk}/classes/{self.class_group.pk}/"
            "students/import-template.xlsx"
        )

        page = self.client.get(import_url)
        response = self.client.get(template_url)

        self.assertContains(page, template_url)
        self.assertContains(page, "下载 Excel 参考模板")
        self.assertContains(page, "上传文件")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        instructions = "\n".join(
            str(cell.value or "")
            for row in workbook["填写说明"].iter_rows()
            for cell in row
        )
        self.assertIn(self.course.name, instructions)
        self.assertIn(str(self.class_group), instructions)
